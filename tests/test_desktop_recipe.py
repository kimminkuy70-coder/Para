"""Snapshot paging and guarded metadata edits; all files are synthetic."""
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from param_manager import collate, locking, namestore, workdirs
from param_manager.desktop_recipe import DesktopRecipe


class RecipeTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.shared=self.root/'shared';self.shared.mkdir()
        self.path=workdirs.collate_path(str(self.shared),'20260921_010101')
        self.machines=[f'AOI-{i:02}' for i in range(1,21)]
        rows=[dict(PI='PI2',Recipe='R1',Zone='Surface',Alg='GlobalRTP',Parameter=f'Bright {i}',**{'비고':'before'},**{m:str(i) for m in self.machines}) for i in range(205)]
        collate.write_collation(self.path,{'PI2':collate.CollateRecipe(recipe='PI2',records=rows)},self.machines)
        self.config=self.root/'config.json';self.config.write_text(json.dumps({'save_dir':str(self.shared)}))
        self.adapter=DesktopRecipe(self.config);self.catalog=self.adapter.open()
        self.patch=patch.object(locking,'VERIFY_DELAY_SEC',0);self.patch.start();self.addCleanup(self.patch.stop)

    def _local(self):
        cfg=json.loads(self.config.read_text());cfg['local_dir']=str(self.root/'local')
        self.config.write_text(json.dumps(cfg))

    def test_zone_filter_kla_hide_and_cell_colors(self):
        from param_manager import refdata
        ip=refdata.ip_path(str(self.shared));refdata.create_blank_ip(ip)
        refdata.save_ip(ip,[{'호기':'AOI-01','IP':'','장비종류':'KLA'},{'호기':'AOI-02','IP':'','장비종류':'Camtek'}])
        cat=self.adapter.open()
        self.assertEqual(cat['zones'],{'PI2':['Surface']});self.assertEqual(cat['machine_types']['AOI-01'],'KLA')
        page=self.adapter.page({'snapshot':cat['version'],'selected_machine':'AOI-02','hide_kla':True})
        self.assertNotIn('AOI-01',page['machines']);self.assertEqual(page['machine_total'],19)
        self.assertEqual(self.adapter.page({'snapshot':cat['version'],'selected_machine':'AOI-02','zone':'Other'})['total'],0)
        self.adapter.edit({'snapshot':cat['version'],'row':0,'kind':'cell','target':'AOI-03','value':'#ffee00'})
        row=self.adapter.page({'snapshot':cat['version'],'selected_machine':'AOI-02','limit':1})['rows'][0]
        self.assertEqual(row['cells'],{'AOI-03':'#FFEE00'})
        store=json.loads((self.shared/'값확인_셀색상.json').read_text(encoding='utf-8'))
        self.assertEqual(list(store.values()),['#FFEE00'])      # tkinter key format: 5 fields + target
        self.assertEqual(len(next(iter(store)).split('\x1f')),6)
        self.adapter.edit({'snapshot':cat['version'],'row':0,'kind':'cell','target':'AOI-03','value':''})
        self.assertEqual(json.loads((self.shared/'값확인_셀색상.json').read_text(encoding='utf-8')),{})
        with self.assertRaises(ValueError):
            self.adapter.edit({'snapshot':cat['version'],'row':0,'kind':'cell','target':'../x','value':'#ffee00'})

    def test_export_filters_to_local_folder(self):
        self._local()
        out=self.adapter.export({'snapshot':self.catalog['version'],'recipes':['PI2'],'machines':['AOI-05','AOI-02'],'query':'Bright 1'})
        self.assertTrue(out['path'].startswith(str(self.root/'local')))
        wb=openpyxl.load_workbook(out['path']);self.addCleanup(wb.close)
        heads=[c.value for c in wb.active[1]]
        self.assertEqual(heads[-2:],['AOI-02','AOI-05'])          # machine order kept from the collation
        self.assertEqual(out['rows'],wb.active.max_row-1)
        for bad in ({'recipes':[],'machines':['AOI-01']},{'recipes':['X'],'machines':['AOI-01']},{'recipes':['PI2'],'machines':['AOI-99']}):
            with self.assertRaises(ValueError):self.adapter.export({'snapshot':self.catalog['version'],**bad})

    def test_recipe_delete_moves_locally_and_drops_sheet(self):
        self._local()
        run=workdirs.form_run_dir(str(self.shared),'PI2','20260920_000000')
        openpyxl.Workbook().save(Path(run)/'PI2_양식.xlsx')
        prev=self.adapter.delete_preview({'recipe':'PI2'})
        self.assertEqual((prev['versions'],prev['files']),(1,1))
        with self.assertRaises(ValueError):self.adapter.delete({'recipe':'PI2','confirm':'pi2'})
        out=self.adapter.delete({'recipe':'PI2','confirm':'PI2'})
        self.assertTrue(out['moved'].startswith(str(self.root/'local')) and Path(out['moved']).is_dir())
        self.assertEqual(out['sheets'],1)
        self.assertEqual(workdirs.list_recipes(str(self.shared)),[])
        self.assertEqual(collate.load_collation(self.path)[0],{})

    def test_row_column_bounds_search_and_baseline(self):
        data=self.adapter.page({'snapshot':self.catalog['version'],'selected_machine':'AOI-20','limit':100})
        self.assertEqual(data['total'],205);self.assertEqual(len(data['rows']),100)
        self.assertEqual(len(data['machines']),12);self.assertEqual(data['rows'][0]['value'],'0')
        data=self.adapter.page({'snapshot':self.catalog['version'],'query':'Bright 204'})
        self.assertEqual(data['total'],1)
        with self.assertRaises(ValueError):self.adapter.page({'snapshot':self.catalog['version'],'limit':101})

    def test_note_is_literal_and_values_preserved(self):
        before=collate.load_collation(self.path)[0]['PI2'][0]['AOI-01']
        self.adapter.edit({'snapshot':self.catalog['version'],'row':0,'kind':'note','value':'=SUM(1,2)'})
        wb=openpyxl.load_workbook(self.path);self.addCleanup(wb.close)
        self.assertEqual(wb['PI2']['F2'].value,'=SUM(1,2)');self.assertEqual(wb['PI2']['F2'].data_type,'s')
        self.assertEqual(collate.load_collation(self.path)[0]['PI2'][0]['AOI-01'],before)

    def test_color_persisted_in_existing_workbook_schema(self):
        next_catalog=self.adapter.edit({'snapshot':self.catalog['version'],'row':0,'kind':'color','value':'#123ABC'})
        saved=namestore.load(namestore.name_path(str(self.shared)))
        self.assertEqual(saved[0]['색상코드'],'#123ABC')
        self.assertEqual(self.adapter.page({'snapshot':next_catalog['version']})['rows'][0]['color'],'#123ABC')

    def test_stale_or_arbitrary_edit_refused(self):
        for params in ({'snapshot':'old','row':0,'kind':'note','value':'bad'},
                       {'snapshot':self.catalog['version'],'row':0,'kind':'AOI-01','value':'bad'},
                       {'snapshot':self.catalog['version'],'row':0,'kind':'note','value':'bad','path':'/etc'}):
            with self.assertRaises(ValueError):self.adapter.edit(params)
        with open(self.path,'ab') as f:f.write(b'changed')
        with self.assertRaises(ValueError):
            self.adapter.edit({'snapshot':self.catalog['version'],'row':0,'kind':'note','value':'bad'})

    def test_foreign_lock_preserves_document(self):
        before=Path(self.path).read_bytes()
        with patch.object(locking,'acquire',return_value=locking.LockState('other',False)):
            with self.assertRaises(ValueError):
                self.adapter.edit({'snapshot':self.catalog['version'],'row':0,'kind':'note','value':'bad'})
        self.assertEqual(Path(self.path).read_bytes(),before)


if __name__=='__main__':unittest.main()
