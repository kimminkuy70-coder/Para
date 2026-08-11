"""레시피 삭제 — **어디까지 지우는가**를 코드로 고정한다(2026-08 사용자 확정).

  지움   ① `양식/{레시피}/` 전체 → 지우지 않고 **로컬 삭제보관으로 이동**(되돌리기)
         ② **최신** 취합본의 그 레시피 시트(값 확인 화면에서 사라지게)
         ③ `감시설정.json` 의 그 레시피(호기별 Job 폴더 지정·목록·plan)
  안 지움 ④ 과거 취합본(이력 확인의 비교 대상)
         ⑤ `변환계수.xlsx`(레시피 키가 없다 — 호기+MAG 기준이라 공용)

가장 중요한 것은 **보관 위치가 저장폴더(OneDrive) 밖**이라는 점이다. 저장폴더 안에
백업을 만들면 지운 파일이 그대로 다시 동기화돼 지운 의미가 없고 동기화만 늘어난다
(2026-08 'Unusual High-Volume Directory Access' 사고와 같은 유형).
"""

import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402

from param_manager import (collate, engine, localdirs, watcher,  # noqa: E402
                           workdirs)

PASS = FAIL = 0
SRC_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       "param_manager")


def run(fn):
    global PASS, FAIL
    print(f"[RUN] {fn.__name__}")
    try:
        fn()
        PASS += 1
        print(f"[PASS] {fn.__name__}\n")
    except Exception as e:  # noqa: BLE001
        FAIL += 1
        import traceback
        print(f"[FAIL] {fn.__name__}: {e}")
        traceback.print_exc()
        print()


def _src(name: str) -> str:
    with open(os.path.join(SRC_DIR, name), encoding="utf-8") as fh:
        return fh.read()


def _make_form(save_dir, recipe, st, aoi="AOI-13"):
    """회차 하나 흉내 — 확정본 + 관련파일 2개."""
    run_dir = workdirs.form_run_dir(save_dir, recipe, st)
    rel = workdirs.related_dir(run_dir)
    for p in (workdirs.form_final_path(run_dir, recipe, aoi, st),
              workdirs.form_original_path(rel, recipe, aoi, st),
              workdirs.form_draft_path(rel, recipe, aoi, st)):
        wb = openpyxl.Workbook()
        wb.active.append(["PI"])
        wb.save(p)
        wb.close()
    return run_dir


def _make_collate(path, recipes, machines=("AOI-13",)):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for r in recipes:
        ws = wb.create_sheet(r)
        ws.append(list(engine.META_FIELDS) + list(machines))
        ws.append([r, "PI", "Z1", "A1", "P1", "", "10"])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    wb.close()
    return path


# --------------------------------------------------------------------------
def test_preview_counts_what_will_be_deleted():
    """삭제 전에 '버전 몇 개 / 파일 몇 개 / 용량'을 정확히 보여줄 수 있어야 한다."""
    with tempfile.TemporaryDirectory() as tmp:
        save = os.path.join(tmp, "저장")
        _make_form(save, "PI3", "20260811_1000")
        _make_form(save, "PI3", "20260811_1200")
        _make_form(save, "RDL1", "20260811_1300")
        p = workdirs.recipe_delete_preview(save, "PI3")
        assert p["exists"] and p["versions"] == 2, p
        assert p["files"] == 6, p            # 회차당 확정1 + 관련파일2
        assert p["bytes"] > 0
        # 다른 레시피는 세지 않는다
        assert workdirs.recipe_delete_preview(save, "RDL1")["versions"] == 1
        # 없는 레시피는 exists=False (예외 없음)
        miss = workdirs.recipe_delete_preview(save, "없는레시피")
        assert miss["exists"] is False and miss["files"] == 0
    print("  삭제 미리보기(버전/파일/용량) OK")


def test_move_to_local_vault_and_restorable():
    """양식 폴더는 **지우지 않고 로컬로 옮긴다** — 그대로 되돌릴 수 있어야 한다."""
    with tempfile.TemporaryDirectory() as tmp:
        save = os.path.join(tmp, "저장")
        local = os.path.join(tmp, "CamtekAOI")
        _make_form(save, "PI3", "20260811_1000")
        _make_form(save, "RDL1", "20260811_1100")
        src = os.path.join(save, workdirs.FORM_DIR, "PI3")
        assert os.path.isdir(src)

        vault = localdirs.new_deleted_slot(local, "PI3")
        moved = workdirs.move_recipe_dir(save, "PI3", vault)
        assert moved == vault
        assert not os.path.exists(src), "저장폴더에서 사라져야 한다"
        assert os.path.isfile(os.path.join(
            vault, "20260811_1000",
            "PI3_AOI-13호기_참조_20260811_1000.xlsx")), "보관본에 파일이 그대로 있어야"
        # 목록에서도 빠진다(양식 만들기·값 업데이트·감시 설정이 이 목록을 쓴다)
        assert workdirs.list_recipes(save) == ["RDL1"]
        # 다른 레시피는 그대로
        assert os.path.isdir(os.path.join(save, workdirs.FORM_DIR, "RDL1"))
        # 없는 레시피를 지우면 None(예외 없음)
        assert workdirs.move_recipe_dir(
            save, "없는레시피", localdirs.new_deleted_slot(local, "x")) is None
    print("  양식 폴더 로컬 보관 이동 + 되돌리기 가능 OK")


def test_vault_must_be_outside_save_dir():
    """보관 위치가 저장폴더 안이면 **거부**한다.

    OneDrive 안에 백업을 만들면 지운 파일이 그대로 다시 동기화된다
    (지운 의미가 없고 동기화만 늘어남 — 2026-08 사고와 같은 유형).
    """
    with tempfile.TemporaryDirectory() as tmp:
        save = os.path.join(tmp, "저장")
        _make_form(save, "PI3", "20260811_1000")
        for bad in (os.path.join(save, "백업"), save,
                    os.path.join(save, "양식", "_보관")):
            try:
                workdirs.move_recipe_dir(save, "PI3", bad)
            except ValueError:
                continue
            assert False, f"저장폴더 안({bad})인데 이동을 허용했다"
        assert os.path.isdir(os.path.join(save, workdirs.FORM_DIR, "PI3")), \
            "거부됐는데 원본이 사라졌다"
        # 삭제보관 기본 위치는 로컬 폴더 아래다
        local = os.path.join(tmp, "CamtekAOI")
        assert localdirs.deleted_dir(local).startswith(local)
        assert localdirs.DELETED in localdirs.new_deleted_slot(local, "PI3")
    print("  보관 위치는 저장폴더 밖(로컬) 강제 OK")


def test_only_latest_collate_is_touched():
    """최신 취합본에서만 시트를 지운다 — 과거 취합본(이력)은 그대로."""
    with tempfile.TemporaryDirectory() as tmp:
        save = os.path.join(tmp, "저장")
        old = _make_collate(workdirs.collate_path(save, "20260810_0900"),
                            ["PI3", "RDL1"])
        new = _make_collate(workdirs.collate_path(save, "20260811_0900"),
                            ["PI3", "RDL1"])
        assert workdirs.latest_collate(save) == new
        n = collate.delete_recipe(new, "PI3")
        assert n == 1
        wb = openpyxl.load_workbook(new)
        assert wb.sheetnames == ["RDL1"], wb.sheetnames
        wb.close()
        # 과거본은 손대지 않는다(이력 확인의 비교 대상)
        wb = openpyxl.load_workbook(old)
        assert set(wb.sheetnames) == {"PI3", "RDL1"}, wb.sheetnames
        wb.close()
    print("  최신 취합본만 시트 제거, 과거 이력 보존 OK")


def test_watch_settings_are_cleaned():
    """감시 설정에서 그 레시피를 빼야 무인 회차가 없는 폴더를 찾지 않는다."""
    with tempfile.TemporaryDirectory() as tmp:
        s, st = watcher.load_settings(tmp)
        s.recipe_paths = watcher.normalize_recipe_paths({
            "AOI-11": {"PI3": r"A\Recipes\PI3", "RDL1": r"A\Recipes\RDL1"},
            "AOI-12": {"PI3": r"B\Recipes\PI3"}})
        s.plan = {"job_keyword": "PI3", "job_name": "J",
                  "recipe_map": {"PI3": ["p"], "RDL1": ["r"]},
                  "recipe_names": ["PI3", "RDL1"]}
        watcher.sync_selection(s)

        res = watcher.drop_recipe(s, "PI3")
        assert res["machines"] == ["AOI-11", "AOI-12"], res
        assert res["empty"] is False
        assert watcher.machine_recipes(s) == {"AOI-11": ["RDL1"]}, s.recipe_paths
        assert s.recipes == ["RDL1"] and s.machines == ["AOI-11"]
        assert s.plan["recipe_map"] == {"RDL1": ["r"]}
        assert s.plan["recipe_names"] == ["RDL1"]

        # 마지막 레시피까지 빼면 '대상 없음'을 알려 준다(빈 회차 방지)
        res2 = watcher.drop_recipe(s, "RDL1")
        assert res2["empty"] is True and watcher.watch_targets(s) == []
        watcher.save_settings(tmp, s, st)
        s2, _ = watcher.load_settings(tmp)
        assert watcher.machine_recipes(s2) == {}
        # 없는 레시피를 빼도 조용히 아무 일 없음
        assert watcher.drop_recipe(s2, "없는레시피")["machines"] == []
    print("  감시 설정 자동 정리(호기별 지정·목록·plan) OK")


def test_delete_scope_is_locked_in_source():
    """삭제 범위가 슬그머니 늘거나 줄지 않도록 소스 규칙으로 고정."""
    src = _src("equip_app.py")
    m = re.search(r"def _delete_recipe_run.*?(?=\n    def )", src, re.S)
    assert m, "_delete_recipe_run 을 찾지 못함"
    body = m.group(0)
    # ① 양식 폴더는 '이동'(되돌리기) — rmtree 로 지우지 않는다
    assert "move_recipe_dir" in body and "new_deleted_slot" in body, \
        "양식 폴더를 로컬 보관으로 옮기지 않음"
    assert "rmtree" not in body, "되돌릴 수 없게 바로 지우고 있음"
    # ② 최신 취합본만 (latest 한 개)
    assert "collate.delete_recipe" in body
    assert "list_collate_files" not in body, "과거 취합본까지 건드리고 있음"
    # ③ 감시 설정 정리
    assert "watcher.drop_recipe" in body and "watcher.save_settings" in body
    # ④ 변환계수는 건드리지 않는다(호기+MAG 기준이라 공용)
    assert "coefstore" not in body, "변환계수.xlsx 를 건드리고 있음"
    # ⑤ 다른 사람이 그 레시피 양식을 만드는 중이면 막는다
    assert "_acquire_global" in body and "_release_global" in body, \
        "전역 잠금 없이 삭제하고 있음"

    d = re.search(r"def _delete_recipe_dialog.*?(?=\n    def )", src, re.S)
    assert d, "_delete_recipe_dialog 을 찾지 못함"
    dbody = d.group(0)
    assert "recipe_delete_preview" in dbody, "무엇이 지워지는지 안 보여줌"
    assert 'typed.get().strip() == level' in dbody, \
        "레시피 이름 확인 입력이 없음(오삭제 방지)"
    print("  삭제 범위·안전장치 소스 규칙 OK")


def test_restore_instructions_shown_before_and_after():
    """되돌리는 법을 **삭제 전 확인창과 삭제 후 완료창 둘 다**에서 알려 준다.

    지운 게 아니라 옮긴 것이라, 어디에 있는지 모르면 되돌릴 수 없다.
    """
    src = _src("equip_app.py")
    # 안내문은 한 곳에서 만든다(두 창의 문구가 어긋나지 않게)
    helper = re.search(r"def _recipe_restore_text.*?(?=\n    def )", src, re.S)
    assert helper, "되돌리기 안내문 helper 가 없음"
    htext = helper.group(0)
    assert "옮깁니다" in htext and "다시 읽기" in htext, \
        "안내문에 '어디로 옮겼는지 / 어떻게 되돌리는지'가 없음"
    assert "workdirs.FORM_DIR" in htext, "되돌려 놓을 위치(양식 폴더)를 알려주지 않음"

    # ① 삭제 전 확인창
    d = re.search(r"def _delete_recipe_dialog.*?(?=\n    def )", src, re.S)
    assert d and "_recipe_restore_text" in d.group(0), \
        "삭제 확인창에 되돌리는 법이 없음"
    # ② 삭제 후 완료창 — 보관 경로 + 되돌리는 절차 + 폴더 열기
    w = re.search(r"def _recipe_deleted_window.*?(?=\n    def )", src, re.S)
    assert w, "삭제 완료창이 없음"
    wtext = w.group(0)
    assert "되돌리는 법" in wtext, "완료창에 되돌리는 법이 없음"
    assert "moved" in wtext and "보관 폴더 열기" in wtext, \
        "완료창에서 보관 폴더를 열어 볼 수 없음"
    assert "_open_in_excel(moved)" in wtext, "보관 폴더 열기가 동작하지 않음"
    # 완료 안내는 이 창으로 간다(요약만 던지고 끝내지 않게)
    r = re.search(r"def _delete_recipe_run.*?(?=\n    def )", src, re.S)
    assert r and "_recipe_deleted_window" in r.group(0)
    print("  되돌리는 법 안내(삭제 전·후 + 폴더 열기) OK")


def test_backup_never_written_into_save_dir():
    """보관 폴더 경로를 저장폴더에서 만들지 않는다(소스 규칙)."""
    src = _src("equip_app.py")
    m = re.search(r"def _delete_recipe_(dialog|run).*?(?=\n    def )", src, re.S)
    assert m
    both = "".join(x.group(0) for x in
                   re.finditer(r"def _delete_recipe_(?:dialog|run).*?(?=\n    def )",
                               src, re.S))
    assert "localdirs.new_deleted_slot(self.local_dir" in both, \
        "보관 위치를 로컬(local_dir)에서 만들지 않음"
    assert "new_deleted_slot(self.save_dir" not in both, \
        "저장폴더 안에 백업을 만들고 있음(OneDrive 동기화 폭주)"
    print("  보관은 로컬(local_dir)에만 OK")


if __name__ == "__main__":
    for t in [test_preview_counts_what_will_be_deleted,
              test_move_to_local_vault_and_restorable,
              test_vault_must_be_outside_save_dir,
              test_only_latest_collate_is_touched,
              test_watch_settings_are_cleaned,
              test_delete_scope_is_locked_in_source,
              test_restore_instructions_shown_before_and_after,
              test_backup_never_written_into_save_dir]:
        run(t)
    print(f"==== {PASS}/{PASS + FAIL} passed ====")
    sys.exit(1 if FAIL else 0)
