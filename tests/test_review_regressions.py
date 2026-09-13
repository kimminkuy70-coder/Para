"""Headless behavioral regressions for version_new; no equipment/network access."""
import ast
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from param_manager import atomicfile, cmwatcher, errlog, updater, watcher, wph
import openpyxl


class ReviewTests(unittest.TestCase):
    def test_log_fallback_never_writes_to_shared_folder(self):
        with tempfile.TemporaryDirectory() as d:
            shared = Path(d) / 'shared'
            shared.mkdir()
            with patch('param_manager.localdirs.active_root', side_effect=OSError), patch.dict(os.environ, LOCALAPPDATA=d):
                self.assertTrue(errlog.write_log(str(shared), 'TEST', 'fallback'))
                self.assertEqual(list(shared.iterdir()), [])
                self.assertTrue((Path(d) / 'CamtekAOI/Logs' / errlog.LOG_NAME).exists())

    def test_same_version_different_binary_preserves_release(self):
        with tempfile.TemporaryDirectory() as d:
            shared = Path(d) / 'shared'
            shared.mkdir()
            exe = Path(d) / 'built.exe'
            exe.write_bytes(b'original')
            first = updater.publish(str(shared), str(exe), '8.1')
            before = Path(updater.manifest_path(str(shared))).read_bytes()
            exe.write_bytes(b'changed')
            with self.assertRaises(ValueError):
                updater.publish(str(shared), str(exe), '8.1')
            self.assertEqual(Path(updater.manifest_path(str(shared))).read_bytes(), before)
            dest = Path(updater.program_dir(str(shared))) / first.filename
            self.assertEqual(dest.read_bytes(), b'original')

    def test_local_atomic_failure_preserves_settings(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / 'settings.json'
            p.write_text('{"old": 1}')
            with patch.object(atomicfile.os, 'replace', side_effect=PermissionError):
                with self.assertRaises(PermissionError):
                    atomicfile.write_json(p, {'new': 2})
            self.assertEqual(json.loads(p.read_text()), {'old': 1})
            self.assertEqual(list(Path(d).iterdir()), [p])

    def test_invalid_manifests_and_legacy_name(self):
        for value in (None, [], 'broken', 3):
            self.assertIsNone(updater.ReleaseInfo.from_dict(value))
        data = dict(version='8.1', sha256='a' * 64, size=10)
        self.assertEqual(updater.ReleaseInfo.from_dict(data).filename, updater.LEGACY_EXE_NAME)
        for name in ('../x.exe', 'C:\\x.exe', 'sub/x.exe', 'x.exe\n', 'x.txt'):
            self.assertIsNone(updater.ReleaseInfo.from_dict(dict(data, filename=name)))

    def test_download_read_failure_is_reported(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / 'new.exe'
            p.write_bytes(b'x')
            rel = updater.ReleaseInfo('8.1', 'new.exe', 'a' * 64, 1)
            with patch.object(updater, 'file_sha256', side_effect=PermissionError('locked')):
                self.assertFalse(updater.verify_download(str(p), rel)[0])

    def test_pending_scan_retries_without_parent_mtime_change(self):
        with tempfile.TemporaryDirectory() as d:
            lot = Path(d) / 'lot'
            sm = lot / 'SM1'
            sm.mkdir(parents=True)
            seen = set()
            with patch.object(cmwatcher, 'lot_dirs_for', return_value=[lot]):
                with patch.object(cmwatcher, '_has_config', return_value=False):
                    first = cmwatcher.scan_new([], [('device', 'lot')], seen, {}, settle_minutes=0)
                self.assertEqual(len(first['pending']), 1)
                with patch.object(cmwatcher, '_has_config', return_value=True):
                    second = cmwatcher.scan_new([], [('device', 'lot')], seen, first['mtimes'], settle_minutes=0)
                    self.assertEqual(len(second['new']), 1)
                    third = cmwatcher.scan_new([], [('device', 'lot')], seen, second['mtimes'], settle_minutes=0)
                    self.assertEqual(third['skipped'], 1)

    def test_wph_deduplicates_and_rejects_paths(self):
        with tempfile.TemporaryDirectory() as d:
            name = 'example_BatchReport.htm'
            (Path(d) / name).touch()
            self.assertEqual(wph._resolve_names(d, '', [name, name]), [name])
            self.assertEqual(wph._resolve_names(d, '', []), [])
            with self.assertRaises(ValueError):
                wph._resolve_names(d, '', ['../' + name])

    def test_wph_more_than_2000_rows_has_formulas_and_statistics(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / 'wph.xlsx'
            rows = [dict(source_file=str(i), wafers=25, avg_scan_sec=100, batch_sec=3000) for i in range(2001)]
            wph.write_wph_excel(p, rows)
            wb = openpyxl.load_workbook(p)
            try:
                ws = wb['01_Raw_Data']
                self.assertEqual(ws['A2002'].value, 2001)
                self.assertTrue(ws['J2002'].value.startswith('='))
                formulas = [c.value for s in wb if s != ws for row in s for c in row if c.data_type == 'f']
                self.assertTrue(any('2002' in f for f in formulas))
            finally:
                wb.close()

    def test_background_completion_uses_main_thread(self):
        tree = ast.parse((Path(__file__).resolve().parents[1] / 'param_manager/equip_app.py').read_text())
        cls = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == 'EquipApp')
        fn = next(n for n in cls.body if isinstance(n, ast.FunctionDef) and n.name == '_run_bg')
        ns = {'threading': threading}
        exec(compile(ast.Module(body=[fn], type_ignores=[]), '<extracted>', 'exec'), ns)
        owner = threading.get_ident()
        callbacks, results = [], []
        class FakeApp:
            def after(self, delay, callback):
                if threading.get_ident() != owner:
                    raise AssertionError('Tk called from worker')
                callbacks.append(callback)
        def done(ok, value):
            self.assertEqual(threading.get_ident(), owner)
            results.append((ok, value))
        for work, expected in ((lambda: 42, True), (lambda: 1 / 0, False)):
            ns['_run_bg'](FakeApp(), work, done)
            deadline = time.monotonic() + 3
            while callbacks and time.monotonic() < deadline:
                callbacks.pop(0)()
                time.sleep(.001)
            self.assertEqual(results[-1][0], expected)
        self.assertEqual(len(results), 2)

    def test_batch_encoding_and_backup_order(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / 'swap.bat'
            updater._write_bat(str(p), '@echo off\r\nrem 한글\n')
            self.assertNotIn(b'\r\r\n', p.read_bytes())
            self.assertIn('한글', p.read_bytes().decode('cp949'))
            p.unlink()
            if os.name != 'nt':
                with self.assertRaises(UnicodeEncodeError):
                    updater._write_bat(str(p), 'echo \U0001f600')
                self.assertFalse(p.exists())
            script = updater.build_swap_script(d, 123, r'C:\100%\old.exe', r'C:\new.exe', r'C:\backup.exe')
            text = Path(script).read_bytes().decode('cp949')
            self.assertIn('100%%', text)
            self.assertLess(text.index('fc /b'), text.index('del /q'))


if __name__ == '__main__':
    unittest.main()
