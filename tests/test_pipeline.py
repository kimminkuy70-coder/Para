"""통합 파이프라인 테스트 — GUI가 오케스트레이션하는 헤드리스 흐름 전체.

parse → 양식(formbuilder) → 취합(collate) → 새 버전(versioning) → 이력(history).
GUI(tkinter)는 이 개발환경에서 못 돌리므로, 그 뼈대 로직을 코드로 검증한다.
"""
import os
import sys
import tempfile
import shutil
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402

from param_manager import (collate, engine, formbuilder, history,  # noqa: E402
                           ini_parser, versioning)

GLOBAL_RTP = "[GLOBAL_RTP]\nMaxFaultsPerWafer = {v}\nApplyDieCalib = 1\n"
ZONE = "[General]\nZoneName = PI Opening\n[Surface]\nHigh_Delta = 25\n"


def _equip(root, equip, wafer="3000"):
    rec = Path(root) / "R_TB500_PI3" / "PI"
    rec.mkdir(parents=True, exist_ok=True)
    (rec / "GlobalRTP.ini").write_text(GLOBAL_RTP.format(v=wafer), encoding="utf-8")
    (rec / "Zones").mkdir(exist_ok=True)
    (rec / "Zones" / "Zone1.ini").write_text(ZONE, encoding="utf-8")
    cfgs = ini_parser.scan_tree(Path(root) / "R_TB500_PI3", default_equipment=equip)
    return ini_parser.build_pivot(cfgs)[0]


def test_full_pipeline():
    with tempfile.TemporaryDirectory() as tmp:
        base = os.path.join(tmp, "공용"); os.makedirs(base)
        # 1) 장비 A 파싱 → initial → (편집: 개명) → 양식 final
        rows_a = _equip(os.path.join(tmp, "A"), "AOI-13")
        init = os.path.join(base, "01_초안.xlsx")
        formbuilder.build_initial_workbook(rows_a, init, level="PI3")
        wb = openpyxl.load_workbook(init); ws = wb[formbuilder.INIT_SHEET]
        heads = [c.value for c in ws[1]]
        fi, gi = heads.index("최종 Parameter") + 1, heads.index("추천 Parameter") + 1
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, gi).value == "Max Defects Per Wafer":
                ws.cell(r, fi).value = "웨이퍼당 최대결함"
        wb.save(init); wb.close()
        form = os.path.join(base, "양식_PI3.xlsx")
        formbuilder.build_final_from_initial(init, form, level="PI3")

        # 2) 장비 B 취합 → 취합 v001 (canonical 갱신)
        canonical = os.path.join(base, "취합", "양식_PI3_PI3_취합.xlsx")
        os.makedirs(os.path.dirname(canonical), exist_ok=True)

        def collate_run(equip, wafer):
            pivot = _equip(os.path.join(tmp, equip + wafer), equip, wafer=wafer)
            res = collate.collate(form, pivot, selected_levels=["PI3"])[0]
            dest = versioning.next_version_path(canonical)
            collate.write_collated(res, dest, source=equip)
            shutil.copy2(dest, canonical)
            return dest, res

        v1, res1 = collate_run("AOI-14", "3000")
        assert res1.mismatches == []
        # 3) 같은 장비 값이 바뀐 두 번째 취합 → v002
        v2, res2 = collate_run("AOI-14", "9999")

        vers = versioning.list_versions(canonical)
        assert len(vers) == 2 and vers == [v1, v2]

        # 4) 이력 비교 v1 → v2: 개명된 파라미터의 AOI-14 값이 3000→9999
        diff = history.diff_files(v1, v2)
        chg = {(c.param, c.machine): c for c in diff.changes}
        key = ("웨이퍼당 최대결함", "AOI-14")
        assert key in chg, list(chg)
        assert chg[key].old == "3000" and chg[key].new == "9999"
        assert chg[key].kind == "값변경"
        print("  pipeline OK: 파싱→양식(개명)→취합→새버전→이력 비교 일관")


if __name__ == "__main__":
    fails = 0
    tests = [(n, f) for n, f in list(globals().items())
             if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        print(f"[RUN] {name}")
        try:
            fn(); print(f"[PASS] {name}\n")
        except Exception as e:  # noqa: BLE001
            fails += 1
            import traceback
            traceback.print_exc()
            print(f"[FAIL] {name}: {e}\n")
    print(f"==== {len(tests) - fails}/{len(tests)} passed ====")
    sys.exit(1 if fails else 0)
