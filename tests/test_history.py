"""history(재설계) 테스트 — 취합 파일 2개(멀티시트) 비교 + 변경내역 엑셀."""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402

from param_manager import collate, engine, history  # noqa: E402


def _collate_file(path, sheet_rows, machines):
    """sheet_rows: {recipe: [(PI,Recipe,Zone,Alg,Parameter, {호기:값})]}."""
    results = {}
    for recipe, rows in sheet_rows.items():
        res = collate.CollateRecipe(recipe=recipe, machines=machines)
        for pi, rc, zone, alg, param, vals in rows:
            rec = {"PI": pi, "Recipe": rc, "Zone": zone, "Alg": alg,
                   "Parameter": param, "비고": ""}
            rec.update(vals)
            res.records.append(rec)
        results[recipe] = res
    collate.write_collation(path, results, machines)


def test_diff_multisheet():
    with tempfile.TemporaryDirectory() as tmp:
        m = ["AOI-24", "AOI-25"]
        old = os.path.join(tmp, "취합_v1.xlsx")
        new = os.path.join(tmp, "취합_v2.xlsx")
        _collate_file(old, {"PI3": [
            ("PI3", "PI", "GlobalRTP", "GLOBAL_RTP", "Max Defects Per Wafer",
             {"AOI-24": "3000", "AOI-25": "3000"}),
            ("PI3", "PI", "PI Opening", "Surface", "삭제될 항목", {"AOI-24": "1"}),
        ]}, m)
        _collate_file(new, {"PI3": [
            ("PI3", "PI", "GlobalRTP", "GLOBAL_RTP", "Max Defects Per Wafer",
             {"AOI-24": "7000", "AOI-25": "3000"}),        # AOI-24 값변경
            ("PI3", "PI", "PI Opening", "Surface", "새 항목", {"AOI-24": "9"}),
        ]}, m)

        diff = history.diff_files(old, new)
        chg = {(c.param, c.machine): c for c in diff.changes}
        assert chg[("Max Defects Per Wafer", "AOI-24")].kind == "값변경"
        assert chg[("Max Defects Per Wafer", "AOI-24")].old == "3000"
        assert chg[("Max Defects Per Wafer", "AOI-24")].new == "7000"
        assert chg[("Max Defects Per Wafer", "AOI-24")].sheet == "PI3"
        assert {r["Parameter"] for r in diff.added_rows} == {"새 항목"}
        assert {r["Parameter"] for r in diff.removed_rows} == {"삭제될 항목"}

        dest = os.path.join(tmp, "변경내역.xlsx")
        history.write_diff_excel(diff, dest, old_label="v1", new_label="v2",
                                 memos={("PI3", "Max Defects Per Wafer", "AOI-24"): "상향"})
        wb = openpyxl.load_workbook(dest)
        ws = wb[history.DIFF_SHEET]
        heads = [c.value for c in ws[1]]
        memo_col = heads.index("비고(메모)")
        memos = [r[memo_col] for r in ws.iter_rows(min_row=2, values_only=True)]
        assert "상향" in memos
        wb.close()
    print("  history OK: 멀티시트 값변경/추가/삭제 + 변경내역 엑셀 + 비고")


if __name__ == "__main__":
    fails = 0
    tests = [(n, f) for n, f in list(globals().items())
             if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        print(f"[RUN] {name}")
        try:
            fn()
            print(f"[PASS] {name}\n")
        except Exception as e:  # noqa: BLE001
            fails += 1
            import traceback
            traceback.print_exc()
            print(f"[FAIL] {name}: {e}\n")
    print(f"==== {len(tests) - fails}/{len(tests)} passed ====")
    sys.exit(1 if fails else 0)
