"""Local synthetic fixture for the browser-to-real-Python integration harness."""
import json
from datetime import datetime, timedelta
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from test_batchreport import report, html_report
from param_manager.desktop_batch import DesktopBatch
from param_manager import desktop_ipc
from param_manager import collate, workdirs, refdata, commonality
import openpyxl


if sys.argv[1] == 'init':
    root = Path(sys.argv[2])
    source = root / 'equipment'
    source.mkdir()
    for i in range(205):
        start = datetime(2026, 9, 1) + timedelta(hours=i)
        status = '<script>window.injected=true</script>' if i == 0 else 'Aborted.' if i % 17 == 0 else 'Pass'
        rep = report('BatchReport_%03d.htm' % i, statuses=(status,), count=1,
                     start=start.strftime('%d-%b-%y %I:%M:%S %p'),
                     end=(start+timedelta(minutes=10)).strftime('%d-%b-%y %I:%M:%S %p'), seconds='00:10:00')
        (source / rep['file_name']).write_text(html_report(rep), encoding='utf-8')
    cfg = {'local_dir': str(root/'local'), 'wph_report_paths': {'AOI-01':str(source),'AOI-02':str(root/'offline-02'),'AOI-03':str(root/'offline-03')},
           'batch_last': {'targets':[{'machine':'AOI-01','query':'','start':'','end':''}],
                          # M07 was retired; older saved settings must still start (A1).
                          'options':{'valid_wafers':1,'metrics':['M01','M02','M03','M07','M10']}}}
    shared=root/'shared';shared.mkdir();cfg['save_dir']=str(shared)
    machines=[f'AOI-{i:02}' for i in range(1,21)]
    records=[dict(PI='PI2',Recipe='R1',Zone='Surface',Alg='GlobalRTP',Parameter=f'Min Defect Bright {i}',
        **{'비고':''},**{m:str(i+(j%3)) for j,m in enumerate(machines)}) for i in range(2000)]
    collate.write_collation(workdirs.collate_path(str(shared),'20260921_010101'),
        {'PI2':collate.CollateRecipe(recipe='PI2',records=records)},machines)
    # A confirmed form + editable candidate for PI2 (one LINEAR variant) for 양식 만들기.
    from param_manager import formbuilder
    st='20260920_010101_000001'
    run=workdirs.form_run_dir(str(shared),'PI2',st)
    openpyxl.Workbook().save(workdirs.form_final_path(run,'PI2','AOI-01',st))
    def pv(variant,param,key,raw,transform='RAW'):
        return dict(layer='PI2',recipe='PI2',mag=variant,zone='Surface',alg='GlobalRTP',param=param,values={},unit='',
                    raws={'양식':raw},use=True,extract={'src_file':'OpticPreset.ini','section':'Scan2d','key':key,'transform':transform,'source_path':''})
    formbuilder.build_initial_workbook([pv('R1','Min Defect Bright 0','K0','5'),pv('R1','Min Defect Bright 1','K1','6'),
        pv('R1-x20','Width','W','3.5','LINEAR')],workdirs.form_original_path(workdirs.related_dir(run),'PI2','AOI-01',st),level='PI2',source='fixture')
    path=refdata.ip_path(str(shared));refdata.create_blank_ip(path)
    wb=openpyxl.load_workbook(path);wb.active.append(['AOI-01','10.0.0.1','Camtek']);wb.save(path);wb.close()
    special=refdata.special_path(str(shared));refdata.create_blank_special(special)
    wb=openpyxl.load_workbook(special);wb.active.append(['2026-09-01','AOI-01','L1','S','PI','점검','진행','','메모']);wb.save(special);wb.close()
    run = workdirs.commonality_run_dir(str(root/'local/Commonality'), 'AOI-01', 'sample')
    cmfile = workdirs.commonality_result_path(run, 'R', 'AOI-01', 'sample')
    cmrecords = [dict(PI='R', Recipe='V', Zone='Z', Alg='A', Parameter=f'P{i}',
                      L1='1', L2='2', L3='1') for i in range(25)]
    commonality.write_lot_result(cmfile, 'R', 'AOI-01',
        collate.CollateRecipe(recipe='R', records=cmrecords), ['L1', 'L2', 'L3'],
        fail_labels=['L2'], low_labels=['L3'])
    (root/'config.json').write_text(json.dumps(cfg), encoding='utf-8')
else:
    original = desktop_ipc.Session
    class FixtureSession(original):
        def __init__(self, output):
            super().__init__(output)
            config = Path(sys.argv[2])/'config.json'
            self.batch = DesktopBatch(config)
            self.commonality.batch.config_path = config
            for adapter in (self.recipe, self.documents, self.form, self.cmsurvey,
                            self.history, self.config, self.opener):
                adapter.config_path = config
    desktop_ipc.Session = FixtureSession
    desktop_ipc.serve(sys.stdin.buffer, sys.stdout.buffer)
