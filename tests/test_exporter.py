"""exporter 테스트 — 레시피별 시트·선택 호기 열·수정값 반영·열너비/줄바꿈."""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402

from param_manager import exporter  # noqa: E402


def test_write_export_sheets_and_values():
    with tempfile.TemporaryDirectory() as tmp:
        dest = os.path.join(tmp, "내보내기.xlsx")
        recs = {
            "PI3": [
                {"PI": "PI3", "Recipe": "PI", "Zone": "Scan Area", "Alg": "Surface",
                 "Parameter": "Min Defect Area", "비고": "", "AOI-24": "123",
                 "AOI-25": "999"},
            ],
            "RDL1": [
                {"PI": "RDL1", "Recipe": "x5", "Zone": "Z", "Alg": "Uniform",
                 "Parameter": "High_Delta", "비고": "메모", "AOI-24": "7",
                 "AOI-25": "8"},
            ],
        }
        exporter.write_export(dest, recs, ["AOI-24", "AOI-25"], title="테스트")
        wb = openpyxl.load_workbook(dest)
        assert set(wb.sheetnames) == {"PI3", "RDL1"}, wb.sheetnames
        ws = wb["PI3"]
        heads = [c.value for c in ws[1]]
        assert heads[:6] == ["PI", "Recipe", "Zone", "Alg", "Parameter", "비고"]
        assert heads[-2:] == ["AOI-24", "AOI-25"]
        # 값(수정 반영된 최종 dict)이 그대로 기록
        row = [c.value for c in ws[2]]
        assert row[4] == "Min Defect Area"
        assert str(row[-2]) == "123" and str(row[-1]) == "999"
        # Parameter 열 너비가 넉넉히 잡혔는지(가독성)
        assert ws.column_dimensions["E"].width >= 40
        assert ws.freeze_panes == "A2"
    print("  exporter OK: 레시피별 시트·선택 호기·수정값·열너비/고정")


def test_write_export_subset_machine():
    """호기 1대만 선택해도 그 열만 나온다."""
    with tempfile.TemporaryDirectory() as tmp:
        dest = os.path.join(tmp, "e.xlsx")
        recs = {"PI3": [{"PI": "PI3", "Recipe": "PI", "Zone": "Z", "Alg": "S",
                         "Parameter": "P", "비고": "", "AOI-24": "1"}]}
        exporter.write_export(dest, recs, ["AOI-24"])
        wb = openpyxl.load_workbook(dest)
        heads = [c.value for c in wb["PI3"][1]]
        assert heads[-1] == "AOI-24" and "AOI-25" not in heads
    print("  exporter OK: 선택 호기 열만 포함")


if __name__ == "__main__":
    fails = 0
    tests = [(n, f) for n, f in list(globals().items())
             if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        print(f"[RUN] {name}")
        try:
            fn()
            print(f"[PASS] {name}\n")
        except Exception as e:  # noqa: BLE001
            fails += 1
            import traceback
            traceback.print_exc()
            print(f"[FAIL] {name}: {e}\n")
    print(f"==== {len(tests) - fails}/{len(tests)} passed ====")
    sys.exit(1 if fails else 0)
