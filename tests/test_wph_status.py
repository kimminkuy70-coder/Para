import sys
import tempfile
from pathlib import Path
import unittest
from datetime import datetime
from zipfile import ZipFile
from xml.etree import ElementTree
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl
from param_manager import wph, wph_status as status, wph_charts


class StatusTests(unittest.TestCase):
    def test_chronological_trends_and_readable_labels(self):
        later, earlier = datetime(2026, 9, 15, 15), datetime(2026, 9, 14, 9)
        base = dict(machine='AOI-1', wafers=25, avg_scan_sec=100)
        rows = [dict(base, batch_sec=3600, batch_end=later, source_file='a.htm'),
                dict(base, batch_sec=1800, batch_end=earlier, source_file='z.htm'),
                dict(base, batch_sec=900, source_file='missing.htm')]
        _, reports, _ = wph_charts.summarize(rows, 25)
        self.assertEqual([r[2] for r in reports], ['z.htm', 'a.htm', 'missing.htm'])
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'time.xlsx'
            wph.write_wph_excel(path, rows)
            wb = openpyxl.load_workbook(path)
            try:
                data, dash = wb['04_그래프데이터'], wb['05_대시보드']
                self.assertEqual(data['P2'].value, earlier)
                self.assertEqual(data['P3'].value, later)
                self.assertEqual(data['Q2'].value, 50)
                self.assertEqual(data['T2'].value, 30)
                self.assertIn('1건 제외', dash['A10'].value)
                for chart in dash._charts[:2]:
                    labels = chart.dataLabels
                    self.assertTrue(labels.showVal)
                    self.assertFalse(labels.showSerName)
                    self.assertFalse(labels.showCatName)
                    self.assertFalse(labels.showLegendKey)
                    self.assertGreaterEqual(chart.anchor.ext.cx, 34 * 360000)
                for chart in dash._charts[2:]:
                    self.assertIn('$P$2:$P$3', chart.series[0].xVal.numRef.f)
                    self.assertEqual(chart.x_axis.numFmt.formatCode, 'yyyy-mm-dd hh:mm')
                    self.assertFalse(chart.x_axis.numFmt.sourceLinked)
                    span = chart.x_axis.scaling.max - chart.x_axis.scaling.min
                    self.assertAlmostEqual(span / chart.x_axis.majorUnit, 4)
                for chart in dash._charts:
                    for axis in (chart.x_axis, chart.y_axis):
                        self.assertFalse(axis.delete)
                        self.assertEqual(axis.tickLblPos, 'low')
                        self.assertEqual(axis.txPr.p[0].pPr.defRPr.sz, 1100)
                        self.assertIsNotNone(axis.title)
                self.assertEqual(dash._charts[0].dataLabels.numFmt, '0.00" WPH"')
                self.assertEqual(dash._charts[1].dataLabels.numFmt, '0"건"')
                self.assertEqual(dash._charts[0].series[0].cat.strRef.strCache.pt[0].v, 'AOI-1')
                self.assertEqual(dash._charts[1].series[0].cat.strRef.strCache.ptCount, 6)
                self.assertEqual(dash._charts[2].x_axis.axPos, 'b')
                self.assertEqual([c.anchor._from.row for c in dash._charts], [11, 47, 83, 119])
            finally:
                wb.close()

    def test_error_types_stay_together_with_readable_axes_and_spacing(self):
        for count in (1, 13, 25):
            with self.subTest(count=count), tempfile.TemporaryDirectory() as d:
                rows = [dict(machine='A', source_file='a', wafer_statuses=[
                    dict(status=f'Unexpected status {i}') for i in range(count)])]
                wb = openpyxl.Workbook()
                status.add_sheets(wb, rows, [])
                path = Path(d) / 'status.xlsx'
                wb.save(path)
                wb.close()
                wb = openpyxl.load_workbook(path)
                try:
                    charts = wb['07_상태요약']._charts
                    self.assertEqual(len(charts), 2)
                    for chart, column in zip(charts, ('J', 'K')):
                        expected_range = f'${column}$2' + (f':${column}${count + 1}' if count > 1 else '')
                        self.assertIn(expected_range, chart.series[0].val.numRef.f)
                        cache = chart.series[0].cat.strRef.strCache
                        self.assertEqual(cache.ptCount, count)
                        self.assertIn(f'Unexpected status {count - 1}', cache.pt[-1].v)
                        self.assertFalse(chart.dataLabels.showCatName)
                        self.assertEqual(chart.dataLabels.numFmt, '0"건"')
                        self.assertEqual(chart.x_axis.tickLblPos, 'low')
                        self.assertFalse(chart.x_axis.delete)
                        self.assertEqual(chart.y_axis.scaling.orientation, 'minMax')
                        self.assertEqual(chart.y_axis.scaling.min, 0)
                        self.assertEqual(chart.y_axis.majorUnit, 1)
                        self.assertGreater(chart.y_axis.scaling.max, 1)
                        self.assertGreaterEqual(chart.anchor.ext.cy, max(20, 6 + count * 1.1) * 360000 - 1)
                    gap_points = (charts[1].anchor._from.row - charts[0].anchor._from.row) * 18
                    self.assertGreater(gap_points, charts[0].anchor.ext.cy / 12700)
                finally:
                    wb.close()

    def test_wph_numeric_chart_data(self):
        rows = [dict(machine='AOI-1', wafers=25, avg_scan_sec=50, batch_sec=1800),
                dict(machine='AOI-1', wafers=25, avg_scan_sec=100, batch_sec=3600),
                dict(machine='AOI-2', wafers=2, avg_scan_sec=50, batch_sec=500),
                dict(machine='AOI-2', wafers=25, avg_scan_sec=float('nan'), batch_sec=500)]
        machines, reports, bands = wph_charts.summarize(rows, 25)
        self.assertEqual(machines['AOI-1'], [2, 50, 5400])
        self.assertEqual(len(machines), 1)
        self.assertEqual([r[1] for r in reports], [50, 25])
        self.assertEqual(bands, [0, 1, 0, 1, 0, 0])

    def test_error_chart_data_matches_detail_rows(self):
        rows = [dict(machine=m, source_file='same.htm', wafer_statuses=[
                    dict(status='Aborted.'), dict(status='Aborted. + Skipped.')])
                for m in ['AOI-1', 'AOI-2']]
        rows.append(dict(source_file='unknown.htm', wafer_statuses=[]))
        wb = openpyxl.Workbook()
        status.add_sheets(wb, rows, [])
        summary = wb['07_상태요약']
        data = list(summary.iter_rows(min_row=2, max_row=3, min_col=9, max_col=11, values_only=True))
        self.assertEqual(data[0], ('작업 중단 / Aborted.', 2, 4))
        self.assertEqual(data[1], ('검사 제외 / Skipped.', 2, 2))
        self.assertEqual(sum(r[2] for r in data), sum(1 for r in wb['09_상태상세'].iter_rows(min_row=2, values_only=True) if r[4] != '상태 확인 불가'))
        wb.close()

    def test_counts_are_per_report_and_per_wafer(self):
        rows = [dict(source_file='a.htm', wafer_statuses=[
            {'status': 'Scan 2D Error. + Aborted. + Skipped.'}, {'status': 'Aborted.'}]),
            dict(source_file='b.htm', wafer_statuses=[{'status': 'Pass'}]),
            dict(source_file='c.htm', wafer_statuses=[])]
        counts, summary, reports, details = status.analyze(rows, [])
        self.assertEqual(counts, {'이슈 포함': 1, '정상 상태만': 1, '상태 확인 불가': 1})
        self.assertEqual(summary[('작업 중단', 'Aborted.')], [1, 2])
        self.assertEqual(len(details), 5)
        self.assertEqual(len(reports), 3)
        self.assertEqual(status.classify('Failed'), [('기타 상태', 'Failed')])
        self.assertEqual(status.classify('Pass'), [])

    def test_html_status_extraction(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / 'x_BatchReport.htm'
            p.write_text('<table><tr><th>No</th><th>Wafer ID</th><th>Scanned Dice</th>'
                         '<th>Yield</th><th>Pass/Fail</th></tr><tr><td>2</td><td>W2</td>'
                         '<td>0</td><td>0</td><td>Wafer Map Import failed.</td></tr></table>')
            row = wph.extract_row(wph.parse_report(p))
            self.assertEqual(row['wafer_statuses'][0], dict(slot='2', wafer_id='W2', status='Wafer Map Import failed.'))

    def test_workbook_and_chart_roundtrip(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / 'out.xlsx'
            rows = [dict(source_file='=report.htm', machine='AOI-1',
                         wafer_statuses=[dict(status='Aborted.', slot='1', wafer_id='=ID')])]
            errors = [dict(machine='AOI-2', source_file='broken.htm', error='decode failure')]
            wph.write_wph_excel(p, rows, parse_errors=errors)
            wb = openpyxl.load_workbook(p)
            try:
                summary = wb['07_상태요약']
                self.assertEqual(summary['B2'].value, 2)
                self.assertEqual(summary['B6'].value, 1)
                self.assertEqual(len(summary._charts), 2)
                self.assertEqual(len(wb['05_대시보드']._charts), 4)
                self.assertEqual(wb.active.title, '05_대시보드')
                self.assertTrue(all(c.anchor._from.col == 0 for c in summary._charts))
                self.assertEqual(wb['08_Report목록']['B2'].data_type, 's')
                self.assertEqual(wb['09_상태상세']['D2'].data_type, 's')
                self.assertEqual(wb['10_파싱오류']['B2'].value, 'broken.htm')
                for chart in summary._charts:
                    self.assertEqual(len(chart.series), 1)
                    self.assertIn('07_상태요약', chart.series[0].val.numRef.f)
                wb.save(Path(d) / 'roundtrip.xlsx')
            finally:
                wb.close()
            with ZipFile(p) as z:
                self.assertIsNone(z.testzip())
                for name in z.namelist():
                    if name.endswith('.xml'):
                        ElementTree.fromstring(z.read(name))

if __name__ == '__main__':
    unittest.main()
