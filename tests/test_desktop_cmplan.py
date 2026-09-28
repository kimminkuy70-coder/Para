"""Commonality 조사 개선 (2026-09): 계획 엑셀 새 양식('이슈 Lot', 생성일자 열 없음)·불러오기,
S/M 찾기 진행 표시와 폴더 목록 재사용, Zone 단위 양식 편집, 확정 시 이름·체크 기억."""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from param_manager import commonality as cm, namestore
from param_manager.desktop_cmrun import DesktopCmRun
from test_commonality import _make_wafer

DEV = "2D@R2-DEVA-1_0855360PD-0A"


class PlanFileTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="rev1-cmplan-"))
        _make_wafer(self.tmp / "eq", "AOI-6", DEV, "6321", "HPG", "CX01", delta=25)
        self.cfg = self.tmp / "c.json"
        self.cfg.write_text(json.dumps({"save_dir": str(self.tmp / "save"), "local_dir": str(self.tmp / "local"),
                                        "commonality_roots": {"AOI-6": str(self.tmp / "eq")}}), encoding="utf-8")
        (self.tmp / "save").mkdir()
        self.run = DesktopCmRun(self.cfg)

    def test_new_template_has_issue_lot_and_no_created_date(self):
        out = self.run.survey.plan_template({"rows": [
            {"디바이스명": "DEVA-1", "공정번호": "6321", "S/M": "HPG", "AOI호기": "AOI-6", cm.ISSUE_KEY: "Y"}]})
        self.assertTrue(out["path"].startswith(str(self.tmp / "local")))      # local, never OneDrive
        wb = openpyxl.load_workbook(out["path"])
        heads = [c.value for c in wb["Lot목록"][1]]
        wb.close()
        self.assertEqual(heads, ["디바이스명", "공정번호", "S/M", "AOI호기", "이슈 Lot"])
        self.assertNotIn("생성일자", heads)
        self.assertNotIn("fail여부", heads)
        read = self.run.survey.read_plan({"path": out["path"]})
        self.assertEqual(read["rows"], [dict(device="DEVA-1", process="6321", sm="HPG", machine="AOI-6", issue=True)])
        self.assertEqual(read["unregistered"], [])

    def test_old_plan_file_still_reads_and_lists_unknown_machines(self):
        path = self.tmp / "old.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Lot목록"
        ws.append(["디바이스명", "공정번호", "S/M", "AOI호기", "fail여부", "생성일자"])
        ws.append(["DEVA-1", "6321", "HPG", "AOI-06", "Y", "2026-08-01"])
        ws.append(["DEVB", "7000", "TVS", "AOI-21,22", "", ""])
        wb.save(path)
        read = self.run.survey.read_plan({"path": str(path)})
        self.assertEqual([r["issue"] for r in read["rows"]], [True, False])
        self.assertEqual(read["unregistered"], ["AOI-21", "AOI-22"])        # AOI-06 == registered AOI-6
        with self.assertRaises(ValueError):
            self.run.survey.read_plan({"path": str(self.tmp / "nope.xlsx")})
        # Old row key 'fail여부' is still accepted by the plan step.
        out = self.run.plan({"machine": "AOI-6", "plan": [
            {"디바이스명": "DEVA-1", "공정번호": "6321", "S/M": "HPG", "AOI호기": "AOI-6", "fail여부": "Y"}]})
        self.assertTrue(out["lots"][0]["fail"])

    def test_result_workbook_labels_issue_lot_and_reads_old_fail(self):
        path = self.tmp / "r.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "PI3"
        wb.active.append(["PI", "Recipe", "Zone", "Alg", "Parameter", "비고", "HPG"])
        meta = wb.create_sheet("_정보")
        meta.append(["Machine", "AOI-6"])
        meta.append(["Fail", "HPG"])                                          # old result file
        wb.save(path)
        self.assertEqual(cm.read_lot_result(str(path))["fails"], {"HPG"})


class SearchProgressTests(unittest.TestCase):
    def test_progress_per_row_and_each_folder_listed_once(self):
        tmp = Path(tempfile.mkdtemp(prefix="rev1-cmscan-"))
        for sm in ("HPG", "TVS", "ASD"):
            _make_wafer(tmp / "eq", "AOI-6", DEV, "6321", sm, "CX01", delta=25)
        roots = cm.scanresult_roots(str(tmp / "eq"), "AOI-6")
        rows = [{"디바이스명": "DEVA-1", "공정번호": "6321", "S/M": sm, "AOI호기": "AOI-6"} for sm in ("HPG", "TVS", "ASD")]
        seen, listed = [], []
        real = os.scandir

        def counting(path):
            listed.append(str(path))
            return real(path)
        with patch("param_manager.commonality.os.scandir", side_effect=counting):
            lots = cm.resolve_plan(roots, rows, progress=lambda i, n, r: seen.append((i, n, r["S/M"])))
        self.assertEqual(seen, [(1, 3, "HPG"), (2, 3, "TVS"), (3, 3, "ASD")])
        self.assertEqual(sorted(l.label for l in lots if l.exists), ["ASD", "HPG", "TVS"])
        # The Scanresult root and the device/process folders are listed once for all three rows.
        self.assertEqual(len(listed), len(set(listed)))


class EditorAndConfirmTests(unittest.TestCase):
    def test_zone_pages_and_confirm_remembers_names_and_candidates(self):
        tmp = Path(tempfile.mkdtemp(prefix="rev1-cmform-"))
        _make_wafer(tmp / "eq", "AOI-6", DEV, "6321", "HPG", "CX01", delta=25)
        save = tmp / "save"
        save.mkdir()
        cfg = tmp / "c.json"
        cfg.write_text(json.dumps({"save_dir": str(save), "local_dir": str(tmp / "local"),
                                   "commonality_roots": {"AOI-6": str(tmp / "eq")}}), encoding="utf-8")
        run = DesktopCmRun(cfg)
        out = run.plan({"machine": "AOI-6", "plan": [{"디바이스명": "DEVA-1", "공정번호": "6321", "S/M": "HPG", "AOI호기": "AOI-6"}]})
        run.copy({"picks": [{"id": out["lots"][0]["id"]}]})
        det = run.detect({"unit": 0, "base": "PI3"})
        parsed = run.parse({"unit": 0, "scales": {s["variant"]: s["coef"] for s in det["scales"]}, "base_form": ""})
        snap = parsed["form"]["version"]
        first = run.form.page({"snapshot": snap, "used_only": False, "limit": 100})
        zones = [z["zone"] for z in first["zones"]]
        self.assertTrue(zones)
        page = run.form.page({"snapshot": snap, "used_only": False, "limit": 3000, "zone": zones[0]})
        self.assertTrue(page["rows"] and all(r["zone"] == zones[0] for r in page["rows"]))
        self.assertEqual(sum(z["total"] for z in first["zones"]), first["total"])
        with self.assertRaises(ValueError):                     # 3000-row pages only in zone view
            run.form.page({"snapshot": snap, "used_only": False, "limit": 3000})
        row = page["rows"][0]
        run.form.edit({"snapshot": snap, "row": row["id"], "kind": "use", "value": True})
        run.form.edit({"snapshot": snap, "row": row["id"], "kind": "name", "value": "내가 정한 이름"})
        conf = run.confirm({"unit": 0, "snapshot": snap})
        self.assertEqual(conf["notes"], [])
        self.assertTrue(os.path.isfile(conf["form"][:-5] + "_원본.xlsx"))   # full candidate list kept
        names = namestore.load(namestore.name_path(str(save)))
        lookup = namestore.make_lookup(names)
        self.assertEqual(lookup(row["alg"], row["orig"]), "내가 정한 이름")
        # The next survey form opens with the remembered name and check.
        again = run.parse({"unit": 0, "scales": {s["variant"]: s["coef"] for s in det["scales"]}, "base_form": ""})
        page2 = run.form.page({"snapshot": again["form"]["version"], "used_only": False, "limit": 3000, "zone": zones[0]})
        same = next(r for r in page2["rows"] if r["orig"] == row["orig"] and r["alg"] == row["alg"])
        self.assertEqual((same["name"], same["use"]), ("내가 정한 이름", True))


class BulkAndDetailTests(unittest.TestCase):
    def test_bulk_select_and_sm_details_and_separate_new_recipe_editor(self):
        from param_manager import desktop_ipc
        import io
        tmp = Path(tempfile.mkdtemp(prefix="rev1-cmbulk-"))
        _make_wafer(tmp / "eq", "AOI-6", DEV, "6321", "HPG", "CX01", delta=25)
        _make_wafer(tmp / "eq", "AOI-6", DEV, "6321", "HPG", "CX02", delta=30)
        cfg = tmp / "c.json"
        cfg.write_text(json.dumps({"save_dir": "", "local_dir": str(tmp / "local"),
                                   "commonality_roots": {"AOI-6": str(tmp / "eq")}}), encoding="utf-8")
        run = DesktopCmRun(cfg)
        out = run.plan({"machine": "AOI-6", "plan": [{"디바이스명": "DEVA-1", "공정번호": "6321", "S/M": "HPG",
                                                       "AOI호기": "AOI-6", cm.ISSUE_KEY: "Y"}]})
        copied = run.copy({"picks": [{"id": out["lots"][0]["id"], "wafers": ["CX01", "CX02"]}]})
        sm = copied["units"][0]["sm_list"]
        self.assertEqual([x["label"] for x in sm], ["HPG·CX01", "HPG·CX02"])
        self.assertEqual({x["slot"] for x in sm}, {"CX01", "CX02"})
        self.assertTrue(all(x["issue"] and x["sm"] == "HPG" and x["lot"] == "6321" for x in sm))
        det = run.detect({"unit": 0, "base": "PI3"})
        snap = run.parse({"unit": 0, "scales": {s["variant"]: s["coef"] for s in det["scales"]}, "base_form": ""})["form"]["version"]
        zones = [z["zone"] for z in run.form.page({"snapshot": snap, "used_only": False, "limit": 100})["zones"]]
        off = run.form.bulk({"snapshot": snap, "value": False, "variant": "", "query": "", "used_only": False})
        self.assertEqual(off["used"], 0)
        on = run.form.bulk({"snapshot": snap, "value": True, "variant": "", "query": "", "used_only": False, "zone": zones[0]})
        page = run.form.page({"snapshot": snap, "used_only": False, "limit": 3000, "zone": zones[0]})
        self.assertEqual(on["used"], len(page["rows"]))
        self.assertTrue(all(r["use"] for r in page["rows"]))
        with self.assertRaises(ValueError):
            run.form.bulk({"snapshot": "stale", "value": True})
        # 신규 Recipe 만들기 and Recipe 양식 편집하기 no longer share one editor.
        session = desktop_ipc.Session(io.BytesIO())
        self.assertIsNot(session.formnew.form, session.form)
        session.close()


if __name__ == "__main__":
    unittest.main()
