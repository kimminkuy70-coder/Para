"""Color · Gray 매칭: 탐색·매칭·크롭 크기·출력 경로 규칙·Pillow 없는 Excel·어댑터 흐름·원본 무변경."""
import base64
import hashlib
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from param_manager import colorgray as cg  # noqa: E402
from param_manager.desktop_colorgray import DesktopColorGray  # noqa: E402
from colorgray_fixture import COLOR_JPEG, GRAY_JPEG, make_wafer  # noqa: E402


class _Config:
    def __init__(self, root):
        self.root = root

    def local_state(self):
        return {"root": self.root}


def _digest(folder):
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in Path(folder).rglob("*") if p.is_file()}


class ColorGrayTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name)
        self.lot = self.base / "LOT1"
        self.w1 = make_wafer(self.lot / "Slot01", "W01")
        self.w2 = make_wafer(self.lot / "Slot02", "W02")
        (self.lot / "empty").mkdir()

    def tearDown(self):
        self.tmp.cleanup()

    def test_discover_and_plan(self):
        found = cg.discover(self.lot)
        self.assertEqual([p.name for p in found], ["W01", "W02"])
        self.assertEqual(cg.discover(self.w1), [self.w1])        # a wafer folder itself
        records, failures, total = cg.plan_wafer(self.w1)
        self.assertEqual(total, 3)
        self.assertEqual(len(records), 1)
        r = records[0]
        self.assertEqual((r["gray_file"], r["candidate_count"]), ("F1.t.1.jpg", 2))   # F1: fault farther from edge
        self.assertAlmostEqual(r["pixel_x"], 1100)
        self.assertEqual((r["crop_w"], r["crop_h"], r["out_w"], r["out_h"]),
                         (round(46 * 0.4674 / cg.GP), round(34 * 0.4674 / cg.GP), 46, 34))
        reasons = dict(failures)
        self.assertIn("C2.jpg", reasons)
        self.assertIn("C3.jpg", reasons)

    def test_jpeg_size(self):
        self.assertEqual(cg.jpeg_size(COLOR_JPEG), (46, 34))
        self.assertEqual(cg.jpeg_size(GRAY_JPEG), (96, 31))
        self.assertIsNone(cg.jpeg_size(b"not a jpeg"))

    def test_adapter_flow_workbook_without_pillow_and_sources_unchanged(self):
        before = _digest(self.lot)
        local = self.base / "local"
        a = DesktopColorGray(_Config(str(local)))
        scan = a.scan({"root": str(self.lot)})
        self.assertEqual(len(scan["wafers"]), 2)
        with self.assertRaises(ValueError):
            a.start({"wafers": [str(self.base / "elsewhere")]})     # only scanned wafers
        with self.assertRaises(ValueError):                         # output inside the source
            a.start({"wafers": [scan["wafers"][0]["path"]], "output": str(self.lot / "out")})
        job = a.start({"wafers": [scan["wafers"][0]["path"]], "output": ""})["job"]
        self.assertTrue(job["output"].startswith(str(local)))
        rec = job["wafers"][0]["records"][0]
        self.assertEqual((rec["color"], rec["gray"], rec["gw"]), ("C1.jpg", "F1.t.1.jpg", cg.GW))
        got = a.read({"wafer": 0, "record": 0, "kind": "gray"})
        self.assertEqual((base64.b64decode(got["data"]), got["next"]), (GRAY_JPEG, None))
        import param_manager.desktop_colorgray as dcg
        old_chunk, dcg.CHUNK = dcg.CHUNK, 500                  # pieces are reassembled in order
        try:
            first = a.read({"wafer": 0, "record": 0, "kind": "gray"})
            rest = a.read({"wafer": 0, "record": 0, "kind": "gray", "offset": first["next"]})
        finally:
            dcg.CHUNK = old_chunk
        self.assertEqual(base64.b64decode(first["data"]) + base64.b64decode(rest["data"]), GRAY_JPEG)
        self.assertIsNone(rest["next"])
        self.assertEqual(dcg.MAX_SOURCE, 5 * 1024 * 1024)
        with self.assertRaises(ValueError):
            a.read({"wafer": 0, "record": 5, "kind": "color"})
        with self.assertRaises(ValueError):
            a.put({"wafer": 0, "record": 0, "kind": "crop", "data": base64.b64encode(b"xx").decode()})
        for kind in ("crop", "color_thumb", "gray_thumb", "crop_thumb"):
            extra = {"box": [1, 2, 3, 4]} if kind == "crop" else {}
            a.put({"wafer": 0, "record": 0, "kind": kind, "data": base64.b64encode(COLOR_JPEG).decode(), **extra})
        done = a.finish({"wafer": 0})
        self.assertEqual((done["matched"], done["failed"]), (1, 2))
        book = Path(done["workbook"])
        self.assertTrue(book.exists())
        with zipfile.ZipFile(book) as z:
            media = [n for n in z.namelist() if n.startswith("xl/media/")]
        self.assertEqual(len(media), 3)
        import openpyxl
        ws = openpyxl.load_workbook(book).active
        self.assertEqual(ws["C2"].hyperlink.target, ".\\_crop_images\\C1_gray_crop.jpeg")
        self.assertIn("Crop: (1, 2, 3, 4)", ws["D2"].value)
        self.assertTrue(Path(done["failures_csv"]).exists())
        self.assertEqual(_digest(self.lot), before)                 # sources only read

    def test_ipc_routes(self):
        import io
        import json
        from param_manager import desktop_ipc
        self.assertTrue({"cgm_scan", "cgm_start", "cgm_read", "cgm_put", "cgm_finish"} <= desktop_ipc.METHODS)
        out = io.BytesIO()
        s = desktop_ipc.Session(out)
        s.colorgray = DesktopColorGray(_Config(str(self.base / "local")))
        s.handle(dict(version=1, id=1, method="cgm_scan", params={"root": str(self.lot)}))
        for w in list(s.workers):
            w.join(10)
        events = [json.loads(line) for line in out.getvalue().decode().splitlines()]
        done = [e for e in events if e["event"] == "completed"]
        self.assertEqual(len(done[0]["colorgray"]["wafers"]), 2)
        s.close()


if __name__ == "__main__":
    unittest.main()
