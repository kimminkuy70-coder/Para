"""formbuilder 테스트 — 편집용 initial 생성 + 편집된 initial→final 양식."""
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402

from param_manager import engine, extract_io, formbuilder, ini_parser  # noqa: E402

GLOBAL_RTP = """[GLOBAL_RTP]
MaxFaultsPerWafer = 3000
ApplyDieCalib = 1
"""
ZONE_INI = """[General]
ZoneName = PI Opening
[Surface]
High_Delta = 25
BrightLength = 5
"""


def _mk_recipe(d: Path):
    d.mkdir(parents=True, exist_ok=True)
    (d / "GlobalRTP.ini").write_text(GLOBAL_RTP, encoding="utf-8")
    (d / "Zones").mkdir(exist_ok=True)
    (d / "Zones" / "Zone1.ini").write_text(ZONE_INI, encoding="utf-8")


def _pivot(tmp):
    recipe = Path(tmp) / "R_TB500_PI3" / "PI"
    _mk_recipe(recipe)
    cfgs = ini_parser.scan_tree(Path(tmp) / "R_TB500_PI3", default_equipment="AOI-13")
    rows, machines = ini_parser.build_pivot(cfgs)
    return rows, machines


def test_build_initial_and_final():
    with tempfile.TemporaryDirectory() as tmp:
        rows, _ = _pivot(tmp)
        assert rows
        init = os.path.join(tmp, "01_초안.xlsx")
        formbuilder.build_initial_workbook(rows, init, level="PI3", source="AOI-13")
        assert os.path.isfile(init)

        wb = openpyxl.load_workbook(init)
        ws = wb[formbuilder.INIT_SHEET]
        heads = [c.value for c in ws[1]]
        assert heads == formbuilder.INIT_HEADERS
        assert ws.max_row - 1 == len(rows)            # 헤더 제외 = 파싱 행 수
        # 최종 Parameter 기본값 = 추천 Parameter
        gi = heads.index("추천 Parameter"); fi = heads.index("최종 Parameter")
        for r in ws.iter_rows(min_row=2, values_only=True):
            assert r[fi] == r[gi]
        wb.close()

        # 사람이 편집한 상황을 흉내: 한 행 사용=N, 한 행 최종 이름 변경
        wb = openpyxl.load_workbook(init)
        ws = wb[formbuilder.INIT_SHEET]
        ui = formbuilder.INIT_HEADERS.index("사용")
        ws.cell(2, ui + 1).value = "N"                # 첫 데이터 행 제외
        ws.cell(3, fi + 1).value = "내가 지정한 이름"     # 둘째 행 이름 변경
        wb.save(init)
        wb.close()

        final = os.path.join(tmp, "02_확정.xlsx")
        res = formbuilder.build_final_from_initial(
            final and init, final, level="PI3", aoi="AOI-13", user="tester")
        assert res["kept"] == len(rows) - 1 and res["dropped"] == 1
        assert res["sheet"] == "PI_ALL"

        # final 은 공용 양식(PI_ALL) + _EXTRACT_MAP 을 가진다
        repo = engine.ParamRepository(final)
        repo.load()
        assert len(repo.rows) == len(rows) - 1
        names = {engine._s(pr.get("Parameter")) for pr in repo.rows}
        assert "내가 지정한 이름" in names
        emap = extract_io.read_extract_map(final)
        assert emap                                   # 재추출 메타 보존
        print("  formbuilder OK: initial 생성 + 사용/이름편집 반영 final + _EXTRACT_MAP")


def test_final_stores_scales():
    with tempfile.TemporaryDirectory() as tmp:
        rows, _ = _pivot(tmp)
        init = os.path.join(tmp, "init.xlsx")
        formbuilder.build_initial_workbook(rows, init, level="PI3")
        final = os.path.join(tmp, "form.xlsx")
        scales = {"PI": 0.8456665875666588, "PI-bubble": 0.7696441409644141}
        formbuilder.build_final_from_initial(init, final, level="PI3", scales=scales)
        got = extract_io.read_scales(final)
        assert got == scales, got
        # '초기 추천값'(대표값)은 최종 스키마에 없어야 함
        assert "초기 추천값" not in engine.META_FIELDS
    print("  formbuilder OK: 변형별 계수 양식에 저장/판독 + 추천값 스키마 제거")


def test_rename_level_draft_and_final():
    """레시피 이름 수정: 초안(양식초안 PI 열·사용법)과 확정본(PI_ALL PI 열·추출_요약)
    안의 레벨 이름이 함께 바뀌고, 확정 재생성 시 새 이름이 유지된다."""
    with tempfile.TemporaryDirectory() as tmp:
        rows, _ = _pivot(tmp)
        init = os.path.join(tmp, "init.xlsx")
        formbuilder.build_initial_workbook(rows, init, level="PI3", source="AOI-13")
        final = os.path.join(tmp, "form.xlsx")
        formbuilder.build_final_from_initial(init, final, level="PI3", aoi="AOI-13")

        # 초안 이름 변경 → PI 열 전부 + 사용법 줄
        n = formbuilder.rename_level(init, "PI3", "PI3-NEW")
        assert n >= len(rows) + 1, n
        wb = openpyxl.load_workbook(init)
        ws = wb[formbuilder.INIT_SHEET]
        pi = formbuilder.INIT_HEADERS.index("PI")
        vals = {r[pi] for r in ws.iter_rows(min_row=2, values_only=True)}
        assert vals == {"PI3-NEW"}, vals
        info = [c.value for c in wb["사용법"]["A"]]
        assert "레시피 레벨: PI3-NEW" in info
        wb.close()

        # 이름 바뀐 초안 → 확정 재생성: 새 이름이 그대로 반영
        final2 = os.path.join(tmp, "form2.xlsx")
        formbuilder.build_final_from_initial(init, final2, level="PI3-NEW")
        repo = engine.ParamRepository(final2)
        repo.load()
        assert {engine._s(pr.get("PI")) for pr in repo.rows} == {"PI3-NEW"}

        # 확정본 직접 이름 변경(초안 없는 옛 버전 대응) → PI_ALL + 추출_요약
        n2 = formbuilder.rename_level(final, "PI3", "PI3-NEW")
        assert n2 > 0
        repo2 = engine.ParamRepository(final)
        repo2.load()
        assert {engine._s(pr.get("PI")) for pr in repo2.rows} == {"PI3-NEW"}
        wb = openpyxl.load_workbook(final)
        sm = wb[extract_io.SHEET_SUMMARY]
        d = {engine._s(r[0]): engine._s(r[1])
             for r in sm.iter_rows(min_row=2, values_only=True) if r and r[0]}
        assert d.get("레시피 레벨") == "PI3-NEW", d
        wb.close()
        # 같은 이름/빈 이름은 no-op
        assert formbuilder.rename_level(final, "PI3-NEW", "PI3-NEW") == 0
        assert formbuilder.rename_level(final, "PI3-NEW", "") == 0
    print("  formbuilder OK: rename_level — 초안/확정/요약 레벨 이름 일괄 변경")


def test_force_level_and_form_params():
    """force_level: PI 열 전체를 새 이름으로 강제(옛 값이 폴더명과 달라도).
    form_params: (Zone,Alg,Parameter) 키 집합으로 두 양식 비교."""
    with tempfile.TemporaryDirectory() as tmp:
        rows, _ = _pivot(tmp)
        init = os.path.join(tmp, "init.xlsx")
        formbuilder.build_initial_workbook(rows, init, level="PI3/x5", source="AOI-13")
        final = os.path.join(tmp, "form.xlsx")
        formbuilder.build_final_from_initial(init, final, level="PI3/x5", aoi="AOI-13")

        # 폴더명(list_recipes)은 'PI3_x5'로 sanitize 됐다고 가정 → rename 은 매칭 실패지만
        # force_level 은 현재 값과 무관하게 전부 새 이름으로 통일
        n = formbuilder.force_level(final, "PI3-NEW")
        assert n > 0
        repo = engine.ParamRepository(final)
        repo.load()
        assert {engine._s(pr.get("PI")) for pr in repo.rows} == {"PI3-NEW"}
        wb = openpyxl.load_workbook(final)
        sm = wb[extract_io.SHEET_SUMMARY]
        d = {engine._s(r[0]): engine._s(r[1])
             for r in sm.iter_rows(min_row=2, values_only=True) if r and r[0]}
        assert d.get("레시피 레벨") == "PI3-NEW"
        wb.close()

        # form_params 비교: 같은 양식이면 차집합 0, 파라미터 추가 시 검출
        keys1 = formbuilder.form_params(final)
        assert keys1 and formbuilder.form_params(final) - keys1 == set()

        recs = [{"PI": "PI3", "Recipe": "PI", "Zone": "Z", "Alg": "S",
                 "Parameter": "OldOnly", "비고": ""},
                {"PI": "PI3", "Recipe": "PI", "Zone": "Z", "Alg": "S",
                 "Parameter": "NewOne", "비고": ""}]
        f2 = os.path.join(tmp, "f2.xlsx")
        extract_io.write_snapshot(f2, recs, machines=[], sheet_name="PI_ALL",
                                  extracts=[None, None], stage="final", level="PI3")
        f1 = os.path.join(tmp, "f1.xlsx")
        extract_io.write_snapshot(f1, recs[:1], machines=[], sheet_name="PI_ALL",
                                  extracts=[None], stage="final", level="PI3")
        added = formbuilder.form_params(f2) - formbuilder.form_params(f1)
        names = {k[2] for k in added}
        assert any("newone" == n for n in names), names
    print("  formbuilder OK: force_level 강제 통일 + form_params 신규항목 검출")


def test_similarity_ranking_and_keys():
    """새 레시피 파라미터 키 집합 + 기존 양식들과 유사도 순위(겹침 많은 순)."""
    rows = [{"zone": "Scan Area", "alg": "Surface", "param": "Contrast Delta - Bright"},
            {"zone": "Scan Area", "alg": "Surface", "param": "Min Defect Width"},
            {"zone": "LIGHT", "alg": "Scan2d", "param": "LightSrcRef_NominalGL"}]
    pk = formbuilder.pivot_param_keys(rows)
    assert len(pk) == 3
    fa = {formbuilder._norm_key3("Scan Area", "Surface", "Contrast Delta - Bright"),
          formbuilder._norm_key3("Scan Area", "Surface", "Min Defect Width"),
          formbuilder._norm_key3("Q", "W", "E")}
    fb2 = {formbuilder._norm_key3("LIGHT", "Scan2d", "LightSrcRef_NominalGL")}
    ranked = formbuilder.rank_similar_forms(pk, {"A": fa, "B": fb2})
    assert ranked[0] == ("A", 2, 3) and ranked[1] == ("B", 1, 1)
    # 기반 선택 파생: 사용 = 키가 기반 양식에 있음
    sel = [formbuilder._norm_key3(r["zone"], r["alg"], r["param"]) in fa for r in rows]
    assert sel == [True, True, False]
    print("  formbuilder OK: 유사도 순위 + 기반 선택 파생")


def test_form_to_pivot_roundtrip():
    """확정 양식 → 편집기용 rows/scales 복원(기존 양식 수정하기 = 프로그램 편집기)."""
    with tempfile.TemporaryDirectory() as tmp:
        final = os.path.join(tmp, "form.xlsx")
        coef = 0.7707763913156815
        recs = [{"PI": "PI3", "Recipe": "PI", "Zone": "Scan Area", "Alg": "Surface",
                 "Parameter": "Min Defect Area (area, µ)", "비고": ""},
                {"PI": "PI3", "Recipe": "PI", "Zone": "Scan Area", "Alg": "Surface",
                 "Parameter": "Contrast Delta - Bright", "비고": ""}]
        exts = [{"src_file": "z.ini", "section": "Surface", "key": "A", "raw": 2,
                 "transform": ini_parser.label_transform("AREA", coef), "source_path": ""},
                {"src_file": "z.ini", "section": "Surface", "key": "B", "raw": 255,
                 "transform": "RAW", "source_path": ""}]
        extract_io.write_snapshot(final, recs, machines=[], sheet_name="PI_ALL",
                                  extracts=exts, stage="final", level="PI3", aoi="AOI-1",
                                  scales={"PI": coef})
        rows, scales = formbuilder.form_to_pivot(final)
        assert len(rows) == 2 and scales.get("PI") == coef
        r0 = next(r for r in rows if "Min Defect Area" in r["param"])
        assert r0["mag"] == "PI" and r0["zone"] == "Scan Area" and r0["use"] is True
        assert r0["extract"]["transform"].startswith("AREA")
        assert engine._s(r0["raws"]["양식"]) == "2"
    print("  formbuilder OK: form_to_pivot 확정양식→편집기 rows/scales 복원")


def test_final_rejects_when_all_unused():
    with tempfile.TemporaryDirectory() as tmp:
        rows, _ = _pivot(tmp)
        init = os.path.join(tmp, "init.xlsx")
        formbuilder.build_initial_workbook(rows, init)
        wb = openpyxl.load_workbook(init)
        ws = wb[formbuilder.INIT_SHEET]
        ui = formbuilder.INIT_HEADERS.index("사용") + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(r, ui).value = "N"
        wb.save(init); wb.close()
        try:
            formbuilder.build_final_from_initial(init, os.path.join(tmp, "f.xlsx"))
            assert False, "빈 선택인데 예외가 안 났다"
        except ValueError:
            pass
    print("  formbuilder OK: 사용 행 없으면 ValueError")


def test_initial_to_pivot_restores_all_candidates():
    """초안('원본')에는 **체크 안 한 항목까지 전부** 있어야 하고, 다시 피벗으로
    복원돼야 한다 — '기존 양식 수정하기'에서 뺀 항목을 되살리는 근거다."""
    with tempfile.TemporaryDirectory() as tmp:
        rows, _ = _pivot(tmp)
        assert len(rows) >= 3, "테스트 데이터가 너무 적다"
        init = os.path.join(tmp, "원본.xlsx")
        formbuilder.build_initial_workbook(rows, init, level="PI3", source="AOI-13")
        # 절반은 사용=N 으로 (사람이 뺀 상황)
        wb = openpyxl.load_workbook(init)
        ws = wb[formbuilder.INIT_SHEET]
        ui = formbuilder.INIT_HEADERS.index("사용") + 1
        n_off = 0
        for r in range(2, ws.max_row + 1):
            if r % 2 == 0:
                ws.cell(r, ui).value = "N"
                n_off += 1
        wb.save(init); wb.close()

        back, used, names = formbuilder.initial_to_pivot(init)
        assert len(back) == len(rows), "후보가 줄었다(사용=N 이 사라지면 안 됨)"
        assert sum(1 for r in back if not r["use"]) == n_off
        assert len(used) == len(rows) - n_off
        assert names and all(len(k) == 3 for k in names)
        # 설정키(재추출 메타)가 살아 있어야 값 매칭이 된다
        assert any(r["extract"].get("key") for r in back)
    print("  formbuilder OK: 초안 → 전체 후보 복원(사용=N 포함)")


def test_edit_existing_form_can_re_add_dropped_params():
    """확정 양식 + 초안 원본 → 편집기 rows: 지금 쓰는 항목은 켜지고,
    예전에 뺀 항목도 **목록에 남아 다시 넣을 수 있어야** 한다."""
    with tempfile.TemporaryDirectory() as tmp:
        rows, _ = _pivot(tmp)
        init = os.path.join(tmp, "원본.xlsx")
        formbuilder.build_initial_workbook(rows, init, level="PI3", source="AOI-13")
        # 첫 행은 빼고, 둘째 행은 이름을 바꿔서 확정(사람이 실제로 하는 편집)
        wb = openpyxl.load_workbook(init)
        ws = wb[formbuilder.INIT_SHEET]
        ui = formbuilder.INIT_HEADERS.index("사용") + 1
        fi = formbuilder.INIT_HEADERS.index("최종 Parameter") + 1
        ws.cell(2, ui).value = "N"
        ws.cell(3, fi).value = "사람이 바꾼 이름"
        wb.save(init); wb.close()
        final = os.path.join(tmp, "확정.xlsx")
        res = formbuilder.build_final_from_initial(init, final, level="PI3",
                                                   aoi="AOI-13")
        assert res["dropped"] == 1

        # 확정본만으로 열면(종전 동작) 뺀 항목이 아예 없다 → 다시 넣을 수 없다
        only_final, _sc = formbuilder.form_to_pivot(final)
        assert len(only_final) == len(rows) - 1

        # 초안 원본으로 열면 전체가 보이고, 쓰는 항목만 체크돼 있다
        cand, _u, _n = formbuilder.initial_to_pivot(init)
        merged, matched = formbuilder.merge_form_into_candidates(cand, final)
        assert len(merged) == len(rows), "후보가 줄면 안 됨"
        assert matched == len(rows) - 1
        assert sum(1 for r in merged if r["use"]) == len(rows) - 1
        assert sum(1 for r in merged if not r["use"]) == 1, "뺀 항목 1개가 후보로 남아야"
        # 이름을 바꾼 행은 **설정키로 매칭**돼 새 항목으로 튀지 않고, 바뀐 이름 유지
        params = [r["param"] for r in merged if r["use"]]
        assert "사람이 바꾼 이름" in params, params
    print("  formbuilder OK: 기존 양식 수정 — 뺀 항목 복원 + 바뀐 이름 설정키 매칭")


def test_merge_keeps_rows_only_in_form():
    """초안에 없고 확정본에만 있는 행(엑셀에서 손으로 추가한 행 등)도 살아남아야."""
    with tempfile.TemporaryDirectory() as tmp:
        rows, _ = _pivot(tmp)
        init = os.path.join(tmp, "원본.xlsx")
        formbuilder.build_initial_workbook(rows, init, level="PI3")
        final = os.path.join(tmp, "확정.xlsx")
        formbuilder.build_final_from_initial(init, final, level="PI3")
        # 확정본에 초안에 없는 행을 하나 추가
        wb = openpyxl.load_workbook(final)
        ws = wb["PI_ALL"]
        heads = [engine._s(c.value) for c in ws[1]]
        newr = [""] * len(heads)
        newr[heads.index("PI")] = "PI3"
        newr[heads.index("Zone")] = "손으로 추가"
        newr[heads.index("Alg")] = "Manual"
        newr[heads.index("Parameter")] = "직접 넣은 항목"
        ws.append(newr)
        wb.save(final); wb.close()

        cand, _u, _n = formbuilder.initial_to_pivot(init)
        merged, _m = formbuilder.merge_form_into_candidates(cand, final)
        assert any(r["param"] == "직접 넣은 항목" and r["use"] for r in merged), \
            "확정본에만 있던 행이 사라졌다"
    print("  formbuilder OK: 확정본에만 있는 행도 보존")


if __name__ == "__main__":
    fails = 0
    tests = [(n, f) for n, f in list(globals().items())
             if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        print(f"[RUN] {name}")
        try:
            fn()
            print(f"[PASS] {name}\n")
        except Exception as e:  # noqa: BLE001
            fails += 1
            import traceback
            traceback.print_exc()
            print(f"[FAIL] {name}: {e}\n")
    print(f"==== {len(tests) - fails}/{len(tests)} passed ====")
    sys.exit(1 if fails else 0)
