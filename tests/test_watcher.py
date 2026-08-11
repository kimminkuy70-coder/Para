"""watcher.py 검증 — 주기/시간대/backoff/안정성 필터/변경 보고서(GUI·장비 비의존)."""

import os
import sys
import tempfile
import time
from datetime import datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402

from param_manager import collate, engine, watcher, workdirs  # noqa: E402

PASS = FAIL = 0


def run(fn):
    global PASS, FAIL
    print(f"[RUN] {fn.__name__}")
    try:
        fn()
        PASS += 1
        print(f"[PASS] {fn.__name__}\n")
    except Exception as e:  # noqa: BLE001
        FAIL += 1
        import traceback
        print(f"[FAIL] {fn.__name__}: {e}")
        traceback.print_exc()
        print()


def _collate_file(path, rows, machines=("AOI-6",)):
    """취합 파일 흉내 — 레시피 시트 1개 + 호기 열."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PI3"
    heads = list(engine.META_FIELDS) + list(machines)
    ws.append(heads)
    for r in rows:
        ws.append([r.get(h) for h in heads])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


def _row(param, value, machine="AOI-6"):
    return {"PI": "PI3", "Recipe": "PI", "Zone": "PAD", "Alg": "Contrast",
            "Parameter": param, "비고": "", machine: value}


# --------------------------------------------------------------------------
def test_settings_roundtrip_and_defaults():
    with tempfile.TemporaryDirectory() as tmp:
        s, st = watcher.load_settings(tmp)            # 파일 없음 → 기본값
        assert s.enabled is False
        assert s.conn_mode == watcher.CONN_SESSION, "기본 접속은 net use 없이"
        assert s.interval_hours == watcher.DEFAULT_INTERVAL_HOURS
        s.enabled = True
        s.interval_hours = 3
        s.conn_mode = watcher.CONN_NETUSE
        s.window_start, s.window_end = 22, 6
        watcher.save_settings(tmp, s, watcher.WatchState(last_run="2026-01-01 00:00:00"))
        s2, st2 = watcher.load_settings(tmp)
        assert s2.enabled and s2.interval_hours == 3
        assert s2.conn_mode == watcher.CONN_NETUSE
        assert (s2.window_start, s2.window_end) == (22, 6)
        assert st2.last_run == "2026-01-01 00:00:00"
    print("  설정 저장/복원 + 기본값(기존 세션 모드) OK")


def test_settings_survives_corrupt_file():
    with tempfile.TemporaryDirectory() as tmp:
        with open(watcher.settings_path(tmp), "w", encoding="utf-8") as fh:
            fh.write("{깨진 JSON")
        s, st = watcher.load_settings(tmp)
        assert s.conn_mode == watcher.CONN_SESSION and st.fail_count == 0
    print("  설정 파일이 깨져도 기본값으로 동작 OK")


def test_time_window():
    s = watcher.WatchSettings(window_start=0, window_end=0)
    assert watcher.in_window(datetime(2026, 8, 4, 3), s), "start==end 면 제한 없음"
    s2 = watcher.WatchSettings(window_start=9, window_end=18)
    assert watcher.in_window(datetime(2026, 8, 4, 10), s2)
    assert not watcher.in_window(datetime(2026, 8, 4, 20), s2)
    s3 = watcher.WatchSettings(window_start=22, window_end=6)   # 자정 넘김
    assert watcher.in_window(datetime(2026, 8, 4, 23), s3)
    assert watcher.in_window(datetime(2026, 8, 4, 2), s3)
    assert not watcher.in_window(datetime(2026, 8, 4, 12), s3)
    print("  실행 시간대 창(자정 넘김 포함) OK")


def test_should_run_interval():
    s = watcher.WatchSettings(enabled=True, interval_hours=6)
    now = datetime(2026, 8, 4, 12, 0, 0)
    assert watcher.should_run(now, s, watcher.WatchState()), "첫 실행은 즉시"
    recent = watcher.WatchState(
        last_run=(now - timedelta(hours=1)).strftime("%Y-%m-%d %H:%M:%S"))
    assert not watcher.should_run(now, s, recent), "주기 전에는 실행 안 함"
    old = watcher.WatchState(
        last_run=(now - timedelta(hours=7)).strftime("%Y-%m-%d %H:%M:%S"))
    assert watcher.should_run(now, s, old)
    s.enabled = False
    assert not watcher.should_run(now, s, old), "꺼져 있으면 실행 안 함"
    print("  주기 판단(첫 실행/대기/도래/비활성) OK")


def test_backoff_delays_after_failures():
    """연속 실패 시 다음 시도를 미룬다 — 바쁜 장비에 몰아치지 않기 위해."""
    s = watcher.WatchSettings(enabled=True, interval_hours=6)
    now = datetime(2026, 8, 4, 12, 0, 0)
    last = (now - timedelta(hours=7)).strftime("%Y-%m-%d %H:%M:%S")
    ok_state = watcher.WatchState(last_run=last, fail_count=0)
    assert watcher.should_run(now, s, ok_state)
    failed = watcher.WatchState(last_run=last, fail_count=2)   # 6h*4=24h 뒤
    assert not watcher.should_run(now, s, failed), "실패가 쌓이면 더 미룬다"
    n1 = watcher.next_run_at(s, ok_state)
    n2 = watcher.next_run_at(s, failed)
    assert n2 > n1
    print("  연속 실패 backoff(재시도 몰아치기 방지) OK")


def test_record_run_updates_state():
    with tempfile.TemporaryDirectory() as tmp:
        s = watcher.WatchSettings(enabled=True)
        st = watcher.WatchState()
        st = watcher.record_run(tmp, s, st, ok=False, note="접속 실패")
        assert st.fail_count == 1 and st.last_run
        st = watcher.record_run(tmp, s, st, ok=False)
        assert st.fail_count == 2
        st = watcher.record_run(tmp, s, st, ok=True, note="정상")
        assert st.fail_count == 0, "성공하면 backoff 초기화"
        _, reloaded = watcher.load_settings(tmp)
        assert reloaded.fail_count == 0 and reloaded.last_result == "정상"
    print("  회차 상태 기록/복원(실패 누적·성공 초기화) OK")


def test_unstable_file_detection():
    """장비가 쓰는 중이던 파일은 그 회차에서 제외 — 거짓 변경 방지."""
    with tempfile.TemporaryDirectory() as tmp:
        stable = os.path.join(tmp, "GlobalRTP.ini")
        moving = os.path.join(tmp, "Zones.ini")
        for p in (stable, moving):
            with open(p, "w", encoding="utf-8") as fh:
                fh.write("[A]\nx=1\n")
        sigs = [(stable, watcher.file_sig(stable)), (moving, watcher.file_sig(moving))]
        time.sleep(0.01)
        with open(moving, "a", encoding="utf-8") as fh:      # 복사 도중 장비가 씀
            fh.write("y=2\n")
        bad = watcher.unstable_files(sigs)
        assert moving in bad and stable not in bad, bad
    print("  찢어진 읽기 감지(복사 중 변한 파일) OK")


def test_drop_unstable_filters_sources():
    with tempfile.TemporaryDirectory() as tmp:
        d1 = os.path.join(tmp, "AOI-6"); os.makedirs(d1)
        d2 = os.path.join(tmp, "AOI-9"); os.makedirs(d2)
        bad_file = os.path.join(d1, "Zones.ini")
        with open(bad_file, "w", encoding="utf-8") as fh:
            fh.write("x")
        sources = [(d1, "PI3", "AOI-6"), (d2, "PI3", "AOI-9")]
        kept = watcher.drop_unstable(sources, [bad_file])
        assert len(kept) == 1 and kept[0][2] == "AOI-9", kept
        assert watcher.drop_unstable(sources, []) == sources
    print("  불안정 파일이 속한 소스 제외 OK")


def test_compare_and_report_detects_change():
    with tempfile.TemporaryDirectory() as tmp:
        prev = _collate_file(os.path.join(tmp, "파라미터 값 취합", "old.xlsx"),
                             [_row("Contrast Delta - Bright", "25"),
                              _row("Min Defect Width", "10")])
        new = _collate_file(os.path.join(tmp, "파라미터 값 취합", "new.xlsx"),
                            [_row("Contrast Delta - Bright", "99"),
                             _row("Min Defect Width", "10")])
        res = watcher.compare_and_report(tmp, prev, new, st="20260804_120000")
        assert res.ran and res.has_change, res
        assert res.changed == 1, res.changed
        assert res.report and os.path.exists(res.report)
        assert "값변경 1건" in res.summary(), res.summary()
        # 보고서에 어느 파라미터가 바뀌었는지 들어 있어야 한다(변경내역 시트)
        wb = openpyxl.load_workbook(res.report)
        text = "\n".join(
            " ".join(str(c) for c in row if c is not None)
            for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
        wb.close()
        assert "Contrast Delta - Bright" in text, text[:400]
        assert "99" in text
    print("  변경 감지 + 보고서 생성(무엇이 바뀌었는지 포함) OK")


def test_no_change_makes_no_report():
    """변경이 없으면 보고서도 알림도 만들지 않는다(알림 피로 방지)."""
    with tempfile.TemporaryDirectory() as tmp:
        rows = [_row("Contrast Delta - Bright", "25")]
        prev = _collate_file(os.path.join(tmp, "c", "old.xlsx"), rows)
        new = _collate_file(os.path.join(tmp, "c", "new.xlsx"), rows)
        res = watcher.compare_and_report(tmp, prev, new)
        assert res.ran and not res.has_change, res
        assert res.report == ""
        assert res.summary() == "변경 없음"
    print("  무변경 → 보고서 없음 OK")


def test_first_run_has_no_baseline():
    with tempfile.TemporaryDirectory() as tmp:
        new = _collate_file(os.path.join(tmp, "c", "new.xlsx"),
                            [_row("Contrast Delta - Bright", "25")])
        res = watcher.compare_and_report(tmp, None, new)
        assert res.ran and not res.has_change and res.report == ""
    print("  첫 실행(기준선 없음) 처리 OK")


def test_connection_check_and_guide():
    """기본 모드는 net use 없이 — 연결 안 된 장비를 미리 알려준다."""
    chk = watcher.check_connections([("AOI-6", "10.255.255.1")], timeout=0.3)
    assert len(chk["missing"]) == 1, chk        # 개발환경엔 장비가 없다
    assert chk["stopped"] is False
    guide = watcher.connection_guide(chk["missing"])
    assert "AOI-6" in guide and "c$" in guide
    assert "비밀번호를 저장하지 않" in guide
    assert watcher.connection_guide([]) == ""
    print("  무인 모드 연결 사전 점검 + 안내문 OK")


def test_connection_check_is_bounded_by_timeout():
    """미응답 장비에서 무한정 붙들리면 화면이 멈춘다(응답없음 버그).
    호스트당 타임아웃이 지켜지는지 확인."""
    targets = [("AOI-6", "10.255.255.1"), ("AOI-9", "10.255.255.2"),
               ("AOI-12", "10.255.255.3")]
    t0 = time.time()
    chk = watcher.check_connections(targets, timeout=0.3)
    elapsed = time.time() - t0
    assert len(chk["missing"]) == 3, chk
    # 3대 × 0.3초 + 여유. 타임아웃이 안 먹으면 수십 초가 걸린다.
    assert elapsed < 5, f"연결 점검이 너무 오래 걸림: {elapsed:.1f}초"
    assert watcher.probe_host("", timeout=0.3) is False      # 빈 IP 방어
    print(f"  연결 점검 타임아웃 준수 OK ({elapsed:.1f}초)")


def test_connection_check_progress_and_cancel():
    """진행 보고 + 사용자 취소(중단) 지원 — 긴 점검을 멈출 수 있어야 한다."""
    targets = [("AOI-6", "10.255.255.1"), ("AOI-9", "10.255.255.2"),
               ("AOI-12", "10.255.255.3")]
    seen = []
    chk = watcher.check_connections(
        targets, timeout=0.2,
        progress=lambda done, total, m: seen.append((done, total, m)),
        should_stop=lambda: len(seen) >= 2)      # 2대 확인 후 취소
    assert chk["stopped"] is True, chk
    assert len(chk["ok"]) + len(chk["missing"]) < 3, chk
    assert seen and seen[0][1] == 3, seen
    print("  연결 점검 진행보고 + 취소 OK")


def test_selected_targets_persist():
    """감시 대상(장비·레시피) 선택이 저장·복원되어야 한다."""
    with tempfile.TemporaryDirectory() as tmp:
        s, st = watcher.load_settings(tmp)
        assert s.machines == [] and s.recipes == [], "기본은 미선택(호출부가 폴백 결정)"
        s.enabled = True
        s.machines = ["AOI-6", "AOI-9"]
        s.recipes = ["PI3", "RDL2"]
        watcher.save_settings(tmp, s, st)
        s2, _ = watcher.load_settings(tmp)
        assert s2.machines == ["AOI-6", "AOI-9"], s2.machines
        assert s2.recipes == ["PI3", "RDL2"], s2.recipes
        # 선택한 장비만 점검 대상이 된다
        targets = [(m, f"10.255.255.{i+1}") for i, m in enumerate(s2.machines)]
        chk = watcher.check_connections(targets, timeout=0.2)
        assert len(chk["ok"]) + len(chk["missing"]) == 2, chk
    print("  감시 대상 장비·레시피 선택 저장/복원 + 선택분만 점검 OK")


def test_stale_selection_is_filtered():
    """선택 저장 후 장비·레시피가 삭제됐으면 실행 시 걸러져야 한다."""
    with tempfile.TemporaryDirectory() as tmp:
        s, st = watcher.load_settings(tmp)
        s.machines = ["AOI-6", "사라진호기"]
        s.recipes = ["PI3", "사라진레시피"]
        watcher.save_settings(tmp, s, st)
        s2, _ = watcher.load_settings(tmp)
        # _watch_run_cycle 과 같은 필터 규칙
        avail_m, avail_r = {"AOI-6", "AOI-9"}, {"PI3", "PI4"}
        machines = [m for m in s2.machines if m in avail_m]
        recipes = [r for r in s2.recipes if r in avail_r]
        assert machines == ["AOI-6"], machines
        assert recipes == ["PI3"], recipes
    print("  삭제된 장비·레시피 선택 자동 제외 OK")


def test_unselected_machines_keep_values():
    """감시 대상을 1대만 골라도 **나머지 호기의 기존 값이 남아 있어야** 한다.

    (버그) 선택한 장비만 취합 열로 쓰면 나머지 호기 값이 통째로 빠져
    변경보고서에 '삭제'로 잘못 찍혔다. 열은 항상 전체 호기여야 한다.
    """
    import shutil
    from param_manager import collate, extract_io, workdirs
    with tempfile.TemporaryDirectory() as tmp:
        machines_all = ["AOI-17", "AOI-19", "AOI-25"]
        # 양식 1개(파라미터 1행)
        recipe = "PI3"
        run_dir = workdirs.form_run_dir(tmp, recipe, "20260101_000000")
        form = workdirs.form_final_path(run_dir, recipe, "AOI-19", "20260101_000000")
        rec = {"PI": recipe, "Recipe": "PI", "Zone": "PAD", "Alg": "Contrast",
               "Parameter": "Contrast Delta - Bright", "비고": ""}
        extract_io.write_snapshot(form, [rec], machines=[], sheet_name="PI_ALL",
                                  extracts=[{"src_file": "Zones/PAD.ini",
                                             "section": "Contrast",
                                             "key": "Delta", "transform": "RAW"}],
                                  stage="final", level=recipe)
        # 직전 취합본: 세 호기 모두 값이 있음
        prev_res = collate.build_collation(tmp, [recipe], [], machines_all)
        made = {r: v for r, v in prev_res.items() if not v.missing_form}
        for r, v in made.items():
            for row in v.records:
                for m, val in zip(machines_all, ["17값", "19값", "25값"]):
                    row[m] = val
        prev_path = workdirs.collate_path(tmp, "20260101_010000")
        collate.write_collation(prev_path, made, machines_all)

        # 감시 회차: AOI-19 만 선택했지만 **열은 전체 호기**로 취합(고친 동작)
        out = collate.build_collation(tmp, [recipe], [], machines_all,
                                      prev_collate_path=prev_path)
        made2 = {r: v for r, v in out.items() if not v.missing_form}
        new_path = workdirs.collate_path(tmp, "20260101_020000")
        collate.write_collation(new_path, made2, machines_all)

        sheets, cols = collate.load_collation(new_path)
        for m in machines_all:
            assert m in cols, f"{m} 열이 사라짐: {cols}"
        row = sheets[recipe][0]
        assert row["AOI-17"] == "17값", row     # 선택 안 한 호기 값 유지
        assert row["AOI-25"] == "25값", row

        # 변경보고서에 '삭제'가 찍히면 안 된다
        res = watcher.compare_and_report(tmp, prev_path, new_path)
        assert res.removed == 0, f"선택 안 한 호기가 삭제로 잡힘: {res.summary()}"
        assert not res.has_change, res.summary()
        shutil.rmtree(os.path.join(tmp, "양식"), ignore_errors=True)
    print("  선택 안 한 호기 값 유지 + 오탐 '삭제' 없음 OK")


def test_plan_roundtrip():
    """무인 수집 계획이 저장·복원되어야 무인 회차가 선택창 없이 돈다."""
    from param_manager.collector import CollectPlan
    with tempfile.TemporaryDirectory() as tmp:
        assert watcher.plan_from_dict({}) is None, "계획 없으면 None"
        plan = CollectPlan(job_keyword="PI", job_name="JobA", setup_name="Setup1",
                           recipe_names=["PI3"], recipe_map={"PI3": ["PI3_FOLDER"]})
        s, st = watcher.load_settings(tmp)
        s.plan = watcher.plan_to_dict(plan)
        watcher.save_settings(tmp, s, st)
        s2, _ = watcher.load_settings(tmp)
        back = watcher.plan_from_dict(s2.plan)
        assert back is not None
        assert back.job_keyword == "PI" and back.setup_name == "Setup1"
        assert back.recipe_map == {"PI3": ["PI3_FOLDER"]}, back.recipe_map
    print("  무인 수집 계획 저장/복원 OK")


def test_report_listing_newest_first():
    """모두가 같은 결과를 보려면 공유 폴더의 보고서를 최신순으로 찾을 수 있어야 한다."""
    with tempfile.TemporaryDirectory() as tmp:
        assert watcher.list_reports(tmp) == []
        assert watcher.latest_report(tmp) is None
        d = watcher.watch_dir(tmp)
        for name in ("변경보고서_20260101_010000.xlsx",
                     "변경보고서_20260804_120000.xlsx",
                     "감시로그.txt"):                      # 보고서 아닌 파일은 제외
            with open(os.path.join(d, name), "w", encoding="utf-8") as fh:
                fh.write("x")
        reports = watcher.list_reports(tmp)
        assert len(reports) == 2, reports
        assert os.path.basename(reports[0]) == "변경보고서_20260804_120000.xlsx"
        assert watcher.latest_report(tmp) == reports[0]
    print("  변경 보고서 목록(최신순·비보고서 제외) OK")


def test_shared_state_carries_result_for_others():
    """감시를 안 돌린 사람도 공유 설정에서 회차 결과를 읽어 알림을 띄울 수 있어야 한다."""
    with tempfile.TemporaryDirectory() as tmp:
        s = watcher.WatchSettings(enabled=True, machines=["AOI-19"], recipes=["PI3"])
        st = watcher.WatchState()
        st = watcher.record_run(tmp, s, st, ok=True, note="값변경 3건")
        # 다른 PC 의 앱이 파일만 읽어도 '켜져 있음 + 마지막 결과'를 알 수 있다
        s2, st2 = watcher.load_settings(tmp)
        assert s2.enabled is True, "설정은 공유 — 모두에게 ON 으로 보여야 함"
        assert st2.last_result == "값변경 3건"
        assert st2.last_run == st.last_run and st2.last_run != ""
    print("  공유 상태로 다른 사용자도 ON·결과 인지 OK")


def test_job_matching_is_per_level_tolerant():
    """레벨 하나가 안 맞아도 **되는 레벨은 수집**되어야 한다.

    (버그) collector 의 plan 재사용은 레벨 하나라도 실패하면 전체를 포기해
    '자동 매칭 실패'로 장비 전체가 건너뛰어졌다. equip_app._watch_collect 의
    auto_match 가 레벨 단위로 다시 시도한다 — 여기서는 그 규칙만 검증.
    """
    from pathlib import Path

    from param_manager import collector
    with tempfile.TemporaryDirectory() as tmp:
        jobs = []
        for name in ("PI3_MAIN", "RDL2_PROD", "기타작업"):
            d = Path(tmp) / name
            d.mkdir()
            jobs.append(d)
        recipe_map = {"PI3": ["PI3_MAIN"]}          # RDL2 는 계획에 없음

        def auto_match(job_dirs, levels):           # _watch_collect 와 같은 규칙
            out = {}
            for lvl in levels:
                names = recipe_map.get(lvl) or []
                sel = []
                if names:
                    sel, _ = collector.match_recipes_by_names(job_dirs, names)
                if not sel:
                    cands = [d for d in job_dirs
                             if collector.contains_keyword(d.name, lvl)]
                    if len(cands) == 1:
                        sel = cands
                if sel:
                    out[lvl] = sel
            return out

        got = auto_match(jobs, ["PI3", "RDL2"])
        assert "PI3" in got, got                    # 계획으로 매칭
        assert "RDL2" in got, "레벨 이름 폴백으로 찾아야 함"
        assert got["RDL2"][0].name == "RDL2_PROD"
        # 없는 레벨은 조용히 제외(추측하지 않음)
        got2 = auto_match(jobs, ["PI3", "PI9"])
        assert set(got2) == {"PI3"}, got2
    print("  Job 폴더 레벨별 매칭(계획→레벨명 폴백·애매하면 제외) OK")


def test_job_relative_and_machine_path():
    """폴더를 한 번 지정하면 IP 만 바꿔 **모든 장비**에 적용돼야 한다."""
    rel = watcher.job_relative(r"\\10.0.0.5\c$\Job\PI3_MAIN\Setup1\Recipes\PI3")
    assert rel == r"PI3_MAIN\Setup1\Recipes\PI3", rel
    # 슬래시 표기·소문자 job 도 인식
    assert watcher.job_relative("//10.0.0.5/c$/job/A/B") == r"A\B"
    # Job 이 없으면 지정 실패로 본다(엉뚱한 경로 저장 방지)
    assert watcher.job_relative(r"D:\어딘가\PI3") == ""
    assert watcher.job_relative("") == ""
    # 같은 상대경로가 다른 장비에 그대로 적용된다
    assert watcher.machine_recipe_dir("10.0.0.9", rel) == \
        r"\\10.0.0.9\c$\Job\PI3_MAIN\Setup1\Recipes\PI3"
    assert watcher.machine_recipe_dir("10.0.0.9", "") == r"\\10.0.0.9\c$\Job"
    print("  감시 폴더 지정: Job 상대경로 추출 + 장비별 경로 조립 OK")


def test_recipe_paths_are_per_machine():
    """장비마다 Job 구조가 다르므로 **호기별로** 지정·조회되어야 한다."""
    with tempfile.TemporaryDirectory() as tmp:
        s, st = watcher.load_settings(tmp)
        assert s.recipe_paths == {}
        s.recipe_paths = {
            "AOI-17": {"PI3": r"PI3_MAIN\Setup1\Recipes\PI3", "RDL2": ""},
            "AOI-19": {"PI3": r"PI_JOB\Recipes\PI3_A"},
        }
        watcher.save_settings(tmp, s, st)
        s2, _ = watcher.load_settings(tmp)
        # 빈 값은 저장하지 않는다(미지정 = 자동 매칭 폴백)
        assert s2.recipe_paths["AOI-17"] == {"PI3": r"PI3_MAIN\Setup1\Recipes\PI3"}
        # 같은 레시피라도 호기마다 다른 경로가 나와야 한다
        assert watcher.path_for(s2.recipe_paths, "AOI-17", "PI3") == \
            r"PI3_MAIN\Setup1\Recipes\PI3"
        assert watcher.path_for(s2.recipe_paths, "AOI-19", "PI3") == \
            r"PI_JOB\Recipes\PI3_A"
        assert watcher.path_for(s2.recipe_paths, "AOI-17", "RDL2") == ""
        assert watcher.path_for(s2.recipe_paths, "AOI-99", "PI3") == "", "미지정 호기"
    print("  감시 폴더 호기별 지정/조회 OK")


def test_recipe_paths_legacy_flat_migrates():
    """구 버전의 전 장비 공통 지정({레시피: 경로})도 계속 동작해야 한다."""
    norm = watcher.normalize_recipe_paths({"PI3": r"A\B\PI3"})
    assert norm == {watcher.ANY_MACHINE: {"PI3": r"A\B\PI3"}}, norm
    # 호기별 지정이 없으면 공통 지정으로 폴백
    assert watcher.path_for(norm, "AOI-17", "PI3") == r"A\B\PI3"
    # 호기별 지정이 있으면 그쪽이 우선
    mixed = watcher.normalize_recipe_paths(
        {"PI3": r"A\B\PI3", "AOI-17": {"PI3": r"X\Y"}})
    assert watcher.path_for(mixed, "AOI-17", "PI3") == r"X\Y"
    assert watcher.path_for(mixed, "AOI-19", "PI3") == r"A\B\PI3"
    print("  구 버전 공통 지정 하위호환(호기별 우선) OK")


def test_interval_presets_roundtrip():
    """주기는 프리셋(30분~24시간)에서 고른다 — 라벨↔시간 왕복이 맞아야 한다."""
    labels = [lab for _v, lab in watcher.INTERVAL_CHOICES]
    assert labels == ["30분", "1시간", "2시간", "4시간", "6시간", "8시간",
                      "12시간", "24시간"], labels
    for v, lab in watcher.INTERVAL_CHOICES:
        assert watcher.interval_label(v) == lab
        assert watcher.interval_from_label(lab) == float(v)
    assert watcher.interval_from_label("없는값") == float(watcher.DEFAULT_INTERVAL_HOURS)
    assert watcher.interval_label(3) == "3시간", "프리셋 밖이면 숫자로"
    assert watcher.interval_label(None) == watcher.interval_label(
        watcher.DEFAULT_INTERVAL_HOURS)
    print("  주기 프리셋 라벨 왕복 OK")


def test_first_run_at_same_start_end():
    """'9시 ~ 9시' = 매일 9시에 시작 — 첫 회차가 다음 9시여야 한다."""
    s = watcher.WatchSettings(window_start=9, window_end=9)
    at = watcher.first_run_at(datetime(2026, 8, 11, 8, 30), s)
    assert (at.day, at.hour, at.minute) == (11, 9, 0), at      # 오늘 9시
    at = watcher.first_run_at(datetime(2026, 8, 11, 9, 30), s)
    assert (at.day, at.hour) == (12, 9), at                    # 이미 지났으면 내일
    at = watcher.first_run_at(datetime(2026, 8, 11, 14, 30), s)
    assert (at.day, at.hour) == (12, 9), at
    # 시간대 제한이 없으면(0~0) 지금 바로
    free = watcher.WatchSettings()
    now = datetime(2026, 8, 11, 14, 30)
    assert watcher.first_run_at(now, free) == now
    # 진짜 창(9~18)이면 지금부터 — 창 밖이면 in_window 가 알아서 미룬다
    win = watcher.WatchSettings(window_start=9, window_end=18)
    early = datetime(2026, 8, 11, 7, 0)
    assert watcher.first_run_at(early, win) == early
    assert not watcher.in_window(early, win)
    assert watcher.in_window(datetime(2026, 8, 11, 9, 0), win)
    print("  시작=끝 시간대의 첫 실행 시각 OK")


def test_machine_recipes_are_per_machine():
    """감시 대상 = 호기별 레시피. 장비마다 다른 레시피를 감시할 수 있어야 한다."""
    s = watcher.WatchSettings()
    s.recipe_paths = watcher.normalize_recipe_paths({
        "AOI-11": {"PI3": r"A\Recipes\PI3"},
        "AOI-12": {"PI3": r"B\Recipes\PI3", "RDL1": r"B\Recipes\RDL1"},
        "AOI-13": {"PI3": ""},                   # 미지정 = 대상 아님
    })
    mr = watcher.machine_recipes(s)
    assert mr == {"AOI-11": ["PI3"], "AOI-12": ["PI3", "RDL1"]}, mr
    assert sorted(watcher.watch_targets(s)) == [
        ("AOI-11", "PI3"), ("AOI-12", "PI3"), ("AOI-12", "RDL1")]
    # 구 목록(machines/recipes)도 지정과 일치하게 맞춰 준다(하위호환 참조용)
    watcher.sync_selection(s)
    assert s.machines == ["AOI-11", "AOI-12"], s.machines
    assert s.recipes == ["PI3", "RDL1"], s.recipes
    # 구 설정(지정 없이 전역 목록만)도 조합으로 읽힌다
    old = watcher.WatchSettings(machines=["AOI-6"], recipes=["PI2", "PI3"])
    assert watcher.machine_recipes(old) == {"AOI-6": ["PI2", "PI3"]}
    assert watcher.machine_recipes(watcher.WatchSettings()) == {}
    print("  호기별 감시 레시피 OK")


def test_job_folder_resolves_subrecipes():
    """**Job 폴더만 지정**하면 그 아래 Setup/Recipes/레시피를 자동으로 찾아야 한다.

    (_collect_fixed_dir 의 대상 폴더 해석 규칙과 동일 — 어느 단계를 지정하든 동작)
    """
    from pathlib import Path

    from param_manager import collector
    with tempfile.TemporaryDirectory() as tmp:
        job = Path(tmp) / "PI3_MAIN"
        recipes_root = job / "Setup1" / "Recipes"
        for rname in ("PI3", "PI3_BUBBLE"):
            d = recipes_root / rname
            (d / "Zones").mkdir(parents=True)
            (d / "GlobalRTP.ini").write_text("[A]\nx=1\n", encoding="utf-8")
        (recipes_root / "빈폴더").mkdir()            # 설정파일 없음 → 제외돼야 함

        def is_recipe_dir(d):
            return any((d / f).is_file() for f in collector.FIXED_FILES) or \
                (d / "Zones").is_dir()

        def resolve(src):                            # _collect_fixed_dir 과 같은 규칙
            if is_recipe_dir(src):
                return [src]
            out = []
            for _s, rr in collector.find_setup_candidates(src):
                out += [d for d in collector.list_dirs(rr) if is_recipe_dir(d)]
            if not out:
                out = [d for d in collector.list_dirs(src) if is_recipe_dir(d)]
            return out

        # ① Job 폴더 지정 → 하위 레시피 2개 자동 발견(빈 폴더 제외)
        got = sorted(d.name for d in resolve(job))
        assert got == ["PI3", "PI3_BUBBLE"], got
        # ② Recipes 폴더를 지정해도 동일
        assert sorted(d.name for d in resolve(recipes_root)) == ["PI3", "PI3_BUBBLE"]
        # ③ Recipe 폴더를 직접 지정하면 그것만
        assert [d.name for d in resolve(recipes_root / "PI3")] == ["PI3"]
        # 실제로 복사 대상 파일이 잡히는지
        planned = collector.plan_files(resolve(job))
        assert planned and any("GlobalRTP.ini" in p[2] for p in planned), planned
    print("  Job 폴더 지정 → 하위 레시피 자동 탐색(어느 단계든 동작) OK")


def test_log_appends():
    with tempfile.TemporaryDirectory() as tmp:
        watcher.append_log(tmp, "회차 시작")
        watcher.append_log(tmp, "변경 없음")
        with open(watcher.log_path(tmp), encoding="utf-8") as fh:
            body = fh.read()
        assert "회차 시작" in body and "변경 없음" in body
        assert body.count("\n") == 2
    print("  감시 로그 누적 OK")


if __name__ == "__main__":
    for t in [test_settings_roundtrip_and_defaults, test_settings_survives_corrupt_file,
              test_time_window, test_should_run_interval,
              test_backoff_delays_after_failures, test_record_run_updates_state,
              test_unstable_file_detection, test_drop_unstable_filters_sources,
              test_compare_and_report_detects_change, test_no_change_makes_no_report,
              test_first_run_has_no_baseline, test_connection_check_and_guide,
              test_connection_check_is_bounded_by_timeout,
              test_connection_check_progress_and_cancel,
              test_selected_targets_persist, test_stale_selection_is_filtered,
              test_unselected_machines_keep_values, test_plan_roundtrip,
              test_report_listing_newest_first,
              test_shared_state_carries_result_for_others,
              test_job_matching_is_per_level_tolerant,
              test_job_relative_and_machine_path, test_recipe_paths_are_per_machine,
              test_recipe_paths_legacy_flat_migrates,
              test_interval_presets_roundtrip, test_first_run_at_same_start_end,
              test_machine_recipes_are_per_machine,
              test_job_folder_resolves_subrecipes,
              test_log_appends]:
        run(t)
    print(f"==== {PASS}/{PASS + FAIL} passed ====")
    sys.exit(1 if FAIL else 0)
