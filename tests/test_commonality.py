"""Commonality 조사 헤드리스 테스트 — 계획 엑셀·폴더 해석·안전복사·구조 diff·
Lot 취합·호기 비교(과반수 이탈 색칠). 로컬 가짜 Scanresult 트리 사용."""
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402

from param_manager import (collate, commonality, engine, formbuilder,  # noqa: E402
                           ini_parser)

ZONE = ("[General]\nZoneName = PI Opening\n"
        "[Surface]\nHigh_Delta = {delta}\nMinDefectWidth_um = 3.5\n")
OPTIC = "[LIGHT]\n[Scan2d1]\nAlg = Scan2d\nLightSrc_Top = 100\n"


def _make_wafer(base, machine, device_folder, lot, sm, wafer, delta=25):
    """가짜 Scanresult 웨이퍼 폴더 생성. 반환: 웨이퍼 config 폴더."""
    wdir = (Path(base) / machine / "Scanresult" / device_folder / lot / sm / wafer)
    (wdir / "Zones").mkdir(parents=True)
    (wdir / "Zones" / "Z.ini").write_text(ZONE.format(delta=delta), encoding="utf-8")
    (wdir / "OpticPreset.ini").write_text(OPTIC, encoding="utf-8")
    (wdir / "RTP.txt").write_text("dummy rtp", encoding="utf-8")
    return wdir


def test_plan_template_roundtrip_and_filter():
    with tempfile.TemporaryDirectory() as tmp:
        p = os.path.join(tmp, "plan.xlsx")
        commonality.create_plan_template(p, rows=[
            {"디바이스명": "S6WH11001-00001", "공정번호": "6321", "S/M": "HPG", "AOI호기": "AOI-6"},
            {"디바이스명": "S6WH11001-00001", "공정번호": "6323", "S/M": "TVS", "AOI호기": "AOI-9"},
        ])
        rows = commonality.read_plan(p)
        assert len(rows) == 2
        assert rows[0]["공정번호"] == "6321"
        only6 = commonality.filter_plan_for_machine(rows, "aoi_6")   # 정규화 매칭
        assert len(only6) == 1 and only6[0]["S/M"] == "HPG"
    print("  commonality OK: 계획 엑셀 왕복 + 호기 필터(정규화)")


def test_resolve_plan_found_and_missing():
    with tempfile.TemporaryDirectory() as tmp:
        _make_wafer(tmp, "AOI-6", "2D@R2-DEVA-1_0855360PD-0A", "6321", "HPG", "CX02")
        _make_wafer(tmp, "AOI-6", "2D@R2-DEVA-1_0855360PD-0A", "6321", "HPG", "CX01")
        root = commonality.scanresult_root(tmp, "AOI-6")
        plan = [
            {"디바이스명": "DEVA-1", "공정번호": "6321", "S/M": "HPG", "AOI호기": "AOI-6"},
            {"디바이스명": "DEVA-1", "공정번호": "9999", "S/M": "HPG", "AOI호기": "AOI-6"},
        ]
        lots = commonality.resolve_plan(root, plan)
        found = [l for l in lots if l.exists]
        missing = [l for l in lots if not l.exists]
        assert len(found) == 1 and len(missing) == 1
        # 이름순 첫 웨이퍼 = CX01
        assert found[0].wafer_dir.name == "CX01"
        assert found[0].has_zones and found[0].has_rtp and found[0].has_optic
        assert "공정" in missing[0].reason
    print("  commonality OK: 폴더 해석(디바이스+LOT+S/M) + 이름순 첫 웨이퍼 + 실패사유")


def test_real_world_folder_variants():
    """실제 환경 변형 흡수: Scanresult_260401 폴더명 + AOI-9 vs AOI-09(0패딩)."""
    with tempfile.TemporaryDirectory() as tmp:
        # 실제 경로: W:/AOI-9/Scanresult_260401/2D@RE-DEV..._x/6412/HCH/wafer
        w = (Path(tmp) / "AOI-9" / "Scanresult_260401"
             / "2D@RE-L6WZ30001-00001_0852445PD-0C" / "6412" / "HCH" / "65349540001")
        (w / "Zones").mkdir(parents=True)
        (w / "Zones" / "Z.ini").write_text(ZONE.format(delta=25), encoding="utf-8")
        (w / "OpticPreset.ini").write_text(OPTIC, encoding="utf-8")
        (w / "RTP.txt").write_text("x", encoding="utf-8")

        # base=tmp, 선택 호기="AOI-09"(0패딩 다름) → 루트가 AOI-9/Scanresult_260401 로 해석
        root = commonality.scanresult_root(tmp, "AOI-09")
        assert root.name == "Scanresult_260401"
        # 계획 필터도 AOI-9 == AOI-09
        plan = [{"디바이스명": "L6WZ30001-00001", "공정번호": "6412", "S/M": "HCH",
                 "AOI호기": "AOI-09"}]
        assert len(commonality.filter_plan_for_machine(plan, "AOI-9")) == 1
        # 폴더 해석 성공(디바이스명 포함 매칭)
        lot = commonality.resolve_lot(root, "L6WZ30001-00001", "6412", "HCH", "AOI-9")
        assert lot.exists and lot.wafer_dir.name == "65349540001"
        # LOT 오타(6421)면 실패 + 사유
        bad = commonality.resolve_lot(root, "L6WZ30001-00001", "6421", "HCH", "AOI-9")
        assert not bad.exists and "공정" in bad.reason
    print("  commonality OK: Scanresult_260401 폴더명 + AOI-9/AOI-09 + LOT 오타 사유")


def test_multiple_scanresult_backups():
    """호기 폴더 아래 Scanresult 백업본이 여러 개면 전부 탐색해 Lot 을 찾는다."""
    with tempfile.TemporaryDirectory() as tmp:
        mdir = Path(tmp) / "AOI-9"
        # 현재본엔 6412 없음, 백업본에만 있음
        (mdir / "Scanresult" / "2D@X-DEVZ_0A" / "9999" / "AAA" / "w0" / "Zones").mkdir(parents=True)
        good = (mdir / "SCANRESULT_BACKUP_260805" / "2D@X-DEVZ_0A" / "6412" / "HCH" / "w1")
        (good / "Zones").mkdir(parents=True)
        (good / "Zones" / "Z.ini").write_text(ZONE.format(delta=25), encoding="utf-8")
        (good / "RTP.txt").write_text("x", encoding="utf-8")
        # 또 다른 백업본
        (mdir / "Scanresult_260402").mkdir(parents=True)

        roots = commonality.scanresult_roots(tmp, "AOI-9")
        assert len(roots) == 3, [p.name for p in roots]   # 3개 모두 탐색 대상
        # 호기 폴더만 줘도(=tmp 아래 AOI-9) 백업본의 6412 를 찾음
        lot = commonality.resolve_lot(roots, "DEVZ", "6412", "HCH", "AOI-9")
        assert lot.exists and lot.wafer_dir.name == "w1", lot.reason
        assert "BACKUP" in str(lot.wafer_dir)
    print("  commonality OK: Scanresult 백업본 다중 탐색(호기 폴더만 지정)")


def test_multiple_recipe_folders_and_multimachine_filter():
    """같은 디바이스가 여러 2D@ 레시피 폴더로 나뉘고, 공정 폴더가 두 번째
    폴더에만 있어도 찾아야 한다. 또 'AOI-4,6,9' 한 칸 여러 호기 필터."""
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp) / "AOI-9" / "Scanresult_260401"
        # 디바이스 L6WZ30001-00001 의 레시피 폴더 2개: _0A(6412 없음), _0C(6412 있음)
        (base / "2D@RE-L6WZ30001-00001_0852445PD-0A" / "9999" / "ZZZ" / "w0"
         / "Zones").mkdir(parents=True)
        good = (base / "2D@RE-L6WZ30001-00001_0852445PD-0C" / "6412" / "HCH"
                / "65349540001")
        (good / "Zones").mkdir(parents=True)
        (good / "Zones" / "Z.ini").write_text(ZONE.format(delta=25), encoding="utf-8")
        (good / "OpticPreset.ini").write_text(OPTIC, encoding="utf-8")
        (good / "RTP.txt").write_text("x", encoding="utf-8")

        root = commonality.scanresult_root(tmp, "AOI-9")
        lot = commonality.resolve_lot(root, "L6WZ30001-00001", "6412", "HCH", "AOI-9")
        assert lot.exists, f"두 번째 레시피 폴더의 6412 를 못 찾음: {lot.reason}"
        assert lot.wafer_dir.name == "65349540001"

        # 여러 호기 한 칸: AOI-4,6,9 → AOI-9 로 필터되어야 함
        plan = [{"디바이스명": "D", "공정번호": "6412", "S/M": "YYH", "AOI호기": "AOI-4,6,9"},
                {"디바이스명": "D", "공정번호": "6412", "S/M": "YYA", "AOI호기": "AOI-05"}]
        got = commonality.filter_plan_for_machine(plan, "AOI-09")
        assert len(got) == 1 and got[0]["S/M"] == "YYH"
    print("  commonality OK: 여러 레시피 폴더 탐색 + 여러 호기 한 칸 필터")


def test_intermediate_level_and_no_numeric_mismatch():
    """공정 폴더가 중간 폴더 아래 있어도 찾고(BFS), 6412 가 64120 에 오매칭 안 됨."""
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp) / "AOI-9" / "Scanresult"
        dev = base / "2D@X-DEVQ-1_0A"
        # 헷갈리게 하는 폴더: 64120(오매칭 유발) + 중간폴더 SUB 아래 진짜 6412
        (dev / "64120" / "AAA" / "w0" / "Zones").mkdir(parents=True)
        good = dev / "SUB" / "6412" / "HCH" / "wafer1"
        (good / "Zones").mkdir(parents=True)
        (good / "Zones" / "Z.ini").write_text(ZONE.format(delta=25), encoding="utf-8")
        (good / "RTP.txt").write_text("x", encoding="utf-8")
        root = commonality.scanresult_root(tmp, "AOI-9")
        lot = commonality.resolve_lot(root, "DEVQ-1", "6412", "HCH", "AOI-9")
        assert lot.exists and lot.wafer_dir.name == "wafer1", lot.reason
        assert "6412" in str(lot.wafer_dir) and "64120" not in str(lot.wafer_dir)
    print("  commonality OK: 중간 폴더 BFS + 숫자 오매칭 방지(6412≠64120)")


def test_read_plan_legacy_header():
    """구 템플릿 헤더(LOT번호)도 공정번호로 읽혀야 한다(하위호환)."""
    with tempfile.TemporaryDirectory() as tmp:
        p = os.path.join(tmp, "old.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Lot목록"
        ws.append(["디바이스명", "LOT번호", "S/M", "AOI호기"])   # 구 헤더
        ws.append(["S6WC61001-00001", 6412, "YYA", "AOI-09"])    # 숫자 6412
        wb.save(p)
        rows = commonality.read_plan(p)
        assert rows[0]["공정번호"] == "6412"     # 별칭 + 문자열화
        assert commonality.filter_plan_for_machine(rows, "AOI-9")[0]["S/M"] == "YYA"
    print("  commonality OK: 구 헤더(LOT번호) 하위호환 + 숫자 공정번호")


def test_copy_lot_read_only():
    with tempfile.TemporaryDirectory() as tmp:
        w = _make_wafer(tmp, "AOI-6", "2D@DEVB_x", "6400", "HPG", "CX10")
        root = commonality.scanresult_root(tmp, "AOI-6")
        lot = commonality.resolve_lot(root, "DEVB", "6400", "HPG", "AOI-6")
        assert lot.exists
        before = {p: (p.read_bytes(), p.stat().st_mtime_ns)
                  for p in w.rglob("*") if p.is_file()}
        dest = os.path.join(tmp, "staging")
        res = commonality.copy_lot(lot, dest, verify=True)
        assert res["count"] == 3       # Z.ini + OpticPreset.ini + RTP.txt
        # 원본 불변
        for p, (content, mtime) in before.items():
            assert p.read_bytes() == content and p.stat().st_mtime_ns == mtime
        # 사본은 로컬 staging 에만
        assert lot.label == "HPG"           # 라벨 = S/M 만
        assert (Path(dest) / "HPG" / "Zones" / "Z.ini").is_file()
        assert (Path(dest) / "HPG" / "RTP.txt").is_file()
    print("  commonality OK: 안전복사(원본 read-only) + 대상만 복사")


def test_multi_recipe_copy_preserves_structure():
    """다중 레시피 안전복사: RecipesInfo/ActiveScenarioOptics/Recipe*- 파일 + Recipe*-Zones/
    구조 보존, 사본에서 레시피별 값 분리 파싱까지."""
    with tempfile.TemporaryDirectory() as tmp:
        w = _make_wafer(tmp, "AOI-6", "2D@DEVM_x", "6400", "HPG", "CX20")  # recipe-1 delta=25
        (w / "RecipesInfo.ini").write_text(
            "[Recipe-1]\nName=PI_Bubble\n[Recipe-2]\nName=PI\n[Recipes]\nCount=2\n",
            encoding="utf-8")
        (w / "ActiveScenarioOptics.ini").write_text(
            "[z]\nScenarioName=Scan2d\nOpticsName=Scan2d1\nOpticId=aaa\n", encoding="utf-8")
        (w / "Recipe2-OpticPreset.ini").write_text(
            "[LIGHT]\n[Scan2d1]\nAlg=Scan2d\nLightSrc_Top=100\n", encoding="utf-8")
        (w / "Recipe2-ActiveScenarioOptics.ini").write_text(
            "[z]\nScenarioName=Scan2d\nOpticsName=Scan2d1\nOpticId=aaa\n", encoding="utf-8")
        (w / "Recipe2-Zones").mkdir()
        (w / "Recipe2-Zones" / "Z.ini").write_text(
            "[General]\nZoneName=PI Opening\n[Surface]\nHigh_Delta=99\n", encoding="utf-8")
        root = commonality.scanresult_root(tmp, "AOI-6")
        lot = commonality.resolve_lot(root, "DEVM", "6400", "HPG", "AOI-6")
        dest = os.path.join(tmp, "staging")
        commonality.copy_lot(lot, dest, verify=True)
        base = Path(dest) / "HPG"
        for rel in ("RecipesInfo.ini", "Recipe2-OpticPreset.ini",
                    "Recipe2-ActiveScenarioOptics.ini", "ActiveScenarioOptics.ini",
                    "Recipe2-Zones/Z.ini", "Zones/Z.ini"):
            assert (base / rel).is_file(), f"복사 누락: {rel}"
        recs = commonality.detect_recipes([("HPG", base)])
        assert [r["name"] for r in recs] == ["PI_Bubble", "PI"], recs
        p1, _ = commonality.parse_lots([("HPG", base)], level="PI_Bubble", recipe_prefix="")
        p2, _ = commonality.parse_lots([("HPG", base)], level="PI", recipe_prefix="Recipe2-")

        def hd(pv):
            for r in pv:
                if r["extract"].get("key") == "High_Delta":
                    return str(r["raws"].get("HPG"))
            return None
        assert hd(p1) == "25" and hd(p2) == "99"     # 레시피별 값 분리(사본)
    print("  commonality OK: 다중 레시피 안전복사 구조 보존 + 사본 파싱 분리")


def test_structure_diff():
    with tempfile.TemporaryDirectory() as tmp:
        w1 = _make_wafer(tmp, "AOI-6", "2D@DEVC_x", "1", "HPG", "CXA")
        w2 = _make_wafer(tmp, "AOI-6", "2D@DEVC_x", "2", "HPG", "CXB")
        # w2 는 Zones ini 에 항목 하나 제거(구조 불일치 유발)
        (w2 / "Zones" / "Z.ini").write_text(
            "[General]\nZoneName = PI Opening\n[Surface]\nHigh_Delta = 25\n",
            encoding="utf-8")
        diff = commonality.structure_diff([("L1", w1), ("L2", w2)], level="PI3")
        assert diff["identical"] is False
        # L2 에 MinDefectWidth 관련 항목이 없음
        assert any("MinDefectWidth" in m for m in diff["lots"]["L2"]["missing"])
        assert diff["lots"]["L1"]["missing"] == []
    print("  commonality OK: Lot 간 구조 diff(빠진 항목 검출)")


def _build_form(tmp, pivot_rows, recipe="PI3"):
    init = os.path.join(tmp, f"init_{recipe}.xlsx")
    formbuilder.build_initial_workbook(pivot_rows, init, level=recipe)
    final = os.path.join(tmp, f"form_{recipe}.xlsx")
    formbuilder.build_final_from_initial(init, final, level=recipe)
    return final


def test_collate_lots_and_result_roundtrip():
    with tempfile.TemporaryDirectory() as tmp:
        w1 = _make_wafer(tmp, "AOI-6", "2D@DEVD_x", "6501", "HPG", "CX1", delta=25)
        w2 = _make_wafer(tmp, "AOI-6", "2D@DEVD_x", "6502", "HPG", "CX2", delta=99)
        lots = [("6501_HPG", w1), ("6502_HPG", w2)]
        pivot, labels = commonality.parse_lots(lots, level="PI3")
        assert set(labels) == {"6501_HPG", "6502_HPG"}
        form = _build_form(tmp, pivot)
        res = commonality.collate_lots("PI3", form, pivot, labels)
        # High_Delta → Contrast Delta - Bright 가 Lot 별로 다른 값
        row = next(r for r in res.records
                   if engine._s(r["Parameter"]) == "Contrast Delta - Bright")
        assert engine._s(row["6501_HPG"]) == "25"
        assert engine._s(row["6502_HPG"]) == "99"
        # 결과 엑셀 왕복
        out = os.path.join(tmp, "조사_AOI-6.xlsx")
        commonality.write_lot_result(out, "PI3", "AOI-6", res, labels,
                                     scan_times={"6501_HPG": "2026-08-01 09:10",
                                                 "6502_HPG": "2026-08-05 21:30"})
        import openpyxl
        wb = openpyxl.load_workbook(out)
        ws = wb[wb.sheetnames[0]]
        head = [c.value for c in ws[1]]
        scan = [c.value for c in ws[2]]
        wb.close()
        # 1행 표시 헤더는 상위/하위 Recipe (내부 키는 PI/Recipe 그대로)
        assert head[0] == "상위 Recipe" and head[1] == "하위 Recipe", head
        assert "PI" not in head and "Recipe" not in head, head
        # **파라미터 첫 행 = Scan일자**, 값은 각 S/M 열 아래
        assert scan[head.index("Parameter")] == commonality.SCAN_ROW_LABEL, scan
        assert scan[-2:] == ["2026-08-01 09:10", "2026-08-05 21:30"], scan
        data = commonality.read_lot_result(out)
        assert data["machine"] == "AOI-6" and set(data["lots"]) == set(labels)
        assert data["scan_times"]["6502_HPG"] == "2026-08-05 21:30"
        # Scan일자 행은 파라미터 목록에 섞이지 않는다
        assert all(engine._s(r.get("Parameter")) != commonality.SCAN_ROW_LABEL
                   for r in data["records"])
        # read 는 내부 키로 되돌린다(다운스트림 비교가 PI/Recipe/Zone 로 접근)
        assert all("PI" in r and "Recipe" in r for r in data["records"])
    print("  commonality OK: Lot 취합 + 호기 결과 엑셀 왕복 + 표시헤더(상위/하위 Recipe)")


def test_build_comparison_outliers():
    with tempfile.TemporaryDirectory() as tmp:
        # AOI-6: 두 Lot(25/99), AOI-9: 한 Lot(25) → 99 만 과반수(25) 이탈
        w6a = _make_wafer(tmp, "AOI-6", "2D@DEVE_x", "6601", "HPG", "CX1", delta=25)
        w6b = _make_wafer(tmp, "AOI-6", "2D@DEVE_x", "6602", "HPG", "CX2", delta=99)
        w9 = _make_wafer(tmp, "AOI-9", "2D@DEVE_x", "6603", "TVS", "CX3", delta=25)

        piv6, lab6 = commonality.parse_lots([("6601_HPG", w6a), ("6602_HPG", w6b)],
                                            level="PI3")
        form = _build_form(tmp, piv6)
        r6 = commonality.collate_lots("PI3", form, piv6, lab6)
        f6 = os.path.join(tmp, "AOI-6.xlsx")
        commonality.write_lot_result(f6, "PI3", "AOI-6", r6, lab6)

        piv9, lab9 = commonality.parse_lots([("6603_TVS", w9)], level="PI3")
        r9 = commonality.collate_lots("PI3", form, piv9, lab9)
        f9 = os.path.join(tmp, "AOI-9.xlsx")
        commonality.write_lot_result(f9, "PI3", "AOI-9", r9, lab9)

        comp = commonality.build_comparison([f6, f9])
        assert comp["columns"][:2] == ["S/M", "호기"]
        assert len(comp["rows"]) == 3          # 세 (Lot,호기) 행
        # Contrast Delta - Bright 열이 변경 파라미터
        cd = next(c for c in comp["columns"] if "Contrast Delta - Bright" in c)
        assert cd in comp["changed_params"]
        # 99 를 가진 행만 이탈
        out_rows = {i for (i, col) in comp["outliers"] if col == cd}
        vals = [engine._s(r.get(cd)) for r in comp["rows"]]
        assert vals.count("99") == 1
        bad = vals.index("99")
        assert out_rows == {bad}
        # 색칠 엑셀 저장(변경열만) — 파일 생성 확인
        cmp_path = os.path.join(tmp, "취합비교.xlsx")
        commonality.write_comparison(cmp_path, comp, changed_only=True)
        wb = openpyxl.load_workbook(cmp_path)
        ws = wb["취합비교"]
        assert ws.cell(row=1, column=1).value == "S/M"
    print("  commonality OK: 호기 비교(행=Lot/호기) + 과반수 이탈 색칠")


def test_sm_variant_expansion_and_fail_flag():
    """S/M 'CFG' → 변형 폴더(CFG X20 / CFG #14 REWORK / CFG-RW_0517S) 모두 후보로.
    fail여부=Y 는 각 항목에 반영."""
    with tempfile.TemporaryDirectory() as tmp:
        dev6412 = Path(tmp) / "AOI-9" / "Scanresult" / "2D@X-DEVZ_0A" / "6412"
        for sm in ("CFG X20", "CFG #14 REWORK", "CFG-RW_0517S"):
            w = dev6412 / sm / "w1"
            (w / "Zones").mkdir(parents=True)
            (w / "Zones" / "Z.ini").write_text(ZONE.format(delta=25), encoding="utf-8")
            (w / "RTP.txt").write_text("x", encoding="utf-8")
        root = commonality.scanresult_root(tmp, "AOI-9")
        variants = commonality.resolve_lot_variants(root, "DEVZ", "6412", "CFG",
                                                    "AOI-9", fail=True)
        got = sorted(v.label for v in variants if v.exists)
        assert got == sorted(["CFG X20", "CFG #14 REWORK", "CFG-RW_0517S"]), got
        assert all(v.fail for v in variants)
        # resolve_plan 도 변형을 펼치고 fail 을 반영
        plan = [{"디바이스명": "DEVZ", "공정번호": "6412", "S/M": "CFG",
                 "AOI호기": "AOI-9", "fail여부": "Y"}]
        lots = commonality.resolve_plan(root, plan)
        assert len([l for l in lots if l.exists]) == 3
        assert all(l.fail for l in lots)
    print("  commonality OK: S/M 변형 다중 후보 + fail여부 반영")


def test_fail_flag_colors_result_and_comparison():
    """fail S/M 이 결과 엑셀→비교표(fail_rows)→노란색 색칠로 이어진다."""
    with tempfile.TemporaryDirectory() as tmp:
        w1 = _make_wafer(tmp, "AOI-6", "2D@DEVF_x", "6700", "AAA", "CX1", delta=25)
        w2 = _make_wafer(tmp, "AOI-6", "2D@DEVF_x", "6700", "BBB", "CX2", delta=25)
        piv, labels = commonality.parse_lots([("AAA", w1), ("BBB", w2)], level="PI3")
        form = _build_form(tmp, piv)
        res = commonality.collate_lots("PI3", form, piv, labels)
        out = os.path.join(tmp, "AOI-6.xlsx")
        commonality.write_lot_result(out, "PI3", "AOI-6", res, labels,
                                     fail_labels=["BBB"])
        data = commonality.read_lot_result(out)
        assert data["fails"] == {"BBB"}
        comp = commonality.build_comparison([out])
        bbb = next(i for i, r in enumerate(comp["rows"]) if r["S/M"] == "BBB")
        aaa = next(i for i, r in enumerate(comp["rows"]) if r["S/M"] == "AAA")
        assert bbb in comp["fail_rows"] and aaa not in comp["fail_rows"]
        cmp_path = os.path.join(tmp, "cmp.xlsx")
        commonality.write_comparison(cmp_path, comp)
        wb = openpyxl.load_workbook(cmp_path)
        ws = wb["취합비교"]
        got = False
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=1).value == "BBB":
                rgb = ws.cell(row=r, column=1).fill.fgColor.rgb or ""
                assert str(rgb).endswith(commonality.FAIL_FILL), rgb
                got = True
        assert got
    print("  commonality OK: fail S/M → 결과·비교표 노란색 색칠")


def test_zone_group_sort_adjacent():
    """비슷한 Zone(AL PAD / PAD)이 비교표 열에서 인접하게 정렬돼야 한다."""
    labels = [
        "Surface / Alg / X",
        "PAD / Surface / C",
        "LIGHT / Scan2d / Latest",
        "AL PAD / Surface / M",
    ]
    ordered = sorted(labels, key=commonality._zone_sort_key)
    # AL PAD 와 PAD 는 서로 이웃(둘 다 'pad' 그룹)
    i_alpad = ordered.index("AL PAD / Surface / M")
    i_pad = ordered.index("PAD / Surface / C")
    assert abs(i_alpad - i_pad) == 1, ordered
    print("  commonality OK: Zone 그룹 정렬(AL PAD ↔ PAD 인접)")


def test_multi_recipe_detect_and_parse():
    """다중 레시피 감지 + 레시피별 parse_lots 값 분리."""
    with tempfile.TemporaryDirectory() as tmp:
        cdir = Path(tmp) / "wafer"
        (cdir / "Zones").mkdir(parents=True)
        (cdir / "Zones" / "Z.ini").write_text(
            "[General]\nZoneName=PI Opening\n[Surface]\nHigh_Delta=11\n", encoding="utf-8")
        (cdir / "OpticPreset.ini").write_text(
            "[LIGHT]\n[Scan2d1]\nAlg=Scan2d\nLightSrc_Top=100\n", encoding="utf-8")
        (cdir / "RecipesInfo.ini").write_text(
            "[Recipe-1]\nName=PI_Bubble\n[Recipe-2]\nName=PI\n[Recipes]\nCount=2\n",
            encoding="utf-8")
        (cdir / "Recipe2-Zones").mkdir()
        (cdir / "Recipe2-Zones" / "Z.ini").write_text(
            "[General]\nZoneName=PI Opening\n[Surface]\nHigh_Delta=99\n", encoding="utf-8")
        (cdir / "Recipe2-OpticPreset.ini").write_text(
            "[LIGHT]\n[Scan2d1]\nAlg=Scan2d\nLightSrc_Top=100\n", encoding="utf-8")

        lot_dirs = [("LotA", cdir)]
        recs = commonality.detect_recipes(lot_dirs)
        assert [r["name"] for r in recs] == ["PI_Bubble", "PI"], recs
        assert [r["prefix"] for r in recs] == ["", "Recipe2-"]

        def hd(pivot):
            for row in pivot:
                if row["extract"].get("key") == "High_Delta":
                    return str(row["raws"].get("LotA"))
            return None
        p1, _ = commonality.parse_lots(lot_dirs, level="PI_Bubble",
                                       recipe_prefix=recs[0]["prefix"])
        p2, _ = commonality.parse_lots(lot_dirs, level="PI",
                                       recipe_prefix=recs[1]["prefix"])
        assert hd(p1) == "11" and hd(p2) == "99"       # 레시피별 값 분리(충돌 없음)
        # 단일 레시피면 None
        (cdir / "RecipesInfo.ini").unlink()
        assert commonality.detect_recipes(lot_dirs) is None
    print("  commonality OK: 다중 레시피 감지 + 레시피별 값 분리 파싱")


def test_wafer_slot_choices_and_switch():
    """슬롯(웨이퍼) 폴더가 여러 개면 **고를 수 있어야** 한다. 기본은 이름순 첫 번째.
    다른 슬롯을 고르면 조사 대상과 대상 파일 판정이 그 폴더로 바뀐다."""
    with tempfile.TemporaryDirectory() as tmp:
        for w, d in (("CX03", 33), ("CX01", 11), ("CX02", 22)):
            _make_wafer(tmp, "AOI-6", "2D@R2-DEVA-1_0855360PD-0A", "6321", "HPG", w,
                        delta=d)
        # 대상 파일이 없는 슬롯도 후보에는 들어간다(고르면 안내가 뜨도록)
        (Path(tmp) / "AOI-6" / "Scanresult" / "2D@R2-DEVA-1_0855360PD-0A" / "6321"
         / "HPG" / "CX99").mkdir()
        root = commonality.scanresult_root(tmp, "AOI-6")
        lot = commonality.resolve_lot(root, "DEVA-1", "6321", "HPG", "AOI-6")

        names = [p.name for p in lot.wafer_choices]
        assert names == ["CX01", "CX02", "CX03", "CX99"], names
        assert lot.wafer_dir.name == "CX01", "기본은 이름순 첫 슬롯"

        pick = next(p for p in lot.wafer_choices if p.name == "CX03")
        commonality.set_wafer(lot, pick)
        assert lot.wafer_dir.name == "CX03"
        assert lot.has_zones and lot.has_rtp and lot.has_optic
        assert lot.reason == ""
        # 고른 슬롯이 실제 조사 대상이 되는지(값이 그 슬롯 것인지) 확인
        pivot, _labels = commonality.parse_lots([(lot.label, lot.wafer_dir)],
                                                level="PI3")
        vals = [str(r["raws"].get(lot.label)) for r in pivot
                if r["extract"].get("key") == "High_Delta"]
        assert vals == ["33"], vals

        # 대상 파일이 없는 슬롯을 고르면 사유가 남는다
        commonality.set_wafer(lot, next(p for p in lot.wafer_choices
                                        if p.name == "CX99"))
        assert not (lot.has_zones or lot.has_rtp or lot.has_optic)
        assert "대상 파일" in lot.reason
    print("  commonality OK: 슬롯 폴더 선택(기본 첫 번째·전환 시 재판정)")

def test_slot_chooser_is_actually_visible():
    """슬롯 선택 콤보박스가 **화면에 보이는 순서**로 배치돼야 한다.

    tkinter pack 은 배치 순서대로 공간을 떼어 간다. expand=True 인 라벨을 먼저
    pack 하면 남은 폭을 전부 차지해, 그 뒤에 side='right' 로 놓은 콤보박스가
    밀려나 보이지 않는다 — 선택 UI 를 넣었는데도 '선택이 안 된다'던 실제 증상.
    (GUI 는 이 환경에서 띄울 수 없어 소스 배치 순서로 고정한다.)
    """
    import re
    src_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                            "param_manager", "equip_app.py")
    with open(src_path, encoding="utf-8") as fh:
        src = fh.read()
    m = re.search(r"def _cm_confirm_lots.*?(?=\n    def )", src, re.S)
    assert m, "_cm_confirm_lots 를 찾지 못함"
    body = m.group(0)
    btn_pack = body.index('btn.pack(side="right"')
    label_pack = body.index('lbl.pack(side="left", fill="x", expand=True)')
    assert btn_pack < label_pack, \
        "확장 라벨보다 슬롯 버튼을 먼저 pack 해야 화면에 보인다"
    # 고른 슬롯이 실제 조사 대상으로 펼쳐지는 경로도 함께 고정
    assert "cm.expand_all(selected)" in src, "슬롯 다중 선택이 조사 대상에 반영되지 않음"
    print("  commonality OK: 슬롯 선택 UI 배치 순서(보이는지) + 다중 선택 반영")

def test_multi_slot_expands_into_separate_targets():
    """슬롯을 여러 개 고르면 **슬롯마다 따로 조사**해야 한다(값이 다를 수 있음).
    취합 열이 겹치지 않게 라벨 뒤에 슬롯명이 붙는다."""
    with tempfile.TemporaryDirectory() as tmp:
        for w, d in (("CX01", 11), ("CX02", 22), ("CX03", 33)):
            _make_wafer(tmp, "AOI-6", "2D@R2-DEVA-1_0855360PD-0A", "6321", "HPG", w,
                        delta=d)
        root = commonality.scanresult_root(tmp, "AOI-6")
        lot = commonality.resolve_lot(root, "DEVA-1", "6321", "HPG", "AOI-6")

        # 고르지 않으면 지금까지와 동일 — 대상 1개, 라벨 그대로
        assert [l.label for l in commonality.expand_wafers(lot)] == [lot.label]

        # 1개만 고르면 라벨을 바꾸지 않는다(이전 결과와 열 이름이 이어짐)
        lot.wafer_picks = [p for p in lot.wafer_choices if p.name == "CX02"]
        one = commonality.expand_wafers(lot)
        assert len(one) == 1 and one[0].label == lot.label
        assert one[0].wafer_dir.name == "CX02"

        # 2개 이상이면 슬롯마다 별도 대상 + 라벨에 슬롯명
        lot.wafer_picks = [p for p in lot.wafer_choices if p.name in ("CX01", "CX03")]
        many = commonality.expand_wafers(lot)
        assert [l.wafer_dir.name for l in many] == ["CX01", "CX03"]
        assert [l.label for l in many] == [f"{lot.label}·CX01", f"{lot.label}·CX03"]
        assert len({l.label for l in many}) == 2, "취합 열 이름이 겹치면 안 됨"
        # 원본 lot 은 그대로(복제본이어야 한다)
        assert lot.label == "HPG"

        # 실제로 슬롯별 값이 따로 조사되는지
        pivot, labels = commonality.parse_lots(
            [(l.label, l.wafer_dir) for l in many], level="PI3")
        vals = {lab: None for lab in labels}
        for r in pivot:
            if r["extract"].get("key") == "High_Delta":
                vals = {lab: str(r["raws"].get(lab)) for lab in labels}
        assert vals == {f"{lot.label}·CX01": "11", f"{lot.label}·CX03": "33"}, vals

        # expand_all 은 목록 전체를 펼친다
        assert len(commonality.expand_all([lot, lot])) == 4
    print("  commonality OK: 슬롯 다중 선택 → 슬롯별 조사 대상 분리")

def test_loose_sm_match_and_scan_time():
    """실제 S/M 폴더명은 계획과 많이 다르다(VUL_TUNNED, SUA RERURN PG8E10 …).
    매칭은 **3단계까지만** — ①정확 일치 ②포함(양방향) ③토큰 겹침.
    그래도 못 찾으면 '폴더 없음'으로 두고 관계없는 폴더를 후보로 올리지 않는다.
    찾은 폴더에는 S/M 폴더 수정시각(Scan 일자)이 붙는다.
    """
    with tempfile.TemporaryDirectory() as tmp:
        dev, lotno = "2D@R2-DEVA-1_0855360PD-0A", "6321"
        for sm in ("VUL_TUNNED", "SUK-3D RE", "SUA RERURN PG8E10",
                   "PG8G17 NFN RETURN 2D+3D 100"):
            _make_wafer(tmp, "AOI-6", dev, lotno, sm, "CX01")
        root = commonality.scanresult_root(tmp, "AOI-6")

        # ① 포함 매칭: 'VUL' → 'VUL_TUNNED'
        got = commonality.resolve_lot_variants(root, "DEVA-1", lotno, "VUL", "AOI-6")
        assert [l.label for l in got] == ["VUL_TUNNED"], [l.label for l in got]
        assert got[0].scan_time, "S/M 폴더 수정시각(Scan 일자)이 있어야 함"

        # 토큰 겹침: 'PG8G17 RETURN' → 'PG8G17 NFN RETURN 2D+3D 100'
        got = commonality.resolve_lot_variants(root, "DEVA-1", lotno,
                                               "PG8G17 RETURN", "AOI-6")
        assert [l.label for l in got] == ["PG8G17 NFN RETURN 2D+3D 100"]

        # 3단계로도 못 찾으면 '폴더 없음' — 관계없는 폴더를 끌어오지 않는다
        got = commonality.resolve_lot_variants(root, "DEVA-1", lotno, "ZZZZ", "AOI-6")
        assert len(got) == 1 and not got[0].exists, [l.label for l in got]
        assert "S/M 폴더 없음" in got[0].reason, got[0].reason
    print("  commonality OK: S/M 매칭 3단계(정확·포함·토큰) + Scan 일자")


def test_scan_time_flows_into_comparison():
    """Scan 일자는 조사 결과 1행에 적히고, 비교표에도 열로 따라와야 한다."""
    with tempfile.TemporaryDirectory() as tmp:
        w = _make_wafer(tmp, "AOI-6", "2D@DEVE_x", "6601", "HPG", "CX1", delta=10)
        pivot, labels = commonality.parse_lots([("HPG", w)], level="PI3")
        form = _build_form(tmp, pivot)
        res = commonality.collate_lots("PI3", form, pivot, labels)
        out = os.path.join(tmp, "조사_AOI-6.xlsx")
        commonality.write_lot_result(out, "PI3", "AOI-6", res, labels,
                                     scan_times={"HPG": "2026-08-06 13:45"})
        comp = commonality.build_comparison([out])
        assert comp["columns"][:3] == ["S/M", "호기", commonality.SCAN_ROW_LABEL]
        assert comp["rows"][0][commonality.SCAN_ROW_LABEL] == "2026-08-06 13:45"
        # 파라미터 열에 Scan일자가 중복으로 끼어들면 안 된다
        assert commonality.SCAN_ROW_LABEL not in comp["columns"][3:]
        dest = os.path.join(tmp, "비교.xlsx")
        commonality.write_comparison(dest, comp)
        import openpyxl
        wb = openpyxl.load_workbook(dest)
        ws = wb["취합비교"]
        assert [c.value for c in ws[1]][:3] == ["S/M", "호기",
                                                commonality.SCAN_ROW_LABEL]
        assert ws.cell(row=2, column=3).value == "2026-08-06 13:45"
        wb.close()
    print("  commonality OK: Scan 일자 → 결과 1행 + 비교표 열")

def test_lot_slot_folder_name_is_not_used_as_variant():
    """commonality 의 config 폴더는 Lot 의 **슬롯 폴더**(CX01 …)다.
    폴더명을 변형 라벨로 쓰면 Lot 마다 변형이 달라져 값이 한 줄로 모이지 않는다
    (양식 만들기는 반대로 폴더명을 라벨로 써야 레시피가 구분된다)."""
    with tempfile.TemporaryDirectory() as tmp:
        w1 = _make_wafer(tmp, "AOI-6", "2D@DEVF_x", "6701", "HPG", "CX1", delta=11)
        w2 = _make_wafer(tmp, "AOI-6", "2D@DEVF_x", "6702", "HPG", "CX9", delta=22)
        pivot, labels = commonality.parse_lots([("L1", w1), ("L2", w2)], level="PI3")
        # 슬롯 이름이 달라도 같은 파라미터가 **한 행**에 두 Lot 값으로 모여야 한다
        hits = [r for r in pivot if r["extract"].get("key") == "High_Delta"]
        assert len(hits) == 1, [r["mag"] for r in hits]
        assert hits[0]["mag"] == "", f"슬롯명이 변형으로 샜다: {hits[0]['mag']!r}"
        assert engine._s(hits[0]["raws"].get("L1")) == "11"
        assert engine._s(hits[0]["raws"].get("L2")) == "22"
        assert set(labels) == {"L1", "L2"}
    print("  commonality OK: 슬롯 폴더명이 변형 라벨로 새지 않음")

def test_preflight_flags_missing_recipe_files():
    """접두 불일치로 **GlobalRTP 만 읽히는** 상태를 양식 만들기 전에 잡아낸다.

    실제 증상: 원본 복사는 멀쩡한데 양식 엑셀에 GlobalRTP 항목만 들어갔다.
    RecipesInfo.ini 가 레시피 2개를 선언하면 Recipe-2 는 `Recipe2-` 접두 파일로
    파싱하는데, GlobalRTP 만 공유본으로 폴백되고 OpticPreset/Zones 는 폴백이
    없어 통째로 빠지기 때문이다.
    """
    with tempfile.TemporaryDirectory() as tmp:
        d = Path(tmp) / "Lot1" / "CX01"
        (d / "Zones").mkdir(parents=True)
        (d / "GlobalRTP.ini").write_text("[GLOBAL_RTP]\nMaxFaultsPerWafer=3000\n",
                                         encoding="utf-8")
        (d / "OpticPreset.ini").write_text(
            "[Scan2d]\nCameraName=TDI\nMag=5\nLightSrcRef_NominalGL=1\n",
            encoding="utf-8")
        (d / "Zones" / "Z1.ini").write_text(
            "[General]\nZoneName=PI Opening\n[Surface]\nHigh_Delta=25\n",
            encoding="utf-8")
        (d / "RecipesInfo.ini").write_text(
            "[Recipe-1]\nName=PI\n[Recipe-2]\nName=PI_Bubble\n", encoding="utf-8")

        pre = commonality.form_preflight([("Lot1", d.parent)])
        assert [r["name"] for r in pre["recipes"]] == ["PI", "PI_Bubble"]
        f1, f2 = pre["files"]["PI"], pre["files"]["PI_Bubble"]
        assert f1["optic"] and f1["zones"] == 1 and not f1["thin"]
        # Recipe2- 접두 파일이 없다 → GlobalRTP 만 남는다(경고 대상)
        assert f2["global"] is True and f2["optic"] is False and f2["zones"] == 0
        assert f2["thin"] is True, f2

        # 접두 파일을 채우면 경고가 사라진다
        (d / "Recipe2-Zones").mkdir()
        (d / "Recipe2-Zones" / "Z1.ini").write_text(
            "[General]\nZoneName=Bub\n[Surface]\nHigh_Delta=30\n", encoding="utf-8")
        (d / "Recipe2-OpticPreset.ini").write_text(
            "[Scan2d]\nCameraName=TDI\nMag=7\nLightSrcRef_NominalGL=2\n",
            encoding="utf-8")
        pre2 = commonality.form_preflight([("Lot1", d.parent)])
        g2 = pre2["files"]["PI_Bubble"]
        assert g2["optic"] and g2["zones"] == 1 and not g2["thin"], g2

        # 단일 레시피(RecipesInfo 없음)도 파일 현황을 준다
        (d / "RecipesInfo.ini").unlink()
        pre3 = commonality.form_preflight([("Lot1", d.parent)])
        assert pre3["recipes"] == []
        assert pre3["files"][""]["optic"] and not pre3["files"][""]["thin"]
    print("  commonality OK: 접두 불일치(GlobalRTP만) 사전 감지")


def test_exact_match_does_not_hide_variants():
    """`ASD` 를 찾을 때 **`ASD` 만** 나오면 안 된다 — 변형 폴더를 다 올려야 한다.

    실제 증상: 정확일치가 하나라도 있으면 거기서 끝내서 `ASD X20`·`ASD REWORK`
    같은 변형이 통째로 숨었다. 어느 것을 조사할지는 사람이 고르는 것이므로
    후보를 전부 보여 줘야 한다.
    """
    with tempfile.TemporaryDirectory() as tmp:
        lot = Path(tmp) / "6412"
        for n in ["ASD", "ASD X20", "ASD REWORK", "ASD-RW_0517S", "BQC"]:
            (lot / n).mkdir(parents=True)
        got = [p.name for p in commonality._find_children(lot, "ASD")]
        assert got[0] == "ASD", got                 # 정확일치가 맨 앞(기본 선택)
        assert set(got) == {"ASD", "ASD X20", "ASD REWORK", "ASD-RW_0517S"}, got
        assert "BQC" not in got, got
        # 대소문자·구분자 무시
        assert [p.name for p in commonality._find_children(lot, "asd")] == got

        # 공정번호처럼 숫자는 **정확일치만**(6412 가 64120/16412 에 걸리면 안 됨)
        dev = Path(tmp) / "dev"
        for n in ["6412", "64120", "16412"]:
            (dev / n).mkdir(parents=True)
        nums = [p.name for p in commonality._find_children(dev, "6412", contains=False)]
        assert nums == ["6412"], nums

        # 정확·포함이 하나도 없을 때만 토큰 겹침(느슨) 단계로 간다
        loose = Path(tmp) / "loose"
        for n in ["SUA RERURN PG8E10", "ZZZ"]:
            (loose / n).mkdir(parents=True)
        tok = [p.name for p in commonality._find_children(loose, "SUA RETURN")]
        assert tok == ["SUA RERURN PG8E10"], tok
    print("  commonality OK: 정확일치가 변형 후보를 가리지 않음")


def test_sm_variants_all_become_lot_candidates():
    """찾은 S/M 변형이 **각각 조사 대상 후보**가 되어 선택창에 오른다."""
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        for sm in ["ASD", "ASD X20", "ASD REWORK"]:
            w = (base / "AOI-9" / "Scanresult" / "2D@R3-DEV1-0001" / "6412"
                 / sm / "CX01")
            (w / "Zones").mkdir(parents=True)
            (w / "Zones" / "Z1.ini").write_text(
                "[General]\nZoneName=Z\n[Surface]\nHigh_Delta=25\n", encoding="utf-8")
            (w / "OpticPreset.ini").write_text(
                "[Scan2d]\nCameraName=TDI\nMag=5\nLightSrcRef_NominalGL=1\n",
                encoding="utf-8")
        roots = commonality.scanresult_roots(str(base), "AOI-9")
        lots = commonality.resolve_lot_variants(roots, "DEV1-0001", "6412", "ASD",
                                                machine="AOI-9")
        labels = [l.label for l in lots if l.exists]
        assert set(labels) == {"ASD", "ASD X20", "ASD REWORK"}, labels
        assert labels[0] == "ASD", labels           # 정확일치가 첫 후보
        assert all(l.has_zones and l.has_optic for l in lots if l.exists)
    print("  commonality OK: S/M 변형이 모두 조사 후보로 오름")


if __name__ == "__main__":
    fails = 0
    tests = [(n, f) for n, f in list(globals().items())
             if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        print(f"[RUN] {name}")
        try:
            fn()
            print(f"[PASS] {name}\n")
        except Exception as e:  # noqa: BLE001
            fails += 1
            import traceback
            traceback.print_exc()
            print(f"[FAIL] {name}: {e}\n")
    print(f"==== {len(tests) - fails}/{len(tests)} passed ====")
    sys.exit(1 if fails else 0)
