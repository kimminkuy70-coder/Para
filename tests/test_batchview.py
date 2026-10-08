"""새 Batch Report 화면(가동률 조사 및 분석 · 찾기 · 취합) 엔진 테스트.

합성 자료는 test_lotmodel 과 같은 형식(실제 Batch Report 424개에서 확인한 열·Wafer ID·S/M)을 쓴다.
장비 접근 없음 — 임시 폴더를 Report 폴더로 쓴다.
"""
import io
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from param_manager import batchview, batchreport_store as store, lotmodel as lm
from param_manager.desktop_batch import DesktopBatch, VIEW_CHUNK
from param_manager.desktop_ipc import Session
from param_manager.desktop_open import DesktopOpen
from test_batchreport import html_report
from test_lotmodel import report, full, records


def lot_records():
    """BAW: 첫 스캔 S25 Error → 1시간 뒤 재스캔 Pass(재스캔으로 완료).
    NSW: 한 번에 25매 모두 Pass(정상 스캔 WPH). SPT: 같은 wafer 를 Pass 두 번(중복 Pass).
    FOCUS: 점검 스캔(Lot 제외). 자정을 넘긴 스캔(NSW 23:30~00:30)."""
    return records(
        ("AOI-1", report("BAW", full(["Scan 2D Error."]), "2026-09-01 10:00", minutes=50)),
        ("AOI-1", report("BAW", [("25", "Pass")], "2026-09-01 11:30", minutes=10)),
        ("AOI-1", report("NSW", full([]), "2026-09-01 23:30", minutes=60)),
        ("AOI-1", report("SPT", full([]), "2026-09-02 08:00", minutes=50)),
        ("AOI-1", report("SPT", [("25", "Pass")], "2026-09-02 09:00", minutes=10)),
        ("AOI-1", report("FOCUS", [("1", "Pass")], "2026-09-02 12:00", minutes=5)),
    )


class ViewNumbers(unittest.TestCase):
    def setUp(self):
        self.view = batchview.View(lot_records())
        self.p = self.view.payload()

    def lot(self, code):
        return next(i for i, lot in enumerate(self.p['lots']) if lot['code'] == code)

    def test_lots_reports_and_states(self):
        self.assertEqual(len(self.p['R']), 6)
        self.assertEqual(self.p['excluded'], [5])                    # FOCUS = 점검 스캔
        baw = self.p['lots'][self.lot('BAW')]
        self.assertEqual((baw['state'], baw['re'], baw['multi']), (lm.LOT_RESCANNED, 1, True))
        self.assertEqual(baw['cz'], [['Scan 2D Error.', 1, False]])
        self.assertEqual(self.p['lots'][self.lot('SPT')]['dup'], 1)
        r0 = self.p['R'][0]
        self.assertEqual((r0['n'], r0['ok'], r0['err'], r0['fe']), (25, 24, 1, 'Scan 2D Error.'))
        self.assertEqual(self.p['range'], ['2026-09-01', '2026-09-02'])

    def test_utilization_three_buckets_and_midnight(self):
        """시간 3칸(사용자 확정 2026-10-05): Error · 중단 없는 Batch Report 는 Batch Time 전체가 웨이퍼 처리,
        Error 가 있으면 Pass 장수 × 1장 처리 시간까지만 처리, 다시 스캔하기까지 대기는 조치(손실)."""
        U = self.p['U']
        days = {}
        for u in U:
            d = days.setdefault((u['d'], u['m']), dict(p=0, ck=0, du=0, e={}))
            d['p'] += u['p']; d['ck'] += u['ck']; d['du'] += u['du']
            for k, x in u['e'].items():
                d['e'][k] = d['e'].get(k, 0) + x[0]
        # 1장 처리 시간 = 정상 25매(NSW 60분 · SPT 첫 50분) Batch Time 합 ÷ 50장 = 132초.
        self.assertEqual(round(self.view.unit(self.view.A[0])), 132)
        # BAW 첫 스캔 50분: Pass 24장 × 132초 = 3168초 > 3000초 → 처리 3000초(남는 시간 없음).
        # 다시 스캔하기까지 10:50 → 11:30 = 2400초는 Scan 2D Error. 조치. 자정을 넘긴 NSW 는 30분씩.
        d1, d2 = days[('2026-09-01', 'AOI-1')], days[('2026-09-02', 'AOI-1')]
        self.assertEqual((d1['p'], d1['e']), (3000 + 600 + 1800, {'Scan 2D Error.': 2400}))
        self.assertEqual((d2['p'], d2['ck']), (1800 + 3000 + 600, 300))
        # SPT 두 번째 스캔 = 이미 Pass 한 wafer 를 다시 스캔해 또 Pass → 처리 시간에 넣고 '그중 다시 스캔' 으로 표시.
        self.assertEqual(d2['du'], 600)
        self.assertEqual([(w['s'], w['t']) for w in self.p['waits']], [('2026-09-01 10:50', 'e')])

    def test_wph_ingredients(self):
        U = [u for u in self.p['U'] if u['n']]
        tot = {k: sum(u[k] for u in U) for k in ('ps', 'dn', 'n', 'ne', 'w25', 's25', 'p')}
        self.assertEqual(tot, dict(ps=76, dn=1, n=5, ne=1, w25=50, s25=6600, p=10800))
        loss = sum(x[0] for u in U for x in u['e'].values())
        # 정상 WPH = 50 × 3600 ÷ 6600 ≈ 27.3, 실제 WPH = 76 × 3600 ÷ (10800 + 2400) ≈ 20.7
        self.assertAlmostEqual(tot['w25'] * 3600 / tot['s25'], 27.27, places=2)
        self.assertAlmostEqual(tot['ps'] * 3600 / (tot['p'] + loss), 20.73, places=2)
        self.assertEqual(sorted(self.p['R'][x['g']]['sm'] for x in self.p['N']), ['NSW', 'SPT'])
        baw = next(x for x in self.p['L'] if x['lot'] == self.lot('BAW'))
        self.assertEqual((baw['ps'], baw['s'], baw['n']), (25, 3000 + 600 + 2400, 2))

    def test_lot_detail_raw_and_aggregate(self):
        li = self.lot('SPT')
        detail = self.view.lot(li)
        s25 = next(w for w in detail['bunches'][0]['wafers'] if w['k'] == 'S25')
        self.assertEqual((s25['v'], s25['pick'], s25['rec'], len(s25['cells'])), (lm.DUPLICATE, 1, 1, 2))
        self.assertEqual(s25['cells'][0][1:5], ['Pass', True, False, 29])
        raw = self.view.raw(0)
        self.assertEqual(raw['h'][0], 'Lot')
        self.assertEqual(len(raw['rows']), 25)
        groups = self.view.aggregate(self.p['lots'][li]['bunches'][0]['att'] + [5])   # 점검 스캔은 빠진다
        self.assertEqual(len(groups), 1)
        w = groups[0]['w']['S25']
        self.assertEqual((len(w['cells']), w['rec'], w['saved']), (2, 1, False))
        with self.assertRaises(ValueError):
            self.view.aggregate([5])

    def test_saved_choice_changes_numbers_not_toggle(self):
        li = self.lot('SPT')
        b = self.view.model['lots'][li]['bunches'][0]
        first = b['attempts'][0]['id']
        saved = batchview.View(lot_records(), overrides={(b['key'], 'S25'): first})
        p = saved.payload()
        self.assertEqual(p['saved'], 1)
        s25 = next(w for w in saved.lot(li)['bunches'][0]['wafers'] if w['k'] == 'S25')
        self.assertEqual((s25['pick'], s25['rec'], s25['ov']), (0, 1, True))
        # 가동률 · WPH 는 어느 Pass 를 쓰든 같다(다시 스캔한 것도 처리 장수에 넣으므로) — Dice 합계만 선택을 따른다.
        self.assertEqual(p['U'], self.p['U'])
        self.assertTrue(saved.aggregate(p['lots'][li]['bunches'][0]['att'])[0]['w']['S25']['saved'])

    def test_3d_scan_is_separate_in_payload(self):
        """S/M ABC-3D = 3D 스캔: 같은 Lot 이지만 줄(bunch)이 따로라 중복 Scan 이 아니고, 화면 · WPH 이름에 3D 가 붙는다."""
        view = batchview.View(records(("AOI-1", report("ABC-3D", full([]), "2026-09-01 10:00")),
                                      ("AOI-1", report("ABC", full([]), "2026-09-01 10:35"))))
        data = view.payload()
        self.assertEqual(len(data["lots"]), 1)
        lot = data["lots"][0]
        self.assertEqual(lot["scans"], ["2D", "3D"])
        self.assertEqual(lot["dup"], 0)
        self.assertEqual([b["k"] for b in lot["bunches"]], ["3D", "2D"])
        self.assertEqual([r["k"] for r in data["R"]], ["3D", "2D"])
        self.assertEqual(sorted(x["r"] for x in data["N"]), [f"{lm.attempt(view.records[1])['job']} · 2D_WBG",
                                                             f"{lm.attempt(view.records[1])['job']} · 2D_WBG · 3D 스캔"])

    def test_day_parts(self):
        from datetime import datetime
        parts = batchview.day_parts(datetime(2026, 9, 1, 23), datetime(2026, 9, 2, 1))
        self.assertEqual(parts, [('2026-09-01', 0.5), ('2026-09-02', 0.5)])


def stop_records():
    """KVA: 정상 25매 5장(Faults 0) 뒤 — ① 10장 Pass 후 Aborted(스캔 도중 멈춘 wafer Faults 870)
    → 10분 뒤 같은 Lot 다시 스캔 ② NMB: 3장 Pass 후 Aborted(Faults 없음) ③ LDF: 정상 wafer 가 적은 레시피(2D)에서 멈춤."""
    items = [("AOI-1", report(code, full([]), f"2026-09-0{i + 1} 01:00", minutes=55)) for i, code in enumerate(("NSA", "NSB", "NSC", "NSD", "NSE"))]
    kva = report("KVA-FOCUS", full(["Pass"] * 10 + ["Aborted."] * 15), "2026-09-08 09:00", minutes=30)
    kva["wafers"][10]["Faults"] = "870"
    nmb = report("NMB", full(["Pass"] * 3 + ["Aborted."] * 22), "2026-09-08 12:00", minutes=10)
    ldf = report("LDF", full(["Pass"] * 2 + ["Aborted."] * 23), "2026-09-08 15:00", minutes=10, recipe="2D")
    ldf["wafers"][1]["Faults"] = "5000"
    items += [("AOI-1", kva), ("AOI-1", report("KVA-WBG", full([]), "2026-09-08 09:40", minutes=55)),
              ("AOI-1", nmb), ("AOI-1", ldf)]
    return records(*items)


class OperatorStop(unittest.TestCase):
    """앞 Error 없이 Aborted. 로 멈춘 Batch Report = 작업자 중단(Error 아님, 사용자 확정 2026-10-05)."""

    def setUp(self):
        self.view = batchview.View(stop_records())
        self.p = self.view.payload()

    def test_stop_is_not_error_but_time_is_lost(self):
        R = self.p['R']
        kva, nmb, ldf = R[5], R[7], R[8]
        self.assertEqual((kva['o'], kva['sk'], kva['ff'], kva['fe'], kva['st']), ('s', 'd', 870, '', 15))
        self.assertEqual((nmb['o'], nmb['sk']), ('s', 'o'))
        # 정상 wafer 2 장뿐인 레시피(2D)는 기준 부족 → Faults 5000 이어도 그 외 중단
        self.assertEqual((ldf['o'], ldf['sk']), ('s', 'o'))
        self.assertEqual(self.p['X']['faults'][R[0]['r']], [0, 150])        # 정상 25매 6장(다시 스캔한 KVA-WBG 포함)
        self.assertIsNone(self.p['X']['faults'].get(ldf['r'], [None, 0])[0])
        stops = {self.p['R'][x['g']]['sm']: (x['k'], x['f'], x['base']) for x in self.p['stops']}
        self.assertEqual(stops['KVA-FOCUS'], ('d', 870, 0))
        self.assertEqual(stops['LDF'], ('o', 5000, None))
        U = [u for u in self.p['U'] if u['d'] == '2026-09-08']
        self.assertEqual(sum(u['ne'] for u in U), 0)                     # Error 로 세지 않는다
        sd = sum(u['sd'][0] for u in U)
        # KVA 30분 − 10장 × 1장 처리 시간(55분 ÷ 25 = 132초) = 480초 + 다시 스캔하기까지 10분 = 1080초
        self.assertEqual(sd, 1800 - 10 * 132 + 600)
        self.assertEqual(sum(u['sd'][1] for u in U), 1)
        self.assertEqual(sum(u['so'][1] for u in U), 2)
        self.assertEqual([w['t'] for w in self.p['waits']], ['s'])
        lot = next(l for l in self.p['lots'] if l['code'] == 'KVA')
        self.assertEqual((lot['cz'], lot['sz']), ([], [['d', 15, False]]))
        self.assertEqual(lot['state'], lm.LOT_RESCANNED)


class DesktopFlow(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        base = Path(self.temp.name)
        self.source = base / 'equipment'
        self.source.mkdir()
        for record in lot_records():
            (self.source / record['report']['file_name']).write_text(html_report(record['report']), encoding='utf-8')
        self.config = base / 'config.json'
        self.config.write_text(json.dumps({'wph_report_paths': {'AOI-1': str(self.source)},
                                           'local_dir': str(base / 'local')}), encoding='utf-8')
        self.batch = DesktopBatch(self.config)
        self.output = io.BytesIO()
        self.session = Session(self.output)
        self.session.batch = self.batch
        self.addCleanup(self.session.close)
        self.rid = 0

    def call(self, method, **params):
        self.rid += 1
        self.session.handle(dict(version=1, id=self.rid, method=method, params=params))
        for worker in list(self.session.workers):
            worker.join(20)
        if self.session.worker:
            self.session.worker.join(20)
        events = [json.loads(line) for line in self.output.getvalue().splitlines()]
        return [e for e in events if e['id'] == self.rid][-1]

    def investigate(self, reuse=True):
        if self.session.job is not None:
            self.call('release', job=self.session.job)          # 화면도 새 조사 전에 앞 조사를 놓는다
        return self.call('investigate', targets=[{'machine': 'AOI-1', 'query': '', 'start': '', 'end': ''}],
                         options={'valid_wafers': 25, 'reuse': reuse})

    def view(self, name, meta):
        data, offset = '', 0
        while True:
            reply = self.call('batch_view', view=name, version=meta['version'], offset=offset)['batch']
            data += reply['data']
            offset += len(reply['data'])
            if offset >= reply['size']:
                return json.loads(data)

    def test_investigate_view_lot_raw_choices_restore(self):
        reply = self.investigate()
        self.assertEqual(reply['event'], 'completed')
        meta = reply['view']
        p = self.view('scope', meta)
        self.assertEqual(len(p['R']), 6)
        self.assertEqual(p['mstat']['AOI-1']['parsed'], 6)
        self.assertEqual(p['scope']['machines'], ['AOI-1'])
        self.assertTrue(p['artifacts']['lots'].endswith('BatchReport_Lot추적.html'))
        li = next(i for i, lot in enumerate(p['lots']) if lot['code'] == 'SPT')
        lot = self.call('batch_lot', view='scope', lot=li)['batch']
        self.assertEqual(len(lot['bunches'][0]['wafers']), 25)
        raw = self.call('batch_raw', view='scope', report=0)['batch']
        self.assertEqual(Path(raw['path']).parent, self.source)
        # 원본 열기: 등록한 Report 폴더 바로 아래 .htm 만 허용(장비 경로라 resolve 안 함).
        self.assertEqual(DesktopOpen(self.config).resolve(raw['path']), Path(raw['path']))
        stray = Path(self.temp.name) / 'x_BatchReport.htm'
        stray.write_text('x', encoding='utf-8')
        with self.assertRaises(ValueError):
            DesktopOpen(self.config).resolve(str(stray))
        # 개발자 기능 [선택 저장]: 앞 스캔(위치 0)을 고르면 로컬 Cache 에 남고 결과가 다시 계산된다.
        saved = self.call('batch_choices', view='scope', changes=[[li, 0, 'S25', 0]])
        self.assertEqual(saved['batch']['saved'], 1)
        meta2 = saved['batch']['views']['scope']
        self.assertGreater(meta2['version'], meta['version'])
        self.assertEqual(meta2['saved'], 1)
        choices = json.loads((Path(self.temp.name) / 'local' / 'Cache' / 'batch_lot_choices.json').read_text(encoding='utf-8'))
        self.assertEqual(len(choices['choices']), 1)
        # 옛 version 으로 조각을 달라고 하면 거절(화면이 다시 받는다).
        self.assertEqual(self.call('batch_view', view='scope', version=meta['version'], offset=0)['event'], 'error')
        # 추천으로 되돌리면 저장에서 지운다.
        self.assertEqual(self.call('batch_choices', view='scope', changes=[[li, 0, 'S25', 1]])['batch']['saved'], 0)
        # 새 엔진(앱 다시 시작): 장비 접근 없이 캐시만으로 지난 조사 결과를 다시 보여 준다.
        fresh = DesktopBatch(self.config)
        restored = fresh.restore()
        self.assertTrue(restored['restored'])
        self.assertEqual(restored['meta']['reports'], 6)
        again = json.loads(fresh._view_items()['scope']['blob'])
        self.assertEqual([r['f'] for r in again['R']], [r['f'] for r in p['R']])   # 같은 id → 같은 저장 선택 키

    def test_reuse_off_rereads(self):
        self.investigate()
        p = self.view('scope', self.investigate()['view'])
        self.assertEqual(p['mstat']['AOI-1'], {'parsed': 0, 'reused': 6, 'records': 6, 'offline': False})
        p = self.view('scope', self.investigate(reuse=False)['view'])
        self.assertEqual(p['mstat']['AOI-1']['parsed'], 6)

    def test_find_neighbors_aggregate_and_excel(self):
        # 1단계(#19): 파일 이름 목록만 — 캐시가 비어 있으면 캐시 단계는 0개, 장비 단계에서 찾는다(원문은 아직 안 읽음).
        listed = self.call('batch_find', machines=['AOI-1'], query='BAW', start='', end='', stage='cache')
        self.assertEqual(listed['event'], 'completed', listed)
        self.assertEqual(listed['batch']['total'], 0)
        listed = self.call('batch_find', machines=['AOI-1'], query='BAW', start='', end='')['batch']
        self.assertEqual(listed['total'], 2)
        self.assertEqual(listed['neighbor_h'], 24)
        self.assertFalse(any(h['c'] for h in listed['hits']))
        self.assertEqual(listed['hits'], sorted(listed['hits'], key=lambda h: h['t'], reverse=True))   # 최근 것 먼저
        self.assertNotIn('find', self.batch._view_items())                                              # 고르기 전에는 읽지 않음
        # 2단계: 고른 것 + 같은 호기에서 앞뒤 24시간 안의 Batch Report 도 읽어 Lot 으로 모은다.
        reply = self.call('batch_find_load', hits=[h['i'] for h in listed['hits']])
        self.assertEqual(reply['event'], 'completed', reply)
        p = self.view('find', reply['batch'])
        self.assertFalse(p['find']['cached_only'])
        self.assertEqual(sorted(p['R'][g]['sm'] for g in p['hits']), ['BAW', 'BAW'])
        self.assertIn('NSW', {r['sm'] for r in p['R']})             # 키워드 밖 · 이어서 스캔 후보
        key = next(k for k in p['scan'] if k.endswith('|BAW'))
        self.assertTrue(p['scan'][key]['pattern'].endswith(os.path.join('Scanresult*', lm.attempt(lot_records()[0])['job'], '6392', 'BAW')))
        groups = self.call('batch_aggregate', view='find', reports=p['hits'])['batch']['groups']
        self.assertEqual(len(groups), 1)
        out = self.call('batch_export', view='find', kind='agg', reports=p['hits'], choices={}, stamp='최신 스캔 자동')
        path = Path(out['batch']['path'])
        self.assertTrue(path.is_file())
        from openpyxl import load_workbook
        wb = load_workbook(path)
        self.assertEqual(wb['Lot 취합']['A1'].value, 'Batch Report 찾기 · 취합 · 선택 기준: 최신 스캔 자동')
        self.assertEqual(wb['Lot 취합'].max_row, 3)
        self.assertEqual(wb['wafer'].max_row, 27)
        li = p['R'][p['hits'][0]]['lot']
        out = self.call('batch_export', view='find', kind='lot', lot=li, drafts={}, stamp='추천')
        self.assertTrue(Path(out['batch']['path']).name.startswith('Lot_BAW_취합_'))
        # 이제 캐시에 있으므로 캐시 단계만으로 찾고(장비 접근 없음), 읽기도 캐시만으로 끝난다(C 안).
        cached = self.call('batch_find', machines=['AOI-1'], query='BAW', start='', end='', stage='cache')['batch']
        self.assertEqual((cached['total'], cached['cached'], cached['listed']), (2, 2, 0))
        again = self.view('find', self.call('batch_find_load', hits=[cached['hits'][0]['i']])['batch'])
        self.assertTrue(again['find']['cached_only'])
        self.assertEqual(len(again['hits']), 1)
        # 잘못된 고르기 · 단계는 거절.
        self.assertEqual(self.call('batch_find_load', hits=[99])['event'], 'error')
        self.assertEqual(self.call('batch_find_load', hits=[])['event'], 'error')
        self.assertEqual(self.call('batch_find', machines=['AOI-1'], query='BAW', start='', end='', stage='x')['event'], 'error')
        # 빈 키워드 · 등록 안 된 호기는 거절.
        self.assertEqual(self.call('batch_find', machines=['AOI-1'], query='', start='', end='')['event'], 'error')
        self.assertEqual(self.call('batch_find', machines=['AOI-9'], query='X', start='', end='')['event'], 'error')

    def test_chunks_reassemble(self):
        meta = self.investigate()['view']
        self.assertLess(meta['size'], VIEW_CHUNK)
        original = VIEW_CHUNK
        import param_manager.desktop_batch as db
        db.VIEW_CHUNK = 100
        try:
            p = self.view('scope', meta)
        finally:
            db.VIEW_CHUNK = original
        self.assertEqual(len(p['R']), 6)

    def test_cache_loader_reads_only_local(self):
        self.investigate()
        target = {'machine': 'AOI-1', 'folder': str(self.source), 'query': '', 'start': '', 'end': ''}
        renamed = self.source.with_name('gone')
        self.source.rename(renamed)                  # 장비 연결이 끊겨도 캐시만으로 읽는다
        out = store.load_cached(Path(self.temp.name) / 'local', [target])
        self.assertEqual(len(out['records']), 6)


if __name__ == '__main__':
    unittest.main()
