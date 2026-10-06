"""저장된 결과 파일(분석 HTML · Excel · 가동률 대시보드 · Lot 추적 HTML) — 앱 화면(v12)과 같은 계산(이슈 #7).

합성 자료는 test_batchview 와 같은 것을 쓴다. 숫자는 payload U 를 앱(batchData.ts)과 같은 식으로 더한 값과 비교한다.
"""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from param_manager import batchsaved as bs, batchview, lotreport
from test_batchview import lot_records


class SavedNumbers(unittest.TestCase):
    def setUp(self):
        self.view = batchview.View(lot_records())
        self.p = self.view.payload()
        self.s = bs.Saved(self.p, machines=['AOI-1'], scope='테스트')

    def test_util_matches_app_formula(self):
        U = self.p['U']
        proc = sum(u['p'] + u['ck'] for u in U)
        loss = sum(sum(x[0] for x in u['e'].values()) + u['sd'][0] + u['so'][0] for u in U)
        c = self.s.util_all()
        self.assertEqual(c['cal'], 2 * 86400)                      # 2일 × 1대 × 24시간
        self.assertAlmostEqual(c['proc'], proc)
        self.assertAlmostEqual(c['loss'], loss)
        self.assertAlmostEqual(c['idle'], c['cal'] - proc - loss)
        days = self.s.util_periods('d')
        self.assertEqual([k for k, _ in days], ['2026-09-02', '2026-09-01'])
        self.assertAlmostEqual(sum(x['proc'] for _, x in days), proc)

    def test_wph_matches_app_formula(self):
        recipes, cells = self.s.wph_cells()
        self.assertTrue(recipes)
        rows = [u for u in self.p['U'] if u['n'] or u['p']]
        w25, s25 = sum(u['w25'] for u in rows), sum(u['s25'] for u in rows)
        a = bs.sum_u(rows)
        self.assertAlmostEqual(bs.wph_normal(a), w25 * 3600 / s25)
        ps = sum(u['ps'] for u in rows)
        self.assertAlmostEqual(bs.wph_actual(a), ps * 3600 / (a['p'] + bs.loss_of(a)))
        self.assertAlmostEqual(bs.cap(a), 24 * bs.wph_actual(a))

    def test_lot_summary_matches_app(self):
        k = {l: n for l, n, *_ in self.s.lot_kpis()}
        self.assertEqual(k['조사 범위 내 Lot Scan 완료'], len(self.p['lots']))
        self.assertEqual(k['재스캔으로 완료된 Lot'], 1)
        self.assertEqual(k['중복 Scan한 wafer가 존재하는 Lot'], 1)
        self.assertEqual(self.s.causes(), [['Scan 2D Error.', 1, 1, 0, 1]])

    def test_interactive_chart_data_and_scripts(self):
        # 이슈 #13: 저장된 HTML 도 앱과 같은 인터랙티브 그래프 — 앱과 같은 U 행 · Batch Report · 조치 대기를 내장하고 스크립트가 그린다.
        import json, re
        page = bs.build_html(self.s, lot_file='BatchReport_Lot추적.html')
        for anchor in ('id="lotchart"', 'id="utilapp"', 'id="wphapp"', 'var BV=', 'function colChart', 'recipeWindow', 'periodWindow'):
            self.assertIn(anchor, page)
        data = json.loads(re.search(r'<script type="application/json" id="bvdata">(.*?)</script>', page).group(1))
        self.assertEqual(data['U'], self.s.U)                       # 숫자는 앱이 받는 것과 같은 U 행에서 계산
        self.assertEqual(data['ids'], ['AOI-1'])
        self.assertEqual(data['lotHref'], 'BatchReport_Lot추적.html')
        self.assertEqual(len(data['R']), len(self.p['R']))
        self.assertEqual([l[0] for l in data['lots']], [l['label'] for l in self.p['lots']])
        self.assertNotIn('</script', re.search(r'id="bvdata">(.*?)</script>', page).group(1))
        dash = bs.build_html(self.s, dashboard=True)
        self.assertEqual(json.loads(re.search(r'id="bvdata">(.*?)</script>', dash).group(1))['lotHref'], '')
        lots = lotreport.build_html(self.view.model, '테스트', view=self.view)
        for anchor in ('id="lotchart"', '__LOTPERIOD', 'var BV=', '#lot='):
            self.assertIn(anchor, lots)

    def test_machine_colors_stable_and_same_as_app_table(self):
        c1 = bs.mach_colors(['AOI-2', 'AOI-1', 'AOI-10'])
        c2 = bs.mach_colors(['AOI-10', 'AOI-2', 'AOI-1', 'AOI-1'])
        self.assertEqual(c1, c2)                                   # 순서 · 중복과 무관
        self.assertEqual(c1['AOI-1'], bs.MACH_PAL[0])
        src = (Path(__file__).resolve().parents[1] / 'frontend' / 'src' / 'batchData.ts').read_text(encoding='utf-8')
        for color in bs.MACH_PAL:                                  # 앱과 같은 색 표
            self.assertIn(f"'{color}'", src)

    def test_html_excel_dashboard_and_no_wph_row(self):
        page = bs.build_html(self.s, lot_file='BatchReport_Lot추적.html')
        for text in ('Lot 추적', '가동률', 'WPH · 생산능력', '지표 정의', '멈춘 이유 요약 · Error', '같은 호기 = 같은 색'):
            self.assertIn(text, page)
        # 이슈 #12: 저장된 HTML 의 WPH 탭에도 앱과 같은 상위 레시피 필터(칩 · 전체 보기) + 숫자 4개를 다시 더할 자료.
        import json, re
        self.assertIn('id="jf"', page)
        self.assertIn('class="jf-chip"', page)
        blob = json.loads(re.search(r'<script type="application/json" id="wphjobs">(.*?)</script>', page).group(1))
        recipes, _ = self.s.wph_cells()
        self.assertEqual(blob['jobs'], sorted({bs._job(r) for r in recipes}))
        self.assertEqual(sum(x['ps'] for x in blob['sum'].values()), bs.sum_u(self.s.wph_rows())['ps'])
        for j in blob['jobs']:
            self.assertIn(f'data-job="{j}"', page)
        dash = bs.build_html(self.s, dashboard=True)
        self.assertIn('content="1800"', dash)
        self.assertNotIn('id="sec-lot"', dash)
        # 정상 25매가 없는 레시피도 '정상 WPH 없음' 행으로 남긴다(사용자 요청 2026-10-05).
        recipes, cells = self.s.wph_cells()
        no_normal = [r for r in recipes if bs.wph_normal(bs.sum_u([u for u in self.p['U'] if u['r'] == r])) is None]
        if no_normal:
            self.assertIn('정상 WPH 없음', page)
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'a.xlsx'
            bs.write_excel(path, self.s, ['알림 1'])
            from openpyxl import load_workbook
            wb = load_workbook(path)
            for name in ('조사 정보 · 정의', '가동률 호기별', '가동률 일별', 'WPH 호기×레시피', 'Lot 목록', 'Batch Report 목록', '작업자 중단'):
                self.assertIn(name, wb.sheetnames)
            self.assertEqual(wb['Lot 목록'].max_row, 2 + len(self.p['lots']))
            wb.close()

    def test_recipe_without_normal_wph_gets_row(self):
        import copy
        p = copy.deepcopy(self.p)
        for u in p['U']:
            u['w25'], u['s25'] = 0, 0                              # 정상 25매 Batch Report 가 하나도 없는 경우
        page = bs.wph_section(bs.Saved(p, machines=['AOI-1']))
        self.assertIn('정상 WPH 없음', page)

    def test_lot_html_uses_view_model(self):
        page = lotreport.build_html(self.view.model, '범위', view=self.view)
        self.assertIn('멈춘 이유 요약 · Error', page)
        self.assertIn('작업자 중단', page)
        self.assertIn('재스캔으로 완료된 Lot', page)
        old = lotreport.build_html(self.view.model, '범위')      # view 없이도 그대로 동작(하위호환)
        self.assertIn('Batch Report Lot 추적', old)


if __name__ == '__main__':
    unittest.main()
