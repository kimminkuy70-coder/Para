"""Headless tests for the desktop form-building adapter (양식 만들기 웹).

Covers catalog → open → page → edit(use/name/transform) → confirm over an
existing candidate under a temp save_dir. No equipment / tkinter / React.
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from param_manager import formbuilder, workdirs
from param_manager.desktop_form import DesktopForm


def _pivot(level, variant, zone, alg, param, key, raw, transform="RAW"):
    return dict(layer=level, recipe=level, mag=variant, zone=zone, alg=alg,
                param=param, values={}, unit="",
                raws={"양식": raw},
                use=True,
                extract={"src_file": "OpticPreset.ini", "section": "Scan2d",
                         "key": key, "transform": transform, "source_path": ""})


class DesktopFormTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="rev1-form-")
        self.save = os.path.join(self.tmp, "저장폴더")
        os.makedirs(self.save)
        self.recipe = "PI3"
        # Build a candidate (원본) with three params under one run version.
        st = "20260921_010101_000001"
        run = workdirs.form_run_dir(self.save, self.recipe, st)
        related = workdirs.related_dir(run)
        # A final workbook must exist directly in the run dir for the version to list.
        openpyxl.Workbook().save(workdirs.form_final_path(run, self.recipe, "AOI-01", st))
        rows = [_pivot("PI3", "PI", "Zone1", "Scan2d", "Min Defect Width", "MinWidth", "12"),
                _pivot("PI3", "PI", "Zone1", "Scan2d", "Max Defects", "MaxDef", "999"),
                _pivot("PI3", "PI-bubble", "Zone2", "Scan2d", "Sensitivity µm", "Sens", "3.5", "LINEAR")]
        formbuilder.build_initial_workbook(
            rows, workdirs.form_original_path(related, self.recipe, "AOI-01", st),
            level="PI3", source="테스트")
        # config with save_dir
        self.cfg = os.path.join(self.tmp, "config.json")
        with open(self.cfg, "w", encoding="utf-8") as fh:
            json.dump({"save_dir": self.save}, fh)
        self.form = DesktopForm(config_path=self.cfg)

    def test_catalog_lists_editable_recipe(self):
        cat = self.form.catalog()
        self.assertTrue(cat["save_dir"])
        recipes = {r["recipe"]: r for r in cat["recipes"]}
        self.assertIn("PI3", recipes)
        self.assertTrue(recipes["PI3"]["can_edit"])
        self.assertTrue(recipes["PI3"]["versions"][0]["has_candidate"])

    def test_open_page_edit_confirm(self):
        self.form.catalog()
        opened = self.form.open({"recipe": "PI3", "stamp": ""})
        self.assertEqual(opened["level"], "PI3")
        self.assertEqual(opened["total"], 3)
        self.assertEqual(opened["used"], 3)
        self.assertIn("PI", opened["variants"])
        self.assertIn("PI-bubble", opened["variants"])
        snap = opened["version"]

        page = self.form.page({"snapshot": snap, "variant": "", "query": "", "used_only": False,
                               "offset": 0, "limit": 100})
        self.assertEqual(page["total"], 3)
        by_orig = {r["orig"]: r for r in page["rows"]}
        self.assertIn("MinWidth", by_orig)

        # Uncheck one, rename another, change a transform.
        self.form.edit({"snapshot": snap, "row": by_orig["MaxDef"]["id"], "kind": "use", "value": False})
        self.form.edit({"snapshot": snap, "row": by_orig["MinWidth"]["id"], "kind": "name", "value": "최소 결함 폭"})
        self.form.edit({"snapshot": snap, "row": by_orig["MinWidth"]["id"], "kind": "transform", "value": "LINEAR"})

        result = self.form.confirm({"snapshot": snap, "machine": "AOI-07"})
        self.assertEqual(result["kept"], 2)          # MaxDef unchecked → dropped
        self.assertTrue(os.path.exists(result["final"]))
        self.assertTrue(os.path.exists(result["original"]))
        # Confirmed workbook holds the renamed param and not the unchecked one.
        wb = openpyxl.load_workbook(result["final"])
        ws = wb[result["sheet"]]
        heads = [c.value for c in ws[1]]
        pcol = heads.index("Parameter")
        params = {ws.cell(r, pcol + 1).value for r in range(2, ws.max_row + 1)}
        self.assertIn("최소 결함 폭", params)
        self.assertNotIn("Max Defects", params)
        wb.close()

    def test_scales_coef_store_and_value_inheritance(self):
        from param_manager import coefstore, collate, refdata
        # Registered machines come from the shared 장비 IP workbook.
        ip = refdata.ip_path(self.save)
        refdata.create_blank_ip(ip)
        refdata.save_ip(ip, [{"호기": "AOI-07", "IP": "10.0.0.7", "장비종류": "Camtek"}])
        coefstore.save(coefstore.coef_path(self.save),
                       [{"호기": "AOI-07", "MAG": "x5", "변형": "PI-bubble", "계수": "0.5", "비고": ""}])
        # A previous collation holds a value that the edited form must inherit.
        prev = workdirs.collate_path(self.save, "20260920_000000")
        rec = {"PI": "PI3", "Recipe": "PI-bubble", "Zone": "Zone2", "Alg": "Scan2d",
               "Parameter": "Sensitivity µm", "비고": "", "AOI-07": "7.7"}
        collate.write_collation(prev, {"PI3": collate.CollateRecipe(recipe="PI3", records=[rec])}, ["AOI-07"])

        catalog = self.form.catalog()
        self.assertEqual(catalog["machines"], ["AOI-07"])
        snap = self.form.open({"recipe": "PI3", "stamp": ""})["version"]
        scales = {s["variant"]: s for s in self.form.scales({"snapshot": snap, "machine": "AOI-07"})["scales"]}
        self.assertEqual((scales["PI-bubble"]["coef"], scales["PI-bubble"]["source"]), (0.5, "변환계수.xlsx"))
        self.assertTrue(scales["PI-bubble"]["needed"])       # LINEAR item in use
        self.assertFalse(scales["PI"]["needed"])             # only RAW items
        for bad in ({"PI-bubble": 0}, {"PI-bubble": float("inf")}, {"PI-bubble": "1"}):
            with self.assertRaises(ValueError):
                self.form.confirm({"snapshot": snap, "machine": "AOI-07", "scales": bad})

        result = self.form.confirm({"snapshot": snap, "machine": "AOI-07", "scales": {"PI-bubble": 1.25}})
        self.assertEqual(result["coef"]["saved"], 1)
        self.assertEqual(coefstore.lookup(coefstore.load(coefstore.coef_path(self.save)), "AOI-07", "PI-bubble"), 1.25)
        merged = result["merge"]
        self.assertTrue(merged["collate"] and os.path.exists(merged["collate"]), merged)
        sheets, machines = collate.load_collation(merged["collate"])
        self.assertEqual(machines, ["AOI-07"])
        values = {r["Parameter"]: r.get("AOI-07") for rows in sheets.values() for r in rows}
        self.assertEqual(values.get("Sensitivity µm"), "7.7")

    def test_used_only_and_query_filter(self):
        self.form.catalog()
        opened = self.form.open({"recipe": "PI3", "stamp": ""})
        snap = opened["version"]
        self.form.edit({"snapshot": snap, "row": 0, "kind": "use", "value": False})
        used = self.form.page({"snapshot": snap, "variant": "", "query": "", "used_only": True,
                               "offset": 0, "limit": 100})
        self.assertEqual(used["total"], 2)
        q = self.form.page({"snapshot": snap, "variant": "", "query": "sens", "used_only": False,
                            "offset": 0, "limit": 100})
        self.assertEqual(q["total"], 1)

    def test_stale_snapshot_rejected(self):
        self.form.catalog()
        self.form.open({"recipe": "PI3", "stamp": ""})
        with self.assertRaises(ValueError):
            self.form.page({"snapshot": "not-the-snapshot", "variant": "", "query": "",
                            "used_only": False, "offset": 0, "limit": 100})

    def test_no_save_dir(self):
        empty = os.path.join(self.tmp, "empty.json")
        with open(empty, "w", encoding="utf-8") as fh:
            json.dump({}, fh)
        form = DesktopForm(config_path=empty)
        self.assertEqual(form.catalog(), {"save_dir": False, "recipes": []})


if __name__ == "__main__":
    unittest.main()
