"""Synthetic regression tests; not the unavailable 419-report production sample."""
import copy
import os
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from param_manager import batchreport as br, batchreport_store as store
from param_manager import batchreport_output as output, batchreport_service as service, wph


def report(name='a.htm', statuses=('Pass',), start='19-Sep-26 10:00:00 AM', end='19-Sep-26 11:00:00 AM', job='JOB/1', count=None, seconds='01:00:00', bad='2', yield_pct='99'):
    return {'file_name': name, 'metadata': [('Batch Start', start), ('Batch End', end), ('Batch Time', seconds),
            ('Wafers Scanned', str(len(statuses) if count is None else count)), ('Job/Setup', job)],
            'wafers': [{'No': str(i), 'Wafer ID': str(i), 'Lot': 'L1', 'Faults': '0', 'Scanned Dice': '100',
                        'Bad Dice': bad, 'Good Dice': '98', 'Yield': yield_pct, 'Pass/Fail': s, 'Recipe(s)': 'R1'}
                       for i, s in enumerate(statuses, 1)], 'table_count': 2}


def records(*reports):
    return [{'id': str(i), 'machine': 'AOI-1', 'report': r} for i, r in enumerate(reports)]


def table(result, title):
    return next(t['rows'] for t in result['tables'] if t['title'] == title)


def html_report(rep):
    from html import escape
    headers = list(rep['wafers'][0])
    return '<html><table>' + ''.join(f'<tr><td>{escape(k)}</td><td>{escape(v)}</td></tr>' for k, v in rep['metadata']) + '</table><table><tr>' + ''.join('<th>' + k + '</th>' for k in headers) + '</tr>' + ''.join('<tr>' + ''.join('<td>' + escape(w[k]) + '</td>' for k in headers) + '</tr>' for w in rep['wafers']) + '</table></html>'


class Metrics(unittest.TestCase):
    def test_classification_and_unknown(self):
        self.assertEqual(br.classify('Pass'), [])
        states = br.classify('Scan 2D Error. + Aborted. + Skipped. + Novel laser failure')
        self.assertEqual({s[0] for s in states}, {'2D Scan 오류', '작업 중단', '검사 제외', 'Unclassified'})
        self.assertEqual(br.classify('Clean Reference Error')[0][2], '조치형')
        self.assertEqual(br.classify('')[0][0], '상태 확인 불가')

    def test_abort_frequency_completion_and_wph_scope(self):
        r = report(statuses=('Scan 2D Error.', 'Aborted.', 'Aborted.', 'Skipped.'), count=1)
        s = report('b.htm', statuses=('Aborted.', 'Aborted.'), count=0)
        c = br.compute(records(r, s), valid_wafers=25)
        # Aborted 보정: (구분, Wafer 수, Lot 수). r·s 는 같은 job·Lot 이라 1 lot.
        self.assertEqual(table(c, 'Aborted 보정'), [('원본 Aborted', 4, 1), ('직접 추정', 1, 1), ('연쇄 추정', 3, 1)])
        self.assertEqual(table(c, 'WPH'), [])
        self.assertEqual(c['summary']['이슈 Batch 수'], 2)
        self.assertEqual(c['batches'][0]['completion'], 25)
        # 유형별 빈도 전체 행: (…, Wafer 발생, Report 수, Lot 수)
        abort = next(r for r in table(c, '유형별 빈도') if r[0] == '전체' and r[3] == '작업 중단')
        self.assertEqual(abort[5:8], (4, 2, 1))
        c = br.compute(records(report(count=10)), selected=['M10'])
        self.assertIsNone(c['batches'][0]['completion'])
        self.assertTrue(all(t['key'] == 'M10' for t in c['tables']))

    def test_lot_grouping_issue_dice_and_recipe_is_job_setup(self):
        a = report('a.htm', statuses=('Pass', 'Pass', 'Pass', 'Aborted.'), job='JOB/1')
        b = report('b.htm', statuses=('Pass',), job='JOB/1',
                   start='19-Sep-26 12:00:00 PM', end='19-Sep-26 12:30:00 PM')  # 같은 lot 재스캔
        d = report('d.htm', statuses=('Pass', 'Pass'), job='JOB/2')             # 정상 lot
        c = br.compute(records(a, b, d), valid_wafers=25)
        self.assertEqual(c['summary']['Lot 수'], 2)
        self.assertEqual(c['summary']['이슈 발생 Lot 수'], 1)
        self.assertEqual(c['summary']['재스캔 Lot 수'], 1)
        issue = dict((r[0], r[1]) for r in table(c, 'Lot 이슈 요약'))
        self.assertEqual((issue['이슈 발생 Lot'], issue['정상 Lot'], issue['재스캔(리포트≥2) Lot']), (1, 1, 1))
        detail = table(c, '문제 Lot 상세')
        self.assertEqual(len(detail), 1)
        self.assertEqual((detail[0][0], detail[0][2]), ('JOB/1', 2))  # 2 스캔 시도
        # M08 은 Recipe(=Job/Setup)별. Good = Scanned - Bad (원문 있으면 원문).
        dice = {r[0]: r for r in table(c, 'Recipe별 정상 Dice 통계')}
        self.assertIn('JOB/1', dice)
        self.assertIn('JOB/2', dice)
        self.assertEqual(dice['JOB/2'][1], 2)          # 표본 수 = Pass wafer 수
        self.assertTrue(br.is_lot_placeholder('LoadPort A'))
        self.assertFalse(br.is_lot_placeholder('KVY'))

    def test_wph_weighted_vs_arithmetic(self):
        c = br.compute(records(report(count=25), report('b.htm', count=25, seconds='00:30:00')), valid_wafers=25)
        total = next(r for r in table(c, 'WPH') if r[0] == '전체')
        self.assertAlmostEqual(total[3], 100 / 3)
        self.assertEqual(total[4], 37.5)

    def test_calendar_and_partial_week(self):
        r = report(start='31-Aug-26 11:30:00 PM', end='01-Sep-26 12:30:00 AM')
        c = br.compute(records(r), now=datetime(2026, 9, 2))
        # 자정 넘긴 1h 배치 → 08-31 에 0.5h, 09-01 에 0.5h 배분(139% 버그 방지).
        days = table(c, '스캔 가동률 · 일')
        day = days[0]
        self.assertEqual(day[1], datetime(2026, 8, 31))
        self.assertAlmostEqual(day[3], 0.5)             # 스캔 시간(h)
        self.assertAlmostEqual(day[5], 0.5 / 24 * 100)  # 가동률(%)
        self.assertAlmostEqual(days[1][3], 0.5)         # 09-01 에도 0.5h
        week = table(c, '스캔 가동률 · 주')[0]
        self.assertEqual(week[4], 48)                   # 기간 길이(h) = 시작~현재
        self.assertEqual(week[-1], '진행 중')
        month = table(c, '스캔 가동률 · 월')[0]
        self.assertEqual(month[4], 31 * 24)

    def test_gaps_use_same_full_job_and_exclusions(self):
        c = br.compute(records(
            report('a.htm', statuses=('Alignment Error.',)),
            report('different.htm', start='19-Sep-26 11:01:00 AM', end='19-Sep-26 11:02:00 AM', job='JOB/2'),
            report('next.htm', start='19-Sep-26 11:10:00 AM', end='19-Sep-26 12:00:00 PM'),
            report('orphan.htm', statuses=('Aborted.',), start='19-Sep-26 01:00:00 PM', end='19-Sep-26 02:00:00 PM')))
        restart = table(c, '재시작 간격(같은 lot 재스캔)')
        self.assertEqual(restart[0][-3:], ('next.htm', 10, '유효'))
        self.assertIn('없음', restart[1][-1])
        # M07 폐지 → 전체 요약이 M06 유형별 표의 '(전체 유효)' 마지막 행으로.
        self.assertEqual(table(c, '유형별 재시작 간격 (유형 중복 허용)')[-1], ('(전체 유효)', 1, 10, 10, 10, 10))

    def test_p95_and_past_only_quality_baseline(self):
        early = report('early.htm', statuses=('Pass', 'Pass'), bad='1')
        later = report('later.htm', start='19-Sep-26 12:00:00 PM', end='19-Sep-26 01:00:00 PM', bad='9', yield_pct='80')
        c = br.compute(records(early, later), min_baseline=2, yield_drop=5)
        self.assertAlmostEqual(br.percentile([1, 2, 3, 4], .95), 3.85)
        anomalies = table(c, '품질 이상 Wafer 상세 (자동 Hold 아님)')
        self.assertEqual(len(anomalies), 1)
        # (Lot, Job/Setup, Batch End, 호기, Report, Wafer ID, Bad, P95, Yield, 하한, 근거)
        self.assertEqual(anomalies[0][4], 'later.htm')
        self.assertEqual(anomalies[0][7], 1)
        self.assertEqual(anomalies[0][9], 94)
        self.assertEqual(len(table(br.compute(records(later), min_baseline=2), '품질 이상 Wafer 상세 (자동 Hold 아님)')), 0)
        self.assertEqual(len(table(c, '품질 이상 Lot 요약')), 1)

    def test_missing_and_nonfinite(self):
        self.assertIsNone(br.number('nan'))
        self.assertIsNone(br.number('-1'))
        self.assertIsNone(br.number('inf'))
        with self.assertRaises(ValueError):
            br.compute([], selected=[])
        c = br.compute(records(report(start='', end='')))
        self.assertEqual(table(c, '스캔 가동률 · 일'), [])

    def test_same_time_is_not_past_and_percent_over_100_is_visible(self):
        a = report('a.htm', statuses=('Pass', 'Pass'), bad='1')
        b = report('b.htm', bad='99')
        c = br.compute(records(a, b), min_baseline=2)
        self.assertEqual(table(c, '품질 이상 Wafer 상세 (자동 Hold 아님)'), [])
        c = br.compute(records(report(seconds='30:00:00')), now=datetime(2026, 9, 21))
        day = table(c, '스캔 가동률 · 일')[0]
        self.assertEqual(day[5], 125)       # 가동률(%): 스캔시간(30h) > 기간(24h)
        self.assertLess(day[6], 0)          # 비스캔 시간(h) 음수
        self.assertIn('100%', day[-1])      # 기간 상태에 100% 초과 표시


class Collection(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        base = Path(self.temp.name)
        self.root, self.source = base / 'local', base / 'source'
        self.source.mkdir()
        self.target = {'machine': 'AOI-1', 'folder': str(self.source), 'query': '', 'names': None}

    def tearDown(self):
        self.temp.cleanup()

    def put(self, name, rep=None):
        (self.source / name).write_text(html_report(rep or report(name)), encoding='utf-8')

    def test_once_late_file_changed_and_selection(self):
        self.put('new.htm')
        with patch.object(wph, 'parse_report', wraps=wph.parse_report) as parse:
            a = store.collect(self.root, [self.target], host_gap=0)
            b = store.collect(self.root, [self.target], host_gap=0)
            self.assertEqual(parse.call_count, 1)
            self.assertEqual((a['parsed'], b['reused']), (1, 1))
            self.put('old_no_date.htm', report('old_no_date.htm', job='OLD/1'))
            self.assertEqual(store.collect(self.root, [self.target], host_gap=0)['parsed'], 1)
            self.put('new.htm', report('new.htm', count=2))
            c = store.collect(self.root, [self.target], host_gap=0)
            self.assertEqual(c['parsed'], 1)
            self.assertEqual(len(c['records']), 2)
        only = dict(self.target, query='old')
        self.assertEqual(len(store.collect(self.root, [only], host_gap=0)['records']), 1)
        none = dict(self.target, names=[])
        self.assertEqual(store.collect(self.root, [none], host_gap=0)['records'], [])
        (self.source / 'old_no_date.htm').unlink()
        self.assertEqual(store.collect(self.root, [only], host_gap=0)['cached_only'], 1)

    def test_fail_retry_backup_dedup_and_same_name_conflict(self):
        path = self.source / 'bad.htm'
        path.write_text('<html>incomplete', encoding='utf-8')
        a = store.collect(self.root, [self.target], host_gap=0)
        self.assertEqual((len(a['records']), len(a['errors'])), (0, 1))
        self.put('bad.htm')
        backup = self.source.parent / 'backup'
        backup.mkdir()
        (backup / 'renamed.htm').write_text(path.read_text(), encoding='utf-8')
        targets = [self.target, dict(self.target, folder=str(backup))]
        a = store.collect(self.root, targets, host_gap=0)
        self.assertEqual(len(a['records']), 1)
        self.assertEqual(len(a['notices']), 1)
        (backup / 'bad.htm').write_text(html_report(report('bad.htm', job='OTHER/2')), encoding='utf-8')
        self.assertEqual(len(store.collect(self.root, targets, host_gap=0)['records']), 2)

    def test_corrupt_cache_and_path_guards(self):
        self.put('a.htm')
        store.collect(self.root, [self.target], host_gap=0)
        cache = next((self.root / '배치분석' / '누적').glob('*.json'))
        cache.write_text('broken', encoding='utf-8')
        with self.assertRaises(ValueError):
            store.collect(self.root, [self.target], host_gap=0)
        self.assertEqual(cache.read_text(), 'broken')
        for root in [self.source, self.source / 'output', self.source.parent / 'OneDrive' / 'output', '//server/share']:
            with self.assertRaises(ValueError):
                store.local_root(root, [self.source])
        with self.assertRaises(ValueError):
            store.collect(self.root.parent / 'clean', [dict(self.target, names=['../bad.htm'])], host_gap=0)

    def test_changed_during_read_not_seen(self):
        self.put('a.htm')
        original = wph.parse_report

        def racing(path):
            rep = original(path)
            path.write_text(path.read_text() + 'changed', encoding='utf-8')
            return rep

        with patch.object(wph, 'parse_report', side_effect=racing):
            result = store.collect(self.root, [self.target], host_gap=0)
        self.assertEqual(len(result['records']), 0)
        self.assertIn('변경', result['errors'][0]['error'])
        self.assertEqual(store.collect(self.root, [self.target], host_gap=0)['parsed'], 1)

    def test_service_outputs_charts_escaping_and_lock(self):
        self.put('a.htm', report('a.htm', statuses=('Unknown <script>alert(1)</script>', 'Aborted.'), job='=1+1'))
        options = {'metrics': list(br.METRICS), 'valid_wafers': 2, 'min_baseline': 2, 'yield_drop': 5}
        result = service.run(self.root, [self.target], options, host_gap=0)
        page = Path(result['html']).read_text()
        self.assertNotIn('<script>alert(1)</script>', page)
        self.assertIn('&lt;script&gt;', page)
        self.assertIn('2026-09-19', page)
        self.assertIn('data-period="주"', page)
        self.assertIn('content="1800"', Path(result['dashboard']).read_text())
        from openpyxl import load_workbook
        wb = load_workbook(result['xlsx'])
        self.assertGreater(len(wb['그래프']._charts), 0)
        self.assertEqual(wb['Batches']['H2'].data_type, 's')
        self.assertEqual(wb['Batches']['H2'].value, '=1+1')
        self.assertEqual(wb['Wafers'].max_row, 3)
        wb.close()
        with zipfile.ZipFile(result['xlsx']) as archive:
            self.assertTrue(any(n.startswith('xl/charts/chart') for n in archive.namelist()))
        self.assertTrue((Path(result['outdir']) / 'WPH_통합.xlsx').exists())
        with service.RUN_LOCK:
            with self.assertRaises(RuntimeError):
                service.run(self.root, [self.target], options, host_gap=0)

    def test_atomic_cache_failure_preserves_previous_snapshot(self):
        self.put('a.htm')
        store.collect(self.root, [self.target], host_gap=0)
        cache = next((self.root / '배치분석' / '누적').glob('*.json'))
        before = cache.read_bytes()
        self.put('b.htm', report('b.htm', job='B/1'))
        with patch('param_manager.atomicfile.os.replace', side_effect=OSError('locked')):
            with self.assertRaises(OSError):
                store.collect(self.root, [self.target], host_gap=0)
        self.assertEqual(cache.read_bytes(), before)
        self.assertEqual(store.collect(self.root, [self.target], host_gap=0)['parsed'], 1)

    def test_period_and_offline_cache_scope(self):
        self.put('a_26-Sep-19_(10.00.00)_BatchReport.htm')
        self.put('b_26-Aug-19_(10.00.00)_BatchReport.htm', report(job='OLD/2'))
        store.collect(self.root, [self.target], host_gap=0)
        target = dict(self.target, start='2026-09-01', end='2026-09-30')
        self.assertEqual(len(store.collect(self.root, [target], host_gap=0)['records']), 1)
        moved = self.source.with_name('disconnected')
        self.source.rename(moved)
        result = store.collect(self.root, [target], host_gap=0)
        self.assertEqual(len(result['records']), 1)
        self.assertEqual(len(result['errors']), 1)
        self.assertEqual(result['cached_only'], 1)


class Scheduling(unittest.TestCase):
    def test_daily_timestamp_and_busy_gate_without_tk_display(self):
        from param_manager.batchreport_ui import BatchReportMixin
        from param_manager import watcher
        from datetime import timedelta

        class App(BatchReportMixin):
            def __init__(self):
                self._cfg = {'batch_auto': True}
                self.local_dir = '/unused'
                self.launched = 0

            def _batch_save(self):
                pass

            def _set_status(self, text):
                self.status = text

            def _run_bg(self, work, done):
                self.launched += 1
                done(True, {'collection': {'errors': []}, 'dashboard_error': '',
                            'result': {'summary': {'Batch 수': 1}}, 'outdir': '/unused'})

            def after(self, ms, fn):
                self.delay = ms

        app = App()
        targets = [{'machine': 'A', 'folder': '/source', 'names': None}]
        options = {'metrics': ['M02'], 'valid_wafers': 25}
        app._batch_launch(targets, options, automatic=True)
        self.assertEqual(app.launched, 1)
        app._batch_tick()
        self.assertEqual(app.launched, 1)  # must NOT run every minute
        state = watcher.WatchState.from_dict(app._cfg['batch_schedule'])
        settings = watcher.WatchSettings(enabled=True, interval_hours=24)
        self.assertTrue(watcher.should_run(datetime.now() + timedelta(hours=25), settings, state))
        app._watch_busy = True
        app._batch_launch(targets, options, automatic=True)
        self.assertEqual(app.launched, 1)


if __name__ == '__main__':
    unittest.main()
