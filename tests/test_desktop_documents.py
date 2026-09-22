import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from param_manager import refdata,locking
from param_manager.desktop_documents import DesktopDocuments

class DocumentsTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.cfg=self.root/'config.json'
        self.cfg.write_text(json.dumps({'save_dir':str(self.root)}))
        self.path=refdata.ip_path(str(self.root));refdata.create_blank_ip(self.path)
        wb=openpyxl.load_workbook(self.path);ws=wb.active
        ws.cell(1,4,'접속ID');ws.cell(1,5,'비밀번호')
        ws.append(['AOI-01','10.0.0.1','Camtek','secret-user','secret-pass'])
        wb.save(self.path);wb.close()
        self.doc=DesktopDocuments(self.cfg);self.catalog=self.doc.open('ip')
        self.patch=patch.object(locking,'VERIFY_DELAY_SEC',0);self.patch.start();self.addCleanup(self.patch.stop)

    def test_theme_fill_and_number_preserved(self):
        from openpyxl.styles import PatternFill
        from openpyxl.styles.colors import Color
        wb=openpyxl.load_workbook(self.path);ws=wb.active
        ws['A2'].fill=PatternFill('solid',fgColor=Color(theme=4))   # theme color: shown as no color
        ws['B2']=42                                                 # numeric cell
        wb.save(self.path);wb.close()
        cat=self.doc.open('ip')
        self.assertEqual(self.doc.page({'snapshot':cat['snapshot']})['rows'][0]['cells'][0]['color'],'')
        cat=self.doc.edit({'snapshot':cat['snapshot'],'row':0,'column':0,'value':'AOI-01X','color':''})
        cat=self.doc.edit({'snapshot':cat['snapshot'],'row':0,'column':1,'value':'43','color':''})
        wb=openpyxl.load_workbook(self.path);self.addCleanup(wb.close);ws=wb.active
        self.assertEqual(ws['A2'].value,'AOI-01X')
        self.assertEqual(ws['A2'].fill.fgColor.theme,4)             # theme fill kept
        self.assertEqual(ws['B2'].value,43)                         # still a number
        self.doc.edit({'snapshot':cat['snapshot'],'row':0,'column':1,'value':'=1+1','color':''})
        wb2=openpyxl.load_workbook(self.path);self.addCleanup(wb2.close)
        self.assertEqual((wb2.active['B2'].value,wb2.active['B2'].data_type),('=1+1','s'))   # never a formula

    def test_choice_column_and_row_delete(self):
        self.assertEqual(self.catalog['columns'][2],{'type':'choice','choices':['Camtek','KLA']})
        with self.assertRaises(ValueError):
            self.doc.edit({'snapshot':self.catalog['snapshot'],'row':0,'column':2,'value':'Other','color':''})
        cat=self.doc.edit({'snapshot':self.catalog['snapshot'],'row':0,'column':2,'value':'KLA','color':''})
        cat=self.doc.append({'snapshot':cat['snapshot'],'values':['AOI-02','10.0.0.2','Camtek']})
        self.assertEqual(cat['total'],2)
        cat=self.doc.delete({'snapshot':cat['snapshot'],'row':0})
        self.assertEqual(cat['total'],1)
        wb=openpyxl.load_workbook(self.path);self.addCleanup(wb.close)
        self.assertEqual([c.value for c in wb.active[2]][:3],['AOI-02','10.0.0.2','Camtek'])
        self.assertIsNone(wb.active['D2'].value)     # deleted row's credentials went with it, not shifted
        with self.assertRaises(ValueError):
            self.doc.delete({'snapshot':cat['snapshot'],'row':5})

    def test_special_done_toggle(self):
        path=refdata.special_path(str(self.root));refdata.create_blank_special(path)
        wb=openpyxl.load_workbook(path);ws=wb.active
        ws.append(['2026-09-01','AOI-01','L1','S','PI','t','진행','Y','메모']);wb.save(path);wb.close()
        cat=self.doc.open('special')
        col=cat['headers'].index('종료 여부')
        self.assertEqual(cat['columns'][col],{'type':'bool'})
        row=self.doc.page({'snapshot':cat['snapshot']})['rows'][0]
        self.assertEqual(row['cells'][col]['value'],'☑')      # legacy 'Y' shown as checked
        with self.assertRaises(ValueError):
            self.doc.edit({'snapshot':cat['snapshot'],'row':0,'column':col,'value':'yes','color':''})
        self.doc.edit({'snapshot':cat['snapshot'],'row':0,'column':col,'value':'☐','color':''})
        rows,_=refdata.load_special(path)
        self.assertFalse(rows[0]['종료 여부'])
        self.assertEqual(rows[0]['특이사항'],'메모')

    def test_credentials_never_exposed(self):
        payload=self.doc.page({'snapshot':self.catalog['snapshot']})
        self.assertEqual(len(payload['rows'][0]['cells']),3)
        self.assertNotIn('secret',json.dumps(payload))

    def test_edit_preserves_other_cells_and_formats(self):
        result=self.doc.edit({'snapshot':self.catalog['snapshot'],'row':0,'column':1,'value':'10.0.0.2','color':'#112233'})
        self.assertNotEqual(result['snapshot'],self.catalog['snapshot'])
        wb=openpyxl.load_workbook(self.path);self.addCleanup(wb.close)
        self.assertEqual(wb.active['B2'].value,'10.0.0.2')
        self.assertEqual(wb.active['E2'].value,'secret-pass')
        self.assertEqual(wb.active['B2'].fill.fgColor.rgb[-6:],'112233')

    def test_stale_path_and_nonlisted_column_rejected(self):
        for bad in ('/etc','unknown'):
            with self.assertRaises(ValueError):self.doc.open(bad)
        self.catalog=self.doc.open('ip')
        with self.assertRaises(ValueError):
            self.doc.edit({'snapshot':self.catalog['snapshot'],'row':0,'column':4,'value':'bad','color':''})
        with open(self.path,'ab') as f:f.write(b'change')
        with self.assertRaises(ValueError):
            self.doc.edit({'snapshot':self.catalog['snapshot'],'row':0,'column':1,'value':'bad','color':''})

    def test_reference_formula_is_text(self):
        path=refdata.ref_path(str(self.root));refdata.create_blank_reference(path)
        cat=self.doc.open('reference')
        self.doc.edit({'snapshot':cat['snapshot'],'row':0,'column':1,'value':'=1+1','color':''})
        wb=openpyxl.load_workbook(path);self.addCleanup(wb.close)
        self.assertEqual(wb.active['B1'].data_type,'s')

    def test_append_preserves_hidden_credentials_and_literal_formula(self):
        result=self.doc.append({'snapshot':self.catalog['snapshot'],'values':['AOI-02','=1+1','Camtek']})
        self.assertEqual(result['total'],2)
        wb=openpyxl.load_workbook(self.path);self.addCleanup(wb.close)
        self.assertEqual(wb.active['E2'].value,'secret-pass')
        self.assertIsNone(wb.active['E3'].value)
        self.assertEqual(wb.active['B3'].value,'=1+1')
        self.assertEqual(wb.active['B3'].data_type,'s')

    def test_append_rejects_invalid_and_stale(self):
        original=Path(self.path).read_bytes()
        for values in ([],['','',''],['x']*4,[True,'',''],['x'*4001,'','']):
            with self.assertRaises(ValueError):
                self.doc.append({'snapshot':self.catalog['snapshot'],'values':values})
        self.assertEqual(Path(self.path).read_bytes(),original)
        with open(self.path,'ab') as f:f.write(b'changed')
        with self.assertRaises(ValueError):
            self.doc.append({'snapshot':self.catalog['snapshot'],'values':['new','','']})

    def test_append_other_lock_keeps_original(self):
        original=Path(self.path).read_bytes()
        locking.acquire(self.path,'another-user')
        self.addCleanup(lambda:locking.release(self.path,'another-user'))
        with self.assertRaises(ValueError):
            self.doc.append({'snapshot':self.catalog['snapshot'],'values':['new','','']})
        self.assertEqual(Path(self.path).read_bytes(),original)

    def test_append_failed_replace_preserves_workbook_and_cleans_temp(self):
        original=Path(self.path).read_bytes()
        with patch('param_manager.desktop_documents.os.replace',side_effect=OSError('blocked')):
            with self.assertRaises(OSError):
                self.doc.append({'snapshot':self.catalog['snapshot'],'values':['new','','']})
        self.assertEqual(Path(self.path).read_bytes(),original)
        self.assertEqual(list(self.root.rglob('.rev1-document-*')),[])

    def test_append_reference_and_missing_header(self):
        path=refdata.ref_path(str(self.root));refdata.create_blank_reference(path)
        cat=self.doc.open('reference')
        result=self.doc.append({'snapshot':cat['snapshot'],'values':['text']*len(cat['headers'])})
        self.assertEqual(result['total'],cat['total']+1)
        wb=openpyxl.load_workbook(self.path);wb.active['B1']='unknown';wb.save(self.path);wb.close()
        cat=self.doc.open('ip')
        with self.assertRaises(ValueError):
            self.doc.append({'snapshot':cat['snapshot'],'values':['new','','']})

if __name__=='__main__':unittest.main()
