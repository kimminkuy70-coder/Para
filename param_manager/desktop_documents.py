"""Allowlisted shared workbooks, paged reads and optimistic locked cell edits."""
import os
from pathlib import Path
import tempfile
from uuid import uuid4
import openpyxl
from openpyxl.styles import PatternFill

from . import engine, locking, namestore, refdata


def refdata_bool(value):
    return engine._s(value).strip() in ('☑','Y','1','True','종료','예')  # same set as refdata.load_special
from .desktop_batch import read_json

CHECKED, UNCHECKED = '☑', '☐'
# Columns with a fixed vocabulary, matching the tkinter editors.
KINDS_META = {'special': {refdata.SPECIAL_BOOL_COL: dict(type='bool')},
              'ip': {'장비종류': dict(type='choice', choices=list(refdata.DEVICE_TYPES))}}
KINDS = {'ip':(refdata.ip_path,refdata.IP_SHEET,refdata.IP_HEADERS),
         'special':(refdata.special_path,refdata.SPECIAL_SHEET,refdata.SPECIAL_HEADERS),
         'reference':(refdata.ref_path,refdata.REF_SHEET,None)}


class DesktopDocuments:
    def __init__(self, config_path=None):
        self.config_path=config_path or Path.home()/'.pi_param_manager.json'
        self.version=None

    def open(self, kind):
        self.version=None
        if not isinstance(kind,str) or kind not in KINDS:raise ValueError('문서 종류를 확인하세요')
        cfg=read_json(self.config_path)
        root=cfg.get('save_dir')
        if not root:return dict(snapshot=None,headers=[],total=0,source='')
        getter,sheet,headers=KINDS[kind]
        self.path=Path(getter(root)).absolute()
        self.root=Path(root).absolute()
        self.check_path()
        if not self.path.is_file():return dict(snapshot=None,headers=[],total=0,source=self.path.name)
        self.stamp=locking.file_stamp(str(self.path))
        wb=openpyxl.load_workbook(self.path)
        try:
            ws=wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
            if ws.max_row>100000 or ws.max_column>100:raise ValueError('문서가 너무 큽니다. 기존 프로그램에서 범위를 확인하세요.')
            self.sheet=ws.title;self.kind=kind
            if headers is None:
                self.headers=[openpyxl.utils.get_column_letter(i) for i in range(1,ws.max_column+1)]
                self.columns=list(range(1,ws.max_column+1));first=1
            else:
                actual=[engine._s(c.value).strip() for c in ws[1]]
                self.headers=list(headers)
                self.columns=[actual.index(h)+1 if h in actual else None for h in headers];first=2
            self.cells=[]
            for r in range(first,ws.max_row+1):
                line=[]
                for c in self.columns:
                    cell=ws.cell(r,c) if c is not None else None
                    color=''
                    if cell and cell.fill.patternType=='solid' and cell.fill.fgColor.type=='rgb':
                        color='#'+cell.fill.fgColor.rgb[-6:]
                    line.append(dict(value=engine._s(cell.value) if cell else '',color=color,editable=c is not None))
                self.cells.append((r,line))
            if self.stamp!=locking.file_stamp(str(self.path)):raise ValueError('문서가 변경되었습니다. 다시 열어 주세요.')
        finally:wb.close()
        meta=KINDS_META.get(kind,{})
        self.columns_meta=[meta.get(h,dict(type='text')) for h in self.headers]
        for _,line in self.cells:
            for cell,info in zip(line,self.columns_meta):
                if info['type']=='bool':
                    # Stored as ☑/☐ by the tkinter app; older files may hold Y/True/종료.
                    cell['value']=CHECKED if refdata_bool(cell['value']) else UNCHECKED
        self.version=uuid4().hex
        return dict(snapshot=self.version,headers=self.headers,columns=self.columns_meta,
                    total=len(self.cells),source=self.path.name)

    def check_path(self):
        if (any(p.is_symlink() or (hasattr(p,'is_junction') and p.is_junction())
                for p in (self.path,*self.path.parents))
                or not self.path.resolve().is_relative_to(self.root.resolve())):
            raise ValueError('공유 문서 연결 경로를 사용할 수 없습니다')

    def page(self, params):
        if set(params)-{'snapshot','offset','limit'} or not self.version or params.get('snapshot')!=self.version:
            raise ValueError('문서를 새로고침하세요')
        offset,limit=params.get('offset',0),params.get('limit',100)
        if type(offset) is not int or offset<0 or type(limit) is not int or not 1<=limit<=100:
            raise ValueError('문서 조회 범위를 확인하세요')
        return dict(rows=[{'id':i,'cells':line} for i,(_,line) in enumerate(self.cells[offset:offset+limit],offset)],total=len(self.cells))

    def edit(self, params):
        if set(params)!={'snapshot','row','column','value','color'} or not self.version or params['snapshot']!=self.version:
            raise ValueError('문서를 새로고침하세요')
        r,c=params['row'],params['column'];value=params['value']
        if type(r) is not int or not 0<=r<len(self.cells) or type(c) is not int or not 0<=c<len(self.columns) or self.columns[c] is None:
            raise ValueError('편집 가능한 셀을 선택하세요')
        if not isinstance(value,str) or len(value)>4000:raise ValueError('내용은 4000자까지 입력하세요')
        color=namestore.normalize_color(params['color'])
        value=self._checked_value(c,value)
        return self._write([(self.cells[r][0],self.columns[c],value,color)])

    def _checked_value(self,c,value):
        info=self.columns_meta[c]
        if info['type']=='bool':
            if value not in (CHECKED,UNCHECKED,''):raise ValueError('종료 여부는 ☑ 또는 ☐ 입니다')
            return value or UNCHECKED
        if info['type']=='choice' and value and value not in info['choices']:
            raise ValueError('허용된 값 중에서 고르세요: '+', '.join(info['choices']))
        return value

    def delete(self, params):
        """Delete one data row (tkinter right-click 행 삭제), same lock/change checks."""
        if set(params)!={'snapshot','row'} or not self.version or params['snapshot']!=self.version:
            raise ValueError('문서를 새로고침하세요')
        r=params['row']
        if type(r) is not int or not 0<=r<len(self.cells):raise ValueError('삭제할 행을 선택하세요')
        return self._write([],delete_row=self.cells[r][0])

    def append(self, params):
        if set(params)!={'snapshot','values'} or not self.version or params['snapshot']!=self.version:
            raise ValueError('문서를 새로고침하세요')
        values=params['values']
        if (not isinstance(values,list) or len(values)!=len(self.columns)
                or any(not isinstance(v,str) or len(v)>4000 for v in values)
                or not any(v.strip() for v in values)):
            raise ValueError('새 행의 내용을 입력하세요. 각 셀은 4000자까지 가능합니다.')
        if any(c is None for c in self.columns):
            raise ValueError('필수 열이 없는 문서입니다. 기존 프로그램에서 양식을 확인하세요.')
        row=self.cells[-1][0]+1 if self.cells else (1 if self.kind=='reference' else 2)
        if row>100000:raise ValueError('문서 최대 행 수를 초과합니다.')
        values=[self._checked_value(i,v) for i,v in enumerate(values)]
        return self._write([(row,c,v,'') for c,v in zip(self.columns,values)])

    def _write(self, updates, delete_row=None):
        self.check_path();path=str(self.path);user=engine.current_user()
        previous=locking.status(path,user);state=locking.acquire(path,user)
        if not state.editable:raise ValueError(locking.holder_message(state,self.path.name))
        wb=None;temporary=None
        try:
            check=locking.check_before_save(path,user,self.stamp)
            if not check['ok']:raise ValueError(check['reason'])
            wb=openpyxl.load_workbook(path);ws=wb[self.sheet]
            for row,column,value,color in updates:
                cell=ws.cell(row,column);cell.value=value;cell.data_type='s'
                cell.fill=PatternFill('solid',fgColor=color[1:]) if color else PatternFill()
            if delete_row is not None:
                ws.delete_rows(delete_row)
            fd,temporary=tempfile.mkstemp(prefix='.rev1-document-',suffix='.xlsx',dir=self.path.parent)
            os.close(fd);wb.save(temporary)
            check=locking.check_before_save(path,user,self.stamp)
            if not check['ok'] or locking.file_stamp(path)!=self.stamp:raise ValueError(check['reason'] or '저장 직전 문서가 변경되었습니다.')
            self.check_path()
            os.replace(temporary,path);temporary=None
        finally:
            if wb:wb.close()
            if temporary and os.path.exists(temporary):os.unlink(temporary)
            if previous.status!='mine':locking.release(path,user)
        return self.open(self.kind)
