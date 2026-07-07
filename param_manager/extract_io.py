"""initial(01_초안)/final(02_확정) 스냅샷 엑셀 기록/판독.

보이는 시트는 공용 파일과 동일한 PI_ALL/RDL_ALL 양식(가독성 통일, 2026-07 확정) —
engine.create_from_records 를 그대로 사용하므로 프로그램에서 다시 열 수도 있다.
여기에 두 시트를 덧붙인다:
  - _EXTRACT_MAP(숨김): Row_ID ↔ 설정파일/섹션/키/Raw/변환방식 — 값 재추출용
  - 추출_요약(표시):    단계/일시/레벨/호기/소스 폴더/항목 수
"""

from __future__ import annotations

import json
import os
from datetime import datetime

import openpyxl

from . import engine

SHEET_MAP = "_EXTRACT_MAP"
SHEET_SUMMARY = "추출_요약"
SCALES_LABEL = "변환계수(JSON)"      # 추출_요약에 저장되는 변형별 계수
MAP_HEADERS = ["Row_ID", "설정파일", "설정 Section", "설정 Parameter",
               "Raw Value", "변환방식", "Unit", "Source Path"]


def write_snapshot(dest_xlsx: str, records: list[dict], machines: list[str],
                   sheet_name: str, extracts: list[dict | None],
                   stage: str, level: str = "", aoi: str = "",
                   source: str = "", user: str | None = None,
                   scales: dict | None = None) -> str:
    """records(공용 스키마 dict 목록)로 스냅샷 파일 생성.

    extracts: records 와 같은 길이/순서의 재추출 메타
              ({src_file, section, key, raw?, transform, source_path} 또는 None).
    stage: "initial" / "final" (요약 표기용).
    scales: 변형별 변환계수 {변형: 계수} — 값 업데이트가 같은 계수로 재파싱하도록 보존.
    반환: dest_xlsx.
    """
    os.makedirs(os.path.dirname(os.path.abspath(dest_xlsx)), exist_ok=True)
    repo = engine.create_from_records(dest_xlsx, records, machines,
                                      user=user, sheet_name=sheet_name)
    # 새로 생성 경로는 records 순서가 repo.rows 순서로 보존된다(Row_ID 정렬 목적).
    row_ids = [pr.row_id for pr in repo.rows]

    wb = openpyxl.load_workbook(dest_xlsx)
    ws = wb.create_sheet(SHEET_MAP)
    ws.sheet_state = "hidden"
    ws.append(MAP_HEADERS)
    for rid, ext in zip(row_ids, extracts):
        ext = ext or {}
        ws.append([rid, ext.get("src_file", ""), ext.get("section", ""),
                   ext.get("key", ""), engine._s(ext.get("raw", "")),
                   ext.get("transform", ""), ext.get("unit", ""),
                   ext.get("source_path", "")])

    sm = wb.create_sheet(SHEET_SUMMARY)
    sm.append(["항목", "값"])
    for k, v in [("단계", stage), ("생성일시", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                 ("레시피 레벨", level), ("호기", aoi), ("소스", source),
                 ("파라미터 수", len(records)), ("시트", sheet_name)]:
        sm.append([k, engine._s(v)])
    if scales:
        sm.append([SCALES_LABEL, json.dumps(scales, ensure_ascii=False)])
    sm.column_dimensions["A"].width = 16
    sm.column_dimensions["B"].width = 80
    wb.save(dest_xlsx)
    return dest_xlsx


def read_scales(xlsx_path: str) -> dict:
    """양식/스냅샷의 추출_요약에서 변형별 변환계수 {변형: 계수} 판독. 없으면 {}."""
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    except Exception:  # noqa: BLE001
        return {}
    if SHEET_SUMMARY not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[SHEET_SUMMARY]
    out: dict = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and engine._s(row[0]) == SCALES_LABEL and len(row) > 1 and row[1]:
            try:
                out = {k: float(v) for k, v in json.loads(row[1]).items()}
            except Exception:  # noqa: BLE001
                out = {}
            break
    wb.close()
    return out


def read_extract_map(xlsx_path: str) -> dict[str, dict]:
    """스냅샷/양식 파일의 _EXTRACT_MAP → {Row_ID: 메타 dict}. 없으면 {}."""
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    except Exception:  # noqa: BLE001
        return {}
    if SHEET_MAP not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[SHEET_MAP]
    out: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        vals = list(row) + [""] * (len(MAP_HEADERS) - len(row))
        out[str(row[0])] = {
            "src_file": vals[1] or "", "section": vals[2] or "",
            "key": vals[3] or "", "raw": vals[4],
            "transform": vals[5] or "", "unit": vals[6] or "",
            "source_path": vals[7] or "",
        }
    wb.close()
    return out
