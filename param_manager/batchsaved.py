"""저장된 결과 파일(분석 HTML · Excel · 가동률 대시보드) — 앱 화면(v12)과 같은 계산 (이슈 #7).

기준은 무조건 프로그램(사용자 지시 2026-10-05): 숫자는 앱이 쓰는 `batchview.View.payload()` 를 그대로 받아
앱(`frontend/src/batchData.ts` · `BatchLot/BatchUtil/BatchWph.tsx`)과 같은 식으로 더한다. 여기서 새로 정의하는 지표는 없다.

- 가동률: 호기마다 하루 24시간 = 웨이퍼 처리(p + 점검 ck) + Error · 중단 및 조치(e + sd + so) + 유휴(나머지).
- WPH: 정상 = w25 × 3600 ÷ s25, 실제 = ps × 3600 ÷ (p + 손실), 하루 생산능력 = 24 × 실제, 처리량 감소 = 1 − 실제 ÷ 정상.
- 호기 색은 `MACH_PAL`(앱 `batchData.MACH_PAL` 과 같은 표 · 같은 순서 — 호기 이름 정렬 순).

파일 접근은 `write_excel`(경로는 호출자가 정함) 하나뿐이고, HTML 은 문자열만 만든다(외부 스크립트 · 인터넷 없음).
"""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import date, datetime, timedelta

from . import lotmodel as lm, lotreport
from .batchreport_output import PAGE_CSS, esc

MACH_PAL = ['#1f77b4', '#2a9d8f', '#e07b39', '#7b5ea7', '#3a7d44', '#c44e7a', '#5b8fc7', '#b8860b',
            '#17859b', '#8c564b', '#4c6ef5', '#6b8e23', '#d1495b', '#31517c', '#9c6ade', '#00a676']
C_PROC, C_LOSS, C_IDLE, C_ERR = '#10b981', '#c2410c', '#dfe5ec', '#c2410c'
C_STOP, C_DEFECT, C_OK, C_WAIT = '#7c8fb0', '#9b4dca', '#10b981', '#d69e2e'
STOPK = {'d': '작업자 중단 · Defect 과다', 'o': '작업자 중단 · 그 외'}
DONE, RE, OPEN = lm.LOT_DONE, lm.LOT_RESCANNED, lm.LOT_OPEN

# 화면 ? 버튼(`metricHelp.ts`)의 정의를 줄여 옮긴 것 — 파일만 받은 사람도 숫자의 뜻을 알 수 있게.
DEFS = [
    ('가동률', '웨이퍼 처리 시간 ÷ 달력 시간. 달력 시간 = 24시간 × 일수 × 호기 수(4조 3교대).'),
    ('웨이퍼 처리', 'Error · 중단 없는 Batch Report는 Batch Time 전체, 있으면 Pass 장수 × 1장 처리 시간(Batch Time 이내). '
                '점검 스캔(S/M 0 · FOCUS 등)도 여기에 넣고 레시피 WPH에서는 뺍니다.'),
    ('Error · 중단 및 조치', 'Error · 작업자 중단이 있는 Batch Report의 나머지 시간 + 같은 호기에서 같은 Lot을 다시 스캔하기까지 '
                       '장비가 아무것도 스캔하지 않은 시간. Error 원문별 · 중단 종류별로 나눕니다.'),
    ('유휴', '달력 시간 − 웨이퍼 처리 − Error · 중단 및 조치. 물량 없음 · PM · 다운은 구분하지 못합니다.'),
    ('1장 처리 시간', '정상 25매 Batch Report(Error · 중단 없이 25매 모두 Pass)의 Batch Time 합 ÷ 장수 합(로봇 이동 포함). '
                 'Avg. Scan Time은 Batch Report 원문 값(순수 스캔).'),
    ('정상 WPH', '정상 25매 Batch Report 장수 합 × 3,600 ÷ Batch Time 합.'),
    ('실제 WPH', 'Pass 장수 × 3,600 ÷ (웨이퍼 처리 + Error · 중단 및 조치) — 장비가 놀던 시간만 뺍니다.'),
    ('처리량 감소', '1 − 실제 WPH ÷ 정상 WPH.'),
    ('하루 생산능력', '24 × 실제 WPH. 레시피 합계 = 그 레시피를 돌린 호기들의 합(24시간 쉬지 않는다는 가정의 최대치).'),
    ('작업자 중단', '앞에 Error 없이 Aborted.로 멈춘 Batch Report — Error로 세지 않지만 시간은 손실로 셉니다. 멈출 때 wafer의 '
               'Faults가 그 레시피 정상 wafer Faults 상위 1%보다 많으면 Defect 과다 중단(정상 wafer 100장 미만인 레시피는 판정하지 않음).'),
    ('스캔 안 한 슬롯', '앞에 Error 없이 Skipped.인 행은 wafer로 세지 않습니다.'),
    ('같은 wafer 다시 스캔', '이미 Pass한 wafer를 다시 스캔한 것도 처리 시간 · 장수에 넣습니다(투입 · 출하 정보가 없어 필요한 '
                        '재스캔인지 알 수 없음 — 재스캔 참고 비율로 따로 표시).'),
]


def mach_colors(ids):
    """호기 → 색. 앱 `batchData.machColors` 와 같은 규칙(이름 정렬 순서로 MACH_PAL 을 돌려 씀)."""
    return {m: MACH_PAL[i % len(MACH_PAL)] for i, m in enumerate(sorted(set(ids)))}


# ---------------------------------------------------------------------- batchData.ts 와 같은 합계
_SUMS = ('p', 'du', 'ck', 'ps', 'dn', 'n', 'ne', 'w25', 's25', 'aw', 'asum')


def empty_agg():
    return dict({k: 0 for k in _SUMS}, e=0.0, sd=0.0, so=0.0, nsd=0, nso=0, err={}, sdl=set(), sol=set(), rec={})


def add_u(a, u):
    for k in _SUMS:
        a[k] += u[k]
    a['sd'] += u['sd'][0]
    a['nsd'] += u['sd'][1]
    a['sdl'].update(u['sd'][2])
    a['so'] += u['so'][0]
    a['nso'] += u['so'][1]
    a['sol'].update(u['so'][2])
    if u['p']:
        a['rec'][u['r']] = a['rec'].get(u['r'], 0) + u['p']
    for k, x in u['e'].items():
        a['e'] += x[0]
        t = a['err'].setdefault(k, [0.0, 0, set()])
        t[0] += x[0]
        t[1] += x[1]
        t[2].update(x[2])
    return a


def sum_u(rows):
    a = empty_agg()
    for u in rows:
        add_u(a, u)
    return a


def add_agg(a, b):
    o = empty_agg()
    for x in (a, b):
        for k in _SUMS + ('e', 'sd', 'so', 'nsd', 'nso'):
            o[k] += x[k]
    return o


def loss_of(a):
    return a['e'] + a['sd'] + a['so']


def wph_normal(a):
    return a['w25'] * 3600 / a['s25'] if a['s25'] else None


def wph_actual(a):
    d = a['p'] + loss_of(a)
    return a['ps'] * 3600 / d if d > 0 else None


def unit_sec(a):
    return a['s25'] / a['w25'] if a['w25'] else None


def avg_scan(a):
    return a['asum'] / a['aw'] if a['aw'] else None


def cap(a):
    w = wph_actual(a)
    return None if w is None else w * 24


def drop(n, w):
    return (1 - w / n) * 100 if n and w is not None else None


def comp(a, cal):
    """시간 3칸 — 가동률 화면 `comp` 와 같음."""
    proc, loss = a['p'] + a['ck'], loss_of(a)
    return dict(proc=proc, loss=loss, idle=max(0.0, cal - proc - loss), cal=cal, a=a)


def pct(x, t):
    return x / t * 100 if t > 0 else None


def num(v, d=0):
    if v is None:
        return '—'
    return f'{v:,.{d}f}' if d else f'{round(v):,}'


def hrs(s):
    return num(s / 3600, 1)


def pc(v):
    return '—' if v is None else f'{v:,.1f}%'


def days_between(a, b):
    if not a or not b:
        return []
    d0, d1 = date.fromisoformat(a), date.fromisoformat(b)
    return [(d0 + timedelta(days=i)).isoformat() for i in range((d1 - d0).days + 1)]


def bucket(d, unit):
    if unit == 'd':
        return d
    if unit == 'm':
        return d[:7]
    x = date.fromisoformat(d)
    return (x - timedelta(days=x.weekday())).isoformat() + ' 주'


def period_name(k, unit):
    """주 = 'YYYY WWn (MM/DD~MM/DD)'(ISO 주, 앱 `periodName` 과 같음)."""
    if unit != 'w':
        return k
    st = date.fromisoformat(k[:10])
    en = st + timedelta(days=6)
    y, w, _ = st.isocalendar()
    return f"{y} WW{w} ({st:%m/%d}~{en:%m/%d})"


# ---------------------------------------------------------------------- 계산 묶음
class Saved:
    """View payload → 저장 파일에 쓸 숫자. 범위 = 조사 결과 전체(처음 ~ 마지막 Batch Report 날짜), 호기 = 조사 범위 호기."""

    def __init__(self, payload, machines=None, scope='', saved=0, created=None):
        self.P = payload
        self.ids = list(dict.fromkeys(machines or sorted({r['m'] for r in payload['R']})))
        self.range = payload['range']
        self.days = days_between(*self.range)
        self.U = [u for u in payload['U'] if self.range[0] <= u['d'] <= self.range[1]]
        self.scope = scope
        self.saved = payload.get('saved', saved)
        self.created = created or datetime.now()
        self.colors = mach_colors(self.ids + [u['m'] for u in self.U])

    # ---- Lot 추적
    def lot_kpis(self):
        lots, R = self.P['lots'], self.P['R']
        chk = len(self.P['excluded'])
        return [('조사 범위 내 Lot Scan 완료', len(lots), 'Lot', 'S/M 영문 3글자 + Lot ID 기준', ''),
                ('조사 범위 내 Batch Report 존재', len(R), '개', f'Lot {len(R) - chk:,} · 점검 스캔 {chk:,}', ''),
                ('재스캔으로 완료된 Lot', sum(l['state'] == RE for l in lots), 'Lot', 'Error · 작업자 중단 뒤 재스캔하여 Pass', 'a'),
                ('Pass하지 못한 Lot', sum(l['state'] == OPEN for l in lots), 'Lot', '다른 호기에서 Scan되었는지 확인', 'r'),
                ('중복 Scan한 wafer가 존재하는 Lot', sum(1 for l in lots if l['dup']), 'Lot', '취합은 가장 나중 Pass(+ 저장된 사람 선택)', 't')]

    def causes(self):
        """멈춘 이유 요약 · Error — [원문, Lot, 재스캔하여 Pass한 Lot, Pass하지 못한 Lot, wafer] (앱 Lot 추적과 같음)."""
        m = {}
        for l in self.P['lots']:
            for c, w, o in l['cz']:
                t = m.setdefault(c, [0, 0, 0, 0])
                t[0] += 1
                t[3] += w
                t[2 if o else 1] += 1
        return sorted(([k, *v] for k, v in m.items()), key=lambda x: (-x[1], -x[4]))

    def stops(self):
        """멈춘 이유 요약 · 작업자 중단 — 종류별 [이름, Lot, 재스캔 Pass Lot, Pass 못 한 Lot, wafer, Batch Report]."""
        m = {'d': [0, 0, 0, 0, 0], 'o': [0, 0, 0, 0, 0]}
        for l in self.P['lots']:
            for k, w, o in l['sz']:
                t = m[k]
                t[0] += 1
                t[3] += w
                t[2 if o else 1] += 1
        for s in self.P['stops']:
            m[s['k']][4] += 1
        return [[STOPK[k], *m[k]] for k in ('d', 'o')]

    # ---- 가동률
    def util_all(self):
        return comp(sum_u([u for u in self.U if u['m'] in self.ids]), len(self.days) * 86400 * max(1, len(self.ids)))

    def util_machines(self):
        cal = len(self.days) * 86400
        return [(m, comp(sum_u([u for u in self.U if u['m'] == m]), cal)) for m in self.ids]

    def util_periods(self, unit):
        o = {}
        for d in self.days:
            o.setdefault(bucket(d, unit), [0, empty_agg()])[0] += 1
        for u in self.U:
            if u['m'] in self.ids:
                k = bucket(u['d'], unit)
                if k in o:
                    add_u(o[k][1], u)
        nm = max(1, len(self.ids))
        return [(k, comp(a, n * 86400 * nm)) for k, (n, a) in sorted(o.items(), reverse=True)]

    def util_days_machines(self):
        """날짜 × 호기 시간 3칸(Excel 용)."""
        by = defaultdict(list)
        for u in self.U:
            by[(u['d'], u['m'])].append(u)
        return [(d, m, comp(sum_u(by.get((d, m), [])), 86400)) for d in self.days for m in self.ids]

    def loss_reasons(self, a):
        rs = [(e, t[0], len(t[2]), C_ERR) for e, t in a['err'].items()]
        rs += [(STOPK['d'], a['sd'], len(a['sdl']), C_DEFECT), (STOPK['o'], a['so'], len(a['sol']), C_STOP)]
        return sorted((r for r in rs if r[1] > 0), key=lambda r: -r[1])

    # ---- WPH
    def wph_rows(self):
        return [u for u in self.U if u['m'] in self.ids and (u['n'] or u['p'])]

    def wph_cells(self):
        o = defaultdict(list)
        for u in self.wph_rows():
            o[(u['r'], u['m'])].append(u)
        cells = {k: sum_u(v) for k, v in o.items()}
        cells = {k: a for k, a in cells.items() if a['n'] > 0}
        recipes = sorted({u['r'] for u in self.wph_rows() if u['n']})
        return recipes, cells


def _job(r):
    return r.split(' · ')[0]


def _step(r):
    return r[len(_job(r)) + 3:]


# ---------------------------------------------------------------------- HTML
SAVED_CSS = """
.doc{max-width:1180px}
.kpis.k5{grid-template-columns:repeat(5,1fr)}
.kpi .u{font-size:13px;font-weight:700;margin-left:2px}
.kpi .s{font-size:11px;color:var(--faint);margin-top:2px}
.kpi .sw{display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:5px;vertical-align:0}
.prows{display:flex;flex-direction:column;gap:5px;margin:8px 0}
.prow{display:grid;grid-template-columns:150px minmax(0,1fr) 70px;gap:10px;align-items:center;font-size:12.5px}
.prow.ph{font-size:11px;color:var(--faint);font-weight:700}
.periodnav button.on{background:var(--navy);color:#fff}
.prow .pl{font-weight:800;color:var(--navy);white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.prow .pv{text-align:right;font-variant-numeric:tabular-nums;font-weight:700}
.stack{display:flex;height:20px;border-radius:5px;overflow:hidden;background:#f3f5f8}
.stack>span{height:100%}
.caprow{display:grid;grid-template-columns:minmax(150px,240px) minmax(0,1fr) 150px;gap:12px;align-items:center;font-size:12.5px;margin:6px 0}
.caprow .cl{font-weight:800;color:var(--navy);overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.caprow .cl small{display:block;font-weight:400;font-size:10.5px;color:var(--faint);overflow:hidden;text-overflow:ellipsis}
.capbar{display:flex;height:26px;border-radius:5px;overflow:hidden;background:#eef1f5}
.capbar>span{height:100%;display:flex;align-items:center;justify-content:center;color:#fff;font-size:11px;font-weight:700;overflow:hidden;white-space:nowrap;border-right:1px solid #fff}
.capv{text-align:right;font-variant-numeric:tabular-nums;color:var(--soft)}.capv b{color:var(--navy);font-size:14px;margin-right:2px}
.capv small{display:block;font-size:11px}.capv small.none{color:var(--rust);font-weight:700}
.capbar>span.capnone{flex:1;background:transparent;color:var(--faint);font-weight:400;justify-content:flex-start;padding-left:10px;border:0}
.keys{display:flex;flex-wrap:wrap;gap:4px 16px;font-size:11.5px;color:var(--soft);margin:8px 0}
.keys i{display:inline-block;width:12px;height:12px;border-radius:2px;margin-right:5px;vertical-align:-1px}
i.msw{display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:6px;vertical-align:-1px}
.hbar{display:grid;grid-template-columns:minmax(160px,320px) minmax(0,1fr) 150px;gap:10px;align-items:center;font-size:12.5px;margin:4px 0}
.hbar .lab{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.hbar .track{height:14px;background:#eef1f5;border-radius:4px;overflow:hidden}.hbar .track>span{display:block;height:100%}
.hbar .val{text-align:right;font-variant-numeric:tabular-nums;color:var(--soft)}
td.num,th.num{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
tr.total td{background:var(--navy-bg)!important;font-weight:700}
tr.grouphead td{background:var(--slate-bg)!important;font-weight:800;color:var(--navy)}
.tscroll{overflow-x:auto}
dl.defs{display:grid;grid-template-columns:170px 1fr;gap:8px 16px;margin:10px 0}
dl.defs dt{font-weight:800;color:var(--navy);font-size:13px}dl.defs dd{margin:0;font-size:13px}
@media(max-width:720px){.kpis.k5{grid-template-columns:repeat(2,1fr)}.prow,.caprow,.hbar{grid-template-columns:1fr}dl.defs{grid-template-columns:1fr}}
"""

TAB_JS = ('function showMetric(k){document.querySelectorAll("section.metric").forEach(function(s){s.hidden=(s.id!=="sec-"+k)});'
          'document.querySelectorAll(".mbtn").forEach(function(b){b.classList.toggle("on",b.dataset.sec===k)});window.scrollTo(0,0);}'
          'function period(p){document.querySelectorAll("[data-period]").forEach(function(e){e.hidden=e.dataset.period!==p});'
          'document.querySelectorAll(".periodnav button").forEach(function(b){b.classList.toggle("on",b.dataset.p===p)});}')


def _kpis(items, cls='k4'):
    out = []
    for label, value, unit, sub, color, tone in items:
        sw = f'<i class="sw" style="background:{color}"></i>' if color else ''
        out.append(f'<div class="kpi {tone}"><div class="n">{esc(value)}<span class="u">{esc(unit)}</span></div>'
                   f'<div class="l">{sw}{esc(label)}</div><div class="s">{esc(sub)}</div></div>')
    return f'<div class="kpis {cls}">' + ''.join(out) + '</div>'


def _stack(c):
    parts = []
    for k, col in (('proc', C_PROC), ('loss', C_LOSS), ('idle', C_IDLE)):
        if c[k] > 0 and c['cal'] > 0:
            parts.append(f'<span style="width:{c[k] / c["cal"] * 100:.2f}%;background:{col}"></span>')
    return '<span class="stack">' + ''.join(parts) + '</span>'


def _tip(name, c):
    return (f'{name}\n웨이퍼 처리 {hrs(c["proc"])}h · {pc(pct(c["proc"], c["cal"]))}\n'
            f'Error · 중단 및 조치 {hrs(c["loss"])}h · {pc(pct(c["loss"], c["cal"]))}\n'
            f'유휴 {hrs(c["idle"])}h · {pc(pct(c["idle"], c["cal"]))}\n달력 시간 {hrs(c["cal"])}h')


TIME_KEYS = ('<div class="keys"><span><i style="background:%s"></i>웨이퍼 처리</span><span><i style="background:%s"></i>Error · 중단 및 조치</span>'
             '<span><i style="background:%s"></i>유휴</span></div>') % (C_PROC, C_LOSS, C_IDLE)


def util_section(s):
    a = s.util_all()
    nm = max(1, len(s.ids))
    parts = [_kpis([
        ('가동률', num(pct(a['proc'], a['cal']), 1), '%', f'웨이퍼 처리 {hrs(a["proc"])}시간', C_PROC, 't'),
        ('Error · 중단 및 조치', num(pct(a['loss'], a['cal']), 1), '%',
         f'Error {hrs(a["a"]["e"])} + 작업자 중단 {hrs(a["a"]["sd"] + a["a"]["so"])}시간', C_LOSS, 'r'),
        ('유휴', num(pct(a['idle'], a['cal']), 1), '%', f'{hrs(a["idle"])}시간', C_IDLE, ''),
        ('달력 시간', hrs(a['cal']), 'h', f'{len(s.days)}일 × {nm}대 · 하루 24시간', '', '')])]
    parts.append('<h2>호기별 시간 구성</h2><div class="sub">100% = 24시간 × 조사 일수. 막대에 마우스를 올리면 시간이 나옵니다.</div>' + TIME_KEYS)
    rows = [f'<div class="prows"><div class="prow ph"><span>호기</span><span>100% = 24시간 × {len(s.days)}일</span><span>가동률</span></div>']
    for m, c in s.util_machines():
        rows.append(f'<div class="prow" title="{esc(_tip(m, c))}"><span class="pl">{esc(m)}</span>{_stack(c)}'
                    f'<span class="pv">{pc(pct(c["proc"], c["cal"]))}</span></div>')
    parts.append(''.join(rows) + '</div>')
    parts.append('<h2>기간별 시간 구성</h2><div class="sub">전체 호기 합계. 100% = 호기 수 × 24시간 × 그 기간 일수(조사 범위 밖 날짜는 뺌).</div>'
                 '<div class="periodnav"><button data-p="d" onclick="period(\'d\')">일</button><button data-p="w" class="on" onclick="period(\'w\')">주</button>'
                 '<button data-p="m" onclick="period(\'m\')">월</button></div>')
    for unit in ('d', 'w', 'm'):
        rows = [f'<div data-period="{unit}"{"" if unit == "w" else " hidden"}><div class="prows">'
                '<div class="prow ph"><span>기간</span><span>시간 구성</span><span>가동률</span></div>']
        for k, c in s.util_periods(unit):
            label = period_name(k, unit) if unit == 'w' else k
            rows.append(f'<div class="prow" title="{esc(_tip(label, c))}"><span class="pl">{esc(label)}</span>{_stack(c)}'
                        f'<span class="pv">{pc(pct(c["proc"], c["cal"]))}</span></div>')
        parts.append(''.join(rows) + '</div></div>')
    # 시간 구성 상세 · 손실 이유
    A = a['a']
    det = [('웨이퍼 처리', a['proc'], C_PROC, False), ('그중 같은 wafer 다시 스캔', A['du'], C_OK, True),
           ('그중 점검 스캔', A['ck'], '#94a3b8', True), ('Error · 중단 및 조치', a['loss'], C_LOSS, False),
           ('Error', A['e'], C_ERR, True), (STOPK['d'], A['sd'], C_DEFECT, True), (STOPK['o'], A['so'], C_STOP, True),
           ('유휴', a['idle'], C_IDLE, False)]
    parts.append('<h2>시간 구성 상세 · 전체 호기</h2><table><thead><tr><th>항목</th><th class="num">시간(h)</th><th class="num">달력 시간 대비</th></tr></thead><tbody>'
                 + ''.join(f'<tr><td>{"&nbsp;&nbsp;└ " if sub else ""}<i class="msw" style="background:{col}"></i>{esc(n)}</td>'
                           f'<td class="num">{hrs(x)}</td><td class="num">{pc(pct(x, a["cal"]))}</td></tr>' for n, x, col, sub in det)
                 + '</tbody></table>')
    rs = s.loss_reasons(A)
    mx = rs[0][1] if rs else 1
    parts.append('<h2>무엇으로 시간을 잃었나</h2><div class="sub">Batch Report의 남은 시간 + 다시 스캔하기까지 조치. '
                 '작업자 중단은 Error가 아니지만 결과를 내지 못한 시간이라 여기에 넣습니다.</div>')
    parts.append(''.join(f'<div class="hbar"><span class="lab" title="{esc(n)}">{esc(n)}</span><span class="track">'
                         f'<span style="width:{x / mx * 100:.1f}%;background:{col}"></span></span>'
                         f'<span class="val">{hrs(x)}h · Lot {lots:,}</span></div>' for n, x, lots, col in rs)
                 or '<p class="info">잃은 시간이 없습니다.</p>')
    rec = sorted(A['rec'].items(), key=lambda x: -x[1])
    rt = sum(x for _, x in rec) or 1
    parts.append('<h2>레시피(Job · Recipe(s))별 웨이퍼 처리 시간</h2>'
                 + (''.join(f'<div class="hbar"><span class="lab" title="{esc(r)}">{esc(r)}</span><span class="track">'
                            f'<span style="width:{x / rt * 100:.1f}%;background:var(--navy)"></span></span>'
                            f'<span class="val">{pc(x / rt * 100)} · {hrs(x)}h</span></div>' for r, x in rec)
                    or '<p class="info">웨이퍼 처리 없음</p>'))
    return ''.join(parts)


def wph_section(s):
    recipes, cells = s.wph_cells()
    A = sum_u(s.wph_rows())
    wn, wa = wph_normal(A), wph_actual(A)
    parts = [_kpis([
        ('정상 WPH', num(wn, 1), '', f'정상 25매 Batch Report {num(A["w25"] / 25)}장', '', ''),
        ('실제 WPH', num(wa, 1), '', f'Pass {num(A["ps"])}장 · 유휴만 뺀 시간', '', 't'),
        ('처리량 감소', num(drop(wn, wa), 1), '%', '1 − 실제 ÷ 정상', '', 'r'),
        ('레시피 · 호기', f'{len(recipes)} · {len({m for _, m in cells})}', '', '하루 생산능력은 아래 레시피 비교 · 표', '', 'a')])]
    col = s.colors
    rows = []
    for r in recipes:
        allc = [a for (rr, _), a in cells.items() if rr == r]
        cs = sorted(((m, a) for (rr, m), a in cells.items() if rr == r and cap(a)), key=lambda x: -cap(x[1]))
        tot = empty_agg()
        for x in allc:
            tot = add_agg(tot, x)
        # WPH · 정상 WPH 가 없는 레시피도 '없음' 행으로 남긴다(앱 레시피 비교와 같음)
        rows.append((r, cs, sum(cap(a) for _, a in cs), wph_normal(tot)))
    mx = max([1] + [x[2] for x in rows])
    parts.append('<h2>레시피별 하루 생산능력</h2><div class="sub">막대 한 칸 = 호기 1대 (24 × 실제 WPH). '
                 '<b>같은 호기는 어느 레시피에서도 같은 색</b>입니다 — 아래 범례.</div>')
    if rows:
        out = []
        for r, cs, total, wn_r in rows:
            segs = ''.join(f'<span style="width:{cap(a) / mx * 100:.2f}%;background:{col[m]}" '
                           f'title="{esc(f"{m} · {_step(r)}{chr(10)}하루 생산능력 {num(cap(a))}장{chr(10)}실제 WPH {num(wph_actual(a), 1)} · 정상 WPH {num(wph_normal(a), 1)}{chr(10)}Batch Report {a[chr(110)]:,}개")}">'
                           f'{esc(m) if cap(a) / mx * 100 > 6 else ""}</span>' for m, a in cs)
            segs = segs or '<span class="capnone">실제 WPH 자료 없음 (Pass한 Batch Report 없음)</span>'
            capv = f'<b>{num(total)}</b>장/일 · {len(cs)}대' if cs else '—'
            norm = '<small class="none">정상 WPH 없음</small>' if wn_r is None else f'<small>정상 WPH {num(wn_r, 1)}</small>'
            out.append(f'<div class="caprow"><span class="cl" title="{esc(r)}"><small>{esc(_job(r))}</small>{esc(_step(r))}</span>'
                       f'<span class="capbar">{segs}</span><span class="capv">{capv}{norm}</span></div>')
        used = sorted({m for _, cs, _, _ in rows for m, _ in cs})
        parts.append(''.join(out) + '<div class="keys">' + ''.join(f'<span><i style="background:{col[m]}"></i>{esc(m)}</span>' for m in used)
                     + '<span>같은 호기 = 같은 색</span></div>')
    else:
        parts.append('<p class="info">조사 범위에 WPH 자료가 없습니다.</p>')
    parts.append('<h2>호기 × 레시피 표</h2><div class="sub">레시피마다 합계 줄 + 호기별 줄. 앱의 [호기 × 레시피 표]와 같은 숫자입니다.</div>')
    body = []
    prev = None
    for r in recipes:
        cs = sorted((m, a) for (rr, m), a in cells.items() if rr == r)
        if not cs:
            continue
        if _job(r) != prev:
            body.append(f'<tr class="grouphead"><td colspan="9">{esc(_job(r))}</td></tr>')
            prev = _job(r)
        allc = empty_agg()
        for _, a in cs:
            allc = add_agg(allc, a)
        body.append(_wph_tr(_step(r), f'합계 ({len(cs)}대)', allc, sum(cap(a) or 0 for _, a in cs), True))
        body.extend(_wph_tr('', m, a, cap(a), False, col[m]) for m, a in cs)
    parts.append('<div class="tscroll"><table><thead><tr><th>레시피</th><th>호기</th><th class="num">1장 처리 시간</th><th class="num">Avg. Scan Time</th>'
                 '<th class="num">정상 WPH</th><th class="num">실제 WPH</th><th class="num">처리량 감소</th><th class="num">하루 생산능력</th>'
                 '<th class="num">Batch Report</th></tr></thead><tbody>'
                 + (''.join(body) or '<tr><td colspan="9">조사 범위에 WPH 자료가 없습니다.</td></tr>') + '</tbody></table></div>')
    return ''.join(parts)


def _wph_tr(step, m, a, total, is_total, color=''):
    n, w, u, sc = wph_normal(a), wph_actual(a), unit_sec(a), avg_scan(a)
    sw = f'<i class="msw" style="background:{color}"></i>' if color else ''
    return (f'<tr class="{"total" if is_total else ""}"><td>{esc(step)}</td><td>{sw}{esc(m)}</td>'
            f'<td class="num">{"—" if u is None else num(u) + "초"}</td><td class="num">{"—" if sc is None else num(sc) + "초"}</td>'
            f'<td class="num">{num(n, 1)}</td><td class="num"><b>{num(w, 1)}</b></td><td class="num">{pc(drop(n, w))}</td>'
            f'<td class="num"><b>{num(total)}</b></td><td class="num">{a["n"]:,}</td></tr>')


def lot_section(s, lot_file=''):
    parts = [_kpis([(l, f'{n:,}', u, sub, '', tone) for l, n, u, sub, tone in s.lot_kpis()], 'k5')]
    parts.append('<h2>멈춘 이유 요약 · Error</h2><div class="sub">wafer마다 Pass/Fail 원문의 첫 문구로 셉니다(앞 Error 뒤 따라온 Aborted. · Skipped.는 그 앞 Error 문구로). '
                 'Lot 판정: 그 Error 난 wafer가 모두 Pass면 재스캔하여 Pass, 하나라도 끝내 못 하면 Pass하지 못한 Lot.</div>')
    rows = ''.join(f'<tr><td>{esc(c)}</td><td class="num">{lots:,}</td><td class="num">{ok:,}</td><td class="num">{op:,}</td><td class="num">{w:,}</td></tr>'
                   for c, lots, ok, op, w in s.causes())
    parts.append('<div class="tscroll"><table><thead><tr><th>Error 원문(첫 문구)</th><th class="num">Lot</th><th class="num">재스캔하여 Pass한 Lot</th>'
                 '<th class="num">Pass하지 못한 Lot</th><th class="num">wafer</th></tr></thead><tbody>'
                 + (rows or '<tr><td colspan="5">Error 없음</td></tr>') + '</tbody></table></div>')
    parts.append('<h2>멈춘 이유 요약 · 작업자 중단</h2><div class="sub">앞에 Error 없이 Aborted.로 멈춘 Batch Report — Error로 세지 않습니다.</div>')
    rows = ''.join(f'<tr><td>{esc(k)}</td><td class="num">{n:,}</td><td class="num">{lots:,}</td><td class="num">{ok:,}</td><td class="num">{op:,}</td><td class="num">{w:,}</td></tr>'
                   for k, lots, ok, op, w, n in s.stops())
    parts.append('<div class="tscroll"><table><thead><tr><th>중단 종류</th><th class="num">Batch Report</th><th class="num">Lot</th><th class="num">재스캔하여 Pass한 Lot</th>'
                 '<th class="num">Pass하지 못한 Lot</th><th class="num">wafer</th></tr></thead><tbody>' + rows + '</tbody></table></div>')
    base = s.P['X']['faults']
    if base:
        parts.append('<details><summary>Defect 과다 판정 기준값(레시피별 정상 wafer Faults 상위 1%)</summary><table><thead><tr><th>레시피</th>'
                     '<th class="num">기준값(Faults)</th><th class="num">정상 wafer 수</th></tr></thead><tbody>'
                     + ''.join(f'<tr><td>{esc(r)}</td><td class="num">{"판정 안 함" if b is None else num(b, 1)}</td><td class="num">{n:,}</td></tr>'
                               for r, (b, n) in sorted(base.items())) + '</tbody></table></details>')
    if lot_file:
        parts.append(f'<div class="note navy"><span class="t">Lot별 상세</span>Lot 목록 · Lot History · wafer 오류 지도는 같은 폴더의 '
                     f'<b>{esc(lot_file)}</b>(Lot 추적 HTML)에서 봅니다.</div>')
    return ''.join(parts)


def defs_section():
    crit = ''.join(f'<dt>{esc(t)}</dt><dd>{esc(d)}</dd>' for t, d, _ in lotreport.CRITERIA)
    return ('<h2>지표 정의</h2><div class="sub">앱 화면의 ? 버튼과 같은 정의입니다(v12.0.0).</div><dl class="defs">'
            + ''.join(f'<dt>{esc(t)}</dt><dd>{esc(d)}</dd>' for t, d in DEFS) + '</dl>'
            '<h2>Lot 판정 기준</h2><dl class="defs">' + crit + '</dl>')


def build_html(s, dashboard=False, lot_file='', notices=()):
    """분석 HTML(dashboard=False: Lot 추적 · 가동률 · WPH · 정의) / 가동률 대시보드(가동률 · WPH, 30분 새로고침)."""
    title = '가동률 대시보드' if dashboard else 'Batch Report 분석'
    tabs = ([] if dashboard else [('lot', 'Lot 추적', lot_section(s, lot_file))]) + [
        ('util', '가동률', util_section(s)), ('wph', 'WPH · 생산능력', wph_section(s)), ('defs', '📘 지표 정의', defs_section())]
    nav = '<div class="metricnav">' + ''.join(
        f'<button class="mbtn{" on" if i == 0 else ""}{" help" if k == "defs" else ""}" data-sec="{k}" onclick="showMetric(\'{k}\')">{esc(t)}</button>'
        for i, (k, t, _) in enumerate(tabs)) + '</div>'
    secs = ''.join(f'<section class="metric" id="sec-{k}"{"" if i == 0 else " hidden"}>{body}</section>'
                   for i, (k, _, body) in enumerate(tabs))
    basis = '추천(가장 나중 Pass)' + (f' + 저장된 사람 선택 {s.saved}건' if s.saved else '')
    meta = (f'<span><b>기간</b> {esc(s.range[0] or "—")} ~ {esc(s.range[1] or "—")}</span>'
            f'<span><b>호기</b> {esc(", ".join(s.ids)) or "—"} ({len(s.ids)}대)</span>'
            f'<span><b>Lot</b> {len(s.P["lots"]):,} · <b>Batch Report</b> {len(s.P["R"]):,}</span>'
            f'<span><b>중복 Pass 선택 기준</b> {esc(basis)}</span><span><b>생성</b> {s.created:%Y-%m-%d %H:%M}</span>')
    notes = ''.join(f'<li>{esc(n)}</li>' for n in notices)
    return ''.join([
        '<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">',
        '<meta http-equiv="Content-Security-Policy" content="default-src &#39;none&#39;; style-src &#39;unsafe-inline&#39;; script-src &#39;unsafe-inline&#39;; img-src data:; connect-src &#39;none&#39;">',
        '<meta http-equiv="refresh" content="1800">' if dashboard else '',
        f'<title>{title}</title><style>{PAGE_CSS}{SAVED_CSS}</style></head><body><div class="doc">',
        '<div class="head"><div class="kick">Camtek AOI Manager · Batch Report 분석 · 앱 화면(v12)과 같은 계산</div>',
        f'<h1>{title}{"<span class=draft>30분 자동 새로고침</span>" if dashboard else ""}</h1><div class="meta">{meta}</div></div>',
        '<div class="pad">',
        f'<div class="info">조사 범위: {esc(s.scope or "—")}</div>',
        '<div class="note navy"><span class="t">자동 새로고침 안내</span>이 페이지는 로컬 파일입니다. 자동 새로고침은 파일만 다시 엽니다. '
        '새 데이터는 앱에서 조사를 다시 하거나 하루 1회 자동 분석으로 생성됩니다.</div>' if dashboard else '',
        nav, secs,
        f'<aside><h2>읽기 오류 · 알림</h2><ul>{notes or "<li>이번 실행에서는 읽기 오류 · 알림이 없습니다.</li>"}</ul></aside>',
        '<div class="foot">※ 숫자는 프로그램의 Batch Report 분석 화면(가동률 조사 및 분석)과 같은 엔진 · 같은 식으로 만든 것입니다. '
        '원본 Batch Report는 읽기만 했습니다. 화면에서 기간 · 호기를 좁혀 보면 그 범위의 숫자가 나옵니다(이 파일은 조사 범위 전체).</div>',
        f'</div></div><script>{TAB_JS}</script></body></html>'])


# ---------------------------------------------------------------------- Excel
def write_excel(path, s, notices=()):
    """분석 Excel — 앱 화면과 같은 숫자를 시트로(값만, 수식 없음)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    r1 = lambda v: None if v is None else round(v, 1)      # noqa: E731
    h1 = lambda x: round(x / 3600, 2)                      # noqa: E731

    def sheet(name, head, rows, widths=None):
        ws = wb.create_sheet(name)
        ws.append([f'{name} · 기간 {s.range[0]} ~ {s.range[1]} · 호기 {len(s.ids)}대 · 앱 화면(v12)과 같은 계산'])
        ws.append(head)
        for row in rows:
            ws.append(list(row))
        for cells in ws.iter_rows(min_row=3):
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = 's'           # 원문에서 온 글자는 수식으로 해석하지 않는다
        ws['A1'].font = Font(bold=True, color='183047')
        for cell in ws[2]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='183047')
        for i in range(1, len(head) + 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 16)
        ws.freeze_panes = 'A3'
        return ws

    a = s.util_all()
    info = [('기간', f'{s.range[0]} ~ {s.range[1]}'), ('호기', ', '.join(s.ids)), ('조사 범위', s.scope),
            ('Lot', len(s.P['lots'])), ('Batch Report', len(s.P['R'])), ('점검 스캔', len(s.P['excluded'])),
            ('중복 Pass 선택 기준', '추천(가장 나중 Pass)' + (f' + 저장된 사람 선택 {s.saved}건' if s.saved else '')),
            ('생성', s.created.strftime('%Y-%m-%d %H:%M')),
            ('가동률(%)', r1(pct(a['proc'], a['cal']))), ('Error · 중단 및 조치(%)', r1(pct(a['loss'], a['cal']))),
            ('유휴(%)', r1(pct(a['idle'], a['cal']))), ('달력 시간(h)', h1(a['cal']))]
    info += [(t, d) for t, d in DEFS] + [('알림', n) for n in notices]
    sheet('조사 정보 · 정의', ['항목', '값'], info, {1: 28, 2: 110})

    util_head = ['웨이퍼 처리(h)', '그중 다시 스캔(h)', '그중 점검(h)', 'Error · 중단 및 조치(h)', 'Error(h)', STOPK['d'] + '(h)',
                 STOPK['o'] + '(h)', '유휴(h)', '달력 시간(h)', '가동률(%)', 'Error · 중단 및 조치(%)', '유휴(%)']

    def util_vals(c):
        A = c['a']
        return [h1(c['proc']), h1(A['du']), h1(A['ck']), h1(c['loss']), h1(A['e']), h1(A['sd']), h1(A['so']), h1(c['idle']),
                h1(c['cal']), r1(pct(c['proc'], c['cal'])), r1(pct(c['loss'], c['cal'])), r1(pct(c['idle'], c['cal']))]
    sheet('가동률 호기별', ['호기'] + util_head, [[m] + util_vals(c) for m, c in s.util_machines()] + [['전체'] + util_vals(a)])
    for unit, name in (('d', '가동률 일별'), ('w', '가동률 주별'), ('m', '가동률 월별')):
        sheet(name, ['기간'] + util_head, [[period_name(k, unit)] + util_vals(c) for k, c in s.util_periods(unit)], {1: 26})
    sheet('가동률 날짜×호기', ['날짜', '호기'] + util_head, [[d, m] + util_vals(c) for d, m, c in s.util_days_machines()])
    sheet('손실 이유', ['이유(Error 원문 · 중단 종류)', '시간(h)', 'Lot'],
          [[n, h1(x), lots] for n, x, lots, _ in s.loss_reasons(a['a'])], {1: 60})

    recipes, cells = s.wph_cells()
    rows = []
    for r in recipes:
        cs = sorted((m, x) for (rr, m), x in cells.items() if rr == r)
        if not cs:
            continue
        allc = empty_agg()
        for _, x in cs:
            allc = add_agg(allc, x)
        for m, x, total in [(f'합계 ({len(cs)}대)', allc, sum(cap(y) or 0 for _, y in cs))] + [(m, x, cap(x)) for m, x in cs]:
            n, w = wph_normal(x), wph_actual(x)
            rows.append([_job(r), _step(r), m, r1(unit_sec(x)), r1(avg_scan(x)), r1(n), r1(w), r1(drop(n, w)),
                         None if total is None else round(total), x['n'], x['ps'], x['dn'], h1(x['p']), h1(loss_of(x))])
    sheet('WPH 호기×레시피', ['Job', 'Recipe(s)', '호기', '1장 처리 시간(초)', 'Avg. Scan Time(초)', '정상 WPH', '실제 WPH',
                            '처리량 감소(%)', '하루 생산능력(장)', 'Batch Report', 'Pass 장수', '그중 다시 스캔 Pass', '웨이퍼 처리(h)',
                            'Error · 중단 및 조치(h)'], rows, {1: 40, 2: 22})
    R, lots = s.P['R'], s.P['lots']
    sheet('정상 25매 Batch Report', ['시작', '호기', '레시피', 'Lot', 'Batch Time(초)', 'Avg. Scan Time(초)', 'WPH', '파일'],
          [[R[x['g']]['s'], x['m'], x['r'], lots[x['lot']]['label'], x['s'], x['a'], round(25 * 3600 / x['s'], 1) if x['s'] else None,
            R[x['g']]['f']] for x in s.P['N']], {3: 40, 8: 60})
    sheet('작업자 중단', ['시작', '호기', '레시피', 'Lot', '종류', '멈출 때 Faults', '기준값(상위 1%)', 'Pass 장수', '파일'],
          [[R[x['g']]['s'], x['m'], x['r'], lots[x['li']]['label'], STOPK[x['k']], x['f'], x['base'], x['ps'], R[x['g']]['f']]
           for x in s.P['stops']], {3: 40, 5: 24, 9: 60})
    sheet('Lot 목록', ['Lot', 'S/M', 'Lot ID', 'Job', '호기', '처음 스캔', '마지막 스캔', '상태', 'Batch Report', 'Pass 없는 wafer',
                      '재스캔 Pass wafer', '중복 Pass', '호기 이동', 'Error 원문(wafer)', '작업자 중단(wafer)'],
          [[l['label'], ', '.join(l['sms']), l['lot_id'], ', '.join(l['jobs']), ' → '.join(l['machines']), l['s'], l['e'], l['state'],
            l['n'], l['open'], l['re'], l['dup'], '예' if l['moved'] else '',
            ' · '.join(f'{c} ×{w}' for c, w, _ in l['cz']), ' · '.join(f'{STOPK[k]} ×{w}' for k, w, _ in l['sz'])] for l in lots],
          {2: 24, 4: 40, 14: 60, 15: 40})
    sheet('Batch Report 목록', ['시작', '끝', '호기', 'Lot', 'S/M', 'Job', 'Recipe(s)', '결과', '행', 'Pass', 'Error', '연쇄',
                               '작업자 중단 행', '스캔 안 한 슬롯', '첫 Error 원문', 'Batch Time(초)', '파일'],
          [[r['s'], r['e'], r['m'], '' if r['lot'] is None else lots[r['lot']]['label'], r['sm'], r['job'], lm.step_label(r['step'], r['k']),
            {'c': '모두 Pass', 'e': 'Error 포함', 's': STOPK.get(r['sk'], '작업자 중단')}[r['o']] if r['lot'] is not None else '점검 스캔',
            r['n'], r['ok'], r['err'], r['chain'], r['st'], r['sp'], r['fe'], r['sec'], r['f']] for r in R],
          {6: 40, 15: 40, 17: 60})
    wb.save(path)


def payload_json(s):
    """테스트 · 디버그용 — 저장 파일이 쓰는 숫자 요약."""
    a = s.util_all()
    return json.dumps(dict(util=[pct(a['proc'], a['cal']), pct(a['loss'], a['cal']), pct(a['idle'], a['cal'])]), ensure_ascii=False)
