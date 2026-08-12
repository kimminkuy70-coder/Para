"""변환계수 저장소 — '변환계수.xlsx'(저장폴더 독립 파일).

변환계수는 **장비 렌즈 특성**이라 장비마다 다르고, **호기별·변형별**로 관리한다
(사용자 확정 2026-08). → 키 = (호기, 변형). 값 = 변환계수.
변형(PI/PI-bubble/x5/x20)이 곧 레시피 레벨의 배율이므로, 같은 호기·같은 변형이면
레시피가 달라도 같은 계수를 쓴다. MAG 열(OpticPreset Scan2d 의 Mag)은 **사람이
어떤 렌즈인지 알아보기 위한 참고열**일 뿐 매칭에 쓰지 않는다 — target optic 이
바뀌면 MAG 이 흔들려 값 업데이트마다 계수가 달라지던 문제 때문에 폐기했다.

  헤더: [호기, MAG, 변형, 계수, 비고]

파일은 장비 IP 주소.xlsx 처럼 **오직 사람이 관리**한다. 값을 읽는 작업(값 업데이트·
commonality)은 조회만 하고 절대 고치지 않는다. 파일을 쓰는 건 사람이 양식에서 계수를
확정할 때뿐(`apply_form_scales`).
"""

from __future__ import annotations

import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from . import engine

COEF_FILENAME = "변환계수.xlsx"
COEF_SHEET = "변환계수"
COEF_HEADERS = ["호기", "MAG", "변형", "계수", "비고"]
_HDR_FILL = "1F4E78"


def coef_path(save_dir: str) -> str:
    return os.path.join(save_dir, COEF_FILENAME)


def _machine_norm(s) -> str:
    """호기 매칭 정규화 — 대소문자/구분자 무시(AOI-K2 == aoi_k2). 숫자 0패딩도 흡수."""
    import re
    txt = str(s or "").lower()
    letters = re.sub(r"[^a-z]", "", txt)
    nums = re.findall(r"\d+", txt)
    return letters + "".join(str(int(n)) for n in nums)


def mag_norm(mag) -> str:
    """MAG 매칭 정규화 — 숫자면 소수 4자리 반올림 문자열, 아니면 소문자 압축."""
    s = engine._s(mag).strip()
    if s == "":
        return ""
    try:
        return f"{round(float(s), 4):g}"
    except ValueError:
        import re
        return re.sub(r"[^0-9a-z]", "", s.lower())


def _style_header(ws) -> None:
    fill = PatternFill("solid", fgColor=_HDR_FILL)
    white = Font(color="FFFFFF", bold=True)
    for c in ws[1]:
        c.fill = fill
        c.font = white
        c.alignment = Alignment(horizontal="center", vertical="center")


def create_blank(path: str) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = COEF_SHEET
    ws.append(COEF_HEADERS)
    _style_header(ws)
    for col, w in zip("ABCDE", (12, 12, 14, 22, 20)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return path


def load(path: str) -> list[dict]:
    """변환계수.xlsx → [{호기, MAG, 변형, 계수, 비고}] (헤더 변동 허용)."""
    if not os.path.isfile(path):
        return []
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[COEF_SHEET] if COEF_SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
    heads = [engine._s(c.value).strip() for c in ws[1]]
    hidx = {h: i for i, h in enumerate(heads) if h}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(v in (None, "") for v in row):
            continue

        def g(h, default=""):
            i = hidx.get(h)
            return engine._s(row[i]).strip() if i is not None and i < len(row) else default
        ho = g("호기")
        if not ho:
            continue
        out.append({"호기": ho, "MAG": g("MAG"), "변형": g("변형"),
                    "계수": g("계수"), "비고": g("비고")})
    wb.close()
    return out


def save(path: str, rows: list[dict]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = COEF_SHEET
    ws.append(COEF_HEADERS)
    for r in rows:
        ws.append([engine._s(r.get(h)) for h in COEF_HEADERS])
    _style_header(ws)
    for col, w in zip("ABCDE", (12, 12, 14, 22, 20)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def _coef_float(v):
    try:
        return float(engine._s(v).strip())
    except ValueError:
        return None


def _variant_norm(s) -> str:
    """변형 매칭 정규화 — 대소문자·구분자(공백/-/_) 무시(PI-bubble == pi_bubble).
    한글·기타 문자는 보존(폴더 라벨을 변형으로 쓰는 경우 대비)."""
    import re
    return re.sub(r"[\s\-_]+", "", str(s or "").strip().lower())


def lookup(rows: list[dict], machine: str, variant) -> float | None:
    """(호기, 변형) → 계수(float). 변형은 대소문자·구분자 무시. 없으면 None.
    변형이 빈 행은 그 호기의 **공통 폴백**(변형 지정이 없을 때만)."""
    mk, vk = _machine_norm(machine), _variant_norm(variant)
    if not mk:
        return None
    fallback = None
    for r in rows:
        if _machine_norm(r.get("호기")) != mk:
            continue
        c = _coef_float(r.get("계수"))
        if c is None:
            continue
        rv = _variant_norm(r.get("변형"))
        if rv == vk:
            return c
        if rv == "" and fallback is None:
            fallback = c                  # 변형 없는 행 = 호기 공통(폴백)
    return fallback


def machine_coefs(rows: list[dict], machine: str) -> list[dict]:
    """한 호기의 계수 행들(표시용). 변형/ MAG 순."""
    mk = _machine_norm(machine)
    out = [r for r in rows if _machine_norm(r.get("호기")) == mk]
    return sorted(out, key=lambda r: (_variant_norm(r.get("변형")), mag_norm(r.get("MAG"))))


def upsert(rows: list[dict], machine: str, variant, coef,
           overwrite: bool = False, note: str = "자동추정", mag: str = "") -> bool:
    """(호기, 변형) 행 추가/갱신. overwrite=False 면 이미 있으면 계수 유지(사람 우선).
    mag 는 참고열(있으면 함께 기록/갱신). note 는 비고에 남길 출처.
    반환: 실제로 추가/갱신했으면 True."""
    if coef is None or engine._s(coef).strip() == "":
        return False
    mk, vk = _machine_norm(machine), _variant_norm(variant)
    for r in rows:
        if _machine_norm(r.get("호기")) == mk and _variant_norm(r.get("변형")) == vk:
            if overwrite:
                r["계수"] = engine._s(coef)
                if mag:
                    r["MAG"] = engine._s(mag)
                if note:
                    r["비고"] = note
                return True
            return False                  # 이미 있음(사람 값 우선)
    rows.append({"호기": engine._s(machine), "MAG": engine._s(mag),
                 "변형": engine._s(variant), "계수": engine._s(coef), "비고": note})
    return True


def _same_coef(a, b) -> bool:
    fa, fb = _coef_float(a), _coef_float(b)
    if fa is None or fb is None:
        return engine._s(a).strip() == engine._s(b).strip()
    return abs(fa - fb) <= 1e-12


def _variant_row(rows: list[dict], machine: str, variant: str) -> dict | None:
    """(호기, 변형)이 정확히 일치하는 기존 행(첫째). 없으면 None."""
    mk, vk = _machine_norm(machine), _variant_norm(variant)
    for r in rows:
        if _machine_norm(r.get("호기")) == mk and _variant_norm(r.get("변형")) == vk:
            return r
    return None


def apply_form_scales(rows: list[dict], machine: str, scales: dict,
                      mags: dict | None = None, note: str = "양식 확정") -> int:
    """**양식에서 사람이 확정한** 변형별 계수를 저장소에 반영한다.

    양식 편집기에서 계수를 고쳐 확정하면 그 값이 정답이므로 **덮어쓴다**
    (기존 값을 보존하지 않는다). 키 = (호기, 변형).

    mags = {변형: MAG} — 알면 참고열에 함께 기록한다(매칭에는 쓰지 않는다).
    **새 변형 행은 MAG(실제 optic 값)를 알 때만 만든다** — MAG 도 없고 기존 행도
    없으면 엉뚱한 행을 만들지 않는다(오타 변형 방지).
    반환: 실제로 값이 바뀐(또는 추가된) 행 수 — 0 이면 저장할 필요가 없다.
    """
    if not engine._s(machine).strip() or not scales:
        return 0
    mags = mags or {}
    n = 0
    for variant, coef in scales.items():
        if _coef_float(coef) is None:
            continue
        v = engine._s(variant).strip()
        mag = engine._s(mags.get(v, mags.get(variant, ""))).strip()
        existing = _variant_row(rows, machine, v)
        if existing is not None:
            if _same_coef(existing.get("계수"), coef):
                continue                  # 값이 그대로면 파일을 다시 쓰지 않는다
            existing["계수"] = engine._s(coef)
            if mag:
                existing["MAG"] = engine._s(mag)
            if note:
                existing["비고"] = note
            n += 1
        elif mag:
            rows.append({"호기": engine._s(machine), "MAG": engine._s(mag),
                         "변형": v, "계수": engine._s(coef), "비고": note})
            n += 1
    return n


def make_lookup(rows: list[dict]):
    """scan_tree/collate(coef_lookup=...) 공용 콜백. (machine, variant)->계수|None."""
    def _cb(machine, variant=""):
        return lookup(rows, machine, variant)
    return _cb
