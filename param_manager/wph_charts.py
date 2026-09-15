"""Native WPH charts over numeric snapshots, independent of Excel recalculation."""
from collections import defaultdict
import math


def summarize(rows, valid_wafers):
    machines = defaultdict(lambda: [0, 0, 0.0])
    reports = []
    bands = [0] * 6
    for index, row in enumerate(rows, 1):
        try:
            wafers, scan, batch = (float(row.get(k)) for k in ('wafers', 'avg_scan_sec', 'batch_sec'))
        except (ValueError, TypeError):
            continue
        if not all(math.isfinite(x) for x in (wafers, scan, batch)) or wafers != valid_wafers or scan <= 0 or batch <= 0:
            continue
        machine = str(row.get('machine') or '(호기 미지정)')
        item = machines[machine]
        item[0] += 1
        item[1] += wafers
        item[2] += batch
        reports.append([index, wafers * 3600 / batch, row.get('source_file', ''), machine])
        bands[min(int(scan // 30), 5)] += 1
    return machines, reports, bands


def add_charts(wb, rows, valid_wafers):
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Font
    data, dashboard = wb['04_그래프데이터'], wb['05_대시보드']
    machines, reports, bands = summarize(rows, valid_wafers)
    blocks = [
        (13, ['호기', '누적 WPH'], [[m, v[1] * 3600 / v[2]] for m, v in machines.items()]),
        (16, ['원본 Report 번호', 'Actual WPH', 'Report 파일명', '호기'], reports),
        (21, ['Avg Scan 시간 구간', 'Report 수'], list(zip(
            ['0~30초 미만', '30~60초 미만', '60~90초 미만', '90~120초 미만', '120~150초 미만', '150초 이상'], bands))),
    ]
    for column, headers, values in blocks:
        for offset, header in enumerate(headers):
            data.cell(1, column + offset, header)
        for r, values_row in enumerate(values or [['유효 데이터 없음', 0]], 2):
            for offset, value in enumerate(values_row):
                cell = data.cell(r, column + offset, value)
                if isinstance(value, str):
                    cell.data_type = 's'
    dashboard['A9'] = '아래 차트는 조사 시점의 유효 Lot 집계입니다. Raw Data 수동 수정 시 다시 조사하세요.'
    dashboard['A9'].font = Font(color='666666')
    for column, count, title, anchor, line in [
        (13, len(machines), '호기별 누적 WPH (총 Wafer / 총 Batch 시간)', 'A11', False),
        (16, len(reports), 'Report별 Actual WPH (원본 목록 순서)', 'K11', True),
        (21, 6, 'Avg Scan 시간 분포 — Report 수', 'A30', False),
    ]:
        chart = LineChart() if line else BarChart()
        chart.title, chart.style, chart.legend = title, 10, None
        chart.width, chart.height = 21, 9
        chart.add_data(Reference(data, min_col=column + 1, min_row=1, max_row=max(1, count) + 1), titles_from_data=True)
        chart.set_categories(Reference(data, min_col=column, min_row=2, max_row=max(1, count) + 1))
        chart.y_axis.title = 'Report 수' if column == 21 else 'WPH'
        if not line:
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True
            chart.dataLabels.numFmt = '0' if column == 21 else '0.00'
        dashboard.add_chart(chart, anchor)
    wb.active = wb.index(dashboard)
