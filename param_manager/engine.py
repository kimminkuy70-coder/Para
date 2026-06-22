"""핵심 엔진 — 화면(GUI)과 분리된 순수 로직 계층.

기존 VBA(ThisWorkbook) 자동화를 파이썬으로 이식한 모듈입니다. tkinter 없이도
단독으로 동작/테스트할 수 있도록 GUI 의존성을 두지 않습니다.

담당 기능
  1) PI_ALL 시트 read/write (openpyxl, .xlsx)
  2) Row_ID 자동 부여 / Display_Order 관리
  3) 변경 감지(편집 전 스냅샷 대비) 및 변경이력 자동 누적
  4) 호기별 요약(변경요약_호기별) 재생성 + 최초값/이전값/변경횟수 보존
  5) 호기 간 값 비교 / 누락 호기 계산
  6) 동시 편집 대비: 저장 직전 디스크 최신본을 다시 읽어 Row_ID 기준 행 병합
  7) 편집 잠금 파일(.editlock) 및 OneDrive 충돌본 감지

실행 중 네트워크 접속은 전혀 없으며, OneDrive 동기화 폴더의 로컬 파일만 다룹니다.
"""

from __future__ import annotations

import getpass
import json
import os
import random
import shutil
import socket
import tempfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Any

import openpyxl

# --------------------------------------------------------------------------
# 스키마 상수 — 기존 엑셀 파일과 호환되도록 정의
# --------------------------------------------------------------------------

SHEET_PI = "PI_ALL"
SHEET_SUM = "변경요약_호기별"
SHEET_LOG = "변경이력"
SHEET_SNAP = "_SNAPSHOT_PI_ALL"
SHEET_SPECIAL = "특이사항"
SHEET_REF = "참고자료"
SHEET_RELATED = "관련 자료"
REF_COLS = 4   # 참고자료 그리드 열 수

# 기본 관리 호기(새 빈 파일/감지 실패 시 사용). 실제 호기 목록은 파일에서 자동 인식.
AOI_UNITS = [
    "AOI-3", "AOI-13", "AOI-15", "AOI-16", "AOI-17", "AOI-18", "AOI-19",
    "AOI-20", "AOI-21", "AOI-22", "AOI-23", "AOI-24", "AOI-25",
]

# 데이터 시트 후보 이름(PI / RDL 등 동일 양식)
SHEET_CANDIDATES = ["PI_ALL", "RDL_ALL"]

# 메타(상위) 항목 — A:G
META_FIELDS = ["PI", "Recipe", "Zone", "Alg", "Parameter", "초기 추천값", "비고"]

# 데이터 시트 뒤쪽 파생/필터 열(호기 열 자동 인식 시 제외 대상)
DERIVED_TAIL = [
    "공통 여부", "값 차이 여부", "누락 호기",
    "Recipe_Filter", "Zone_Filter", "Alg_Filter",
    "SourceSheet", "SourceRow", "Parameter_Key",
    "Row_ID", "Display_Order",
]


def editable_fields(aoi_units) -> list:
    return list(META_FIELDS) + list(aoi_units)


def pi_headers(aoi_units) -> list:
    return list(META_FIELDS) + list(aoi_units) + list(DERIVED_TAIL)


def snap_headers(aoi_units) -> list:
    return ["Row_ID"] + list(META_FIELDS) + list(aoi_units)


def detect_sheet_name(wb) -> str | None:
    """데이터 시트 이름 자동 인식: PI_ALL/RDL_ALL 우선, 없으면 헤더로 추정."""
    for name in SHEET_CANDIDATES:
        if name in wb.sheetnames:
            return name
    for ws in wb.worksheets:
        hdr = {_s(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)}
        if "Parameter" in hdr and "PI" in hdr:
            return ws.title
    return None


def detect_aoi_units(wb, sheet_name) -> list:
    """헤더에서 호기 열을 자동 인식(메타/파생 열을 제외한 나머지, 순서 유지)."""
    if not sheet_name or sheet_name not in wb.sheetnames:
        return list(AOI_UNITS)
    ws = wb[sheet_name]
    meta = set(META_FIELDS)
    derived = set(DERIVED_TAIL)
    seen = set()
    units = []
    for c in range(1, ws.max_column + 1):
        name = _s(ws.cell(1, c).value)
        if not name or name in meta or name in derived or name in seen:
            continue
        seen.add(name)
        units.append(name)
    return units or list(AOI_UNITS)


# 화면에서 편집 가능한 전체 필드(기본)
EDITABLE_FIELDS = META_FIELDS + AOI_UNITS

# PI_ALL 전체 헤더(기본)
PI_HEADERS = pi_headers(AOI_UNITS)

LOG_HEADERS = [
    "Version", "Update Date", "Updated By", "Sheet", "Recipe", "Zone", "Alg",
    "Parameter", "AOI", "Old Value", "New Value", "Change Type", "Reason",
    "Source", "Comment", "비고",
]

SUM_HEADERS = [
    "PI", "AOI", "Recipe", "Zone", "Alg", "Parameter", "초기 추천값", "비고",
    "현재값", "최초값", "이전값", "변경횟수", "최초 변경일시", "마지막 변경일시",
    "마지막 변경자", "마지막 변경유형", "Source", "Comment", "Summary_Key",
    "Source_Cell", "Row_ID", "Stable_Key", "Display_Order", "Status",
    "PI_Order", "AOI_Order",
]

SNAP_HEADERS = snap_headers(AOI_UNITS)

# 특이사항 시트 — 종료 여부는 불리언(체크박스)
SPECIAL_HEADERS = ["일자", "호기", "라트 번호", "S/M", "Layer", "목적", "진행 상황", "종료 여부", "특이사항"]
SPECIAL_BOOL_COL = "종료 여부"

# 빈 칸 취급 값 (호기 값이 비었다고 볼 토큰)
_EMPTY_TOKENS = {"", "-", "—", "–"}

LOCK_STALE_MINUTES = 30  # 이 시간이 지난 잠금은 만료로 간주


def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def today_str() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def format_kdate(value: Any) -> tuple[str, bool]:
    """일자 값을 'YYYY년 M월 D일' 형식으로 변환.

    허용 입력: 8자리(YYYYMMDD), YYYY-MM-DD / YYYY/MM/DD / YYYY.MM.DD,
    'YYYY년 M월 D일', datetime/date 객체.
    반환: (표시문자열, 정상여부). 빈 값은 ("", True), 형식 오류는 ("일자 오류", False).
    """
    import re
    from datetime import date as _date
    from datetime import datetime as _dt

    if value is None:
        return "", True
    if isinstance(value, (_dt, _date)):
        return f"{value.year}년 {value.month}월 {value.day}일", True
    s = str(value).strip()
    if s == "":
        return "", True
    candidates: list[tuple[int, int, int]] = []
    km = re.match(r"^(\d{4})\s*년\s*(\d{1,2})\s*월\s*(\d{1,2})\s*일?$", s)
    if km:
        candidates.append((int(km.group(1)), int(km.group(2)), int(km.group(3))))
    sm = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$", s)
    if sm:
        candidates.append((int(sm.group(1)), int(sm.group(2)), int(sm.group(3))))
    if re.fullmatch(r"\d{8}", s):
        candidates.append((int(s[:4]), int(s[4:6]), int(s[6:8])))
    for y, mo, d in candidates:
        try:
            dt = _date(y, mo, d)
            return f"{dt.year}년 {dt.month}월 {dt.day}일", True
        except ValueError:
            continue
    return "일자 오류", False


def current_user() -> str:
    """Windows/OS 로그인 계정명 (VBA Environ("USERNAME") 대체)."""
    for key in ("USERNAME", "USER", "LOGNAME"):
        v = os.environ.get(key)
        if v:
            return v
    try:
        return getpass.getuser()
    except Exception:
        return "unknown"


def new_row_id() -> str:
    return "RID_" + datetime.now().strftime("%Y%m%d%H%M%S") + "_" + f"{random.randint(0, 99999):05d}"


def _s(v: Any) -> str:
    """비교/표시용 정규화 문자열."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def is_blank(v: Any) -> bool:
    return _s(v) in _EMPTY_TOKENS


def pi_order(pi_val: str) -> int:
    digits = "".join(ch for ch in _s(pi_val) if ch.isdigit())
    return int(digits) if digits else 999


def aoi_order(aoi: str, aoi_units=None) -> int:
    units = aoi_units if aoi_units is not None else AOI_UNITS
    try:
        return list(units).index(aoi) + 1
    except ValueError:
        return 999


# --------------------------------------------------------------------------
# 데이터 모델
# --------------------------------------------------------------------------

@dataclass
class ParamRow:
    """데이터 시트 한 행 = 하나의 파라미터(메타 + 호기 값들)."""

    values: dict[str, Any] = field(default_factory=dict)  # 필드명 -> 값
    row_id: str = ""
    display_order: int = 0
    aoi_units: list = field(default_factory=lambda: list(AOI_UNITS))  # 이 행이 쓰는 호기 목록

    def get(self, field_name: str) -> Any:
        return self.values.get(field_name)

    def set(self, field_name: str, value: Any) -> None:
        self.values[field_name] = value

    def snapshot(self) -> dict[str, str]:
        """편집 가능한 필드만 비교용 문자열 dict 로."""
        return {f: _s(self.values.get(f)) for f in editable_fields(self.aoi_units)}

    def parameter_key(self) -> str:
        return "|".join(_s(self.values.get(f)) for f in ("PI", "Recipe", "Zone", "Alg", "Parameter"))

    def missing_units(self) -> list[str]:
        return [u for u in self.aoi_units if is_blank(self.values.get(u))]

    def filled_values(self) -> list[Any]:
        return [self.values.get(u) for u in self.aoi_units if not is_blank(self.values.get(u))]

    def has_value_diff(self) -> bool:
        vals = {_s(v) for v in self.filled_values()}
        return len(vals) > 1

    def is_common(self) -> bool:
        """모든 호기에 값이 채워져 있고 값이 동일하면 공통."""
        if self.missing_units():
            return False
        return not self.has_value_diff()


@dataclass
class ChangeRecord:
    row_id: str
    field: str          # 변경된 필드명 (메타명 또는 호기명)
    old: Any
    new: Any
    is_aoi: bool


# --------------------------------------------------------------------------
# 잠금(소프트 락) — 동시 편집 충돌 완화
# --------------------------------------------------------------------------

@dataclass
class LockInfo:
    user: str
    host: str
    time: str
    pid: int

    @property
    def datetime(self) -> datetime | None:
        try:
            return datetime.strptime(self.time, "%Y-%m-%d %H:%M:%S")
        except Exception:
            return None

    def is_stale(self) -> bool:
        dt = self.datetime
        if dt is None:
            return True
        return datetime.now() - dt > timedelta(minutes=LOCK_STALE_MINUTES)


def lock_path_for(xlsx_path: str) -> str:
    return xlsx_path + ".editlock"


def read_lock(xlsx_path: str) -> LockInfo | None:
    p = lock_path_for(xlsx_path)
    if not os.path.exists(p):
        return None
    try:
        with open(p, encoding="utf-8") as fh:
            d = json.load(fh)
        return LockInfo(d.get("user", "?"), d.get("host", "?"), d.get("time", ""), int(d.get("pid", 0)))
    except Exception:
        return None


def write_lock(xlsx_path: str, user: str) -> None:
    p = lock_path_for(xlsx_path)
    payload = {"user": user, "host": socket.gethostname(), "time": now_str(), "pid": os.getpid()}
    with open(p, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False)


def release_lock(xlsx_path: str, user: str) -> None:
    """내가 잡은 잠금만 해제."""
    info = read_lock(xlsx_path)
    if info and info.user == user and info.pid == os.getpid():
        try:
            os.remove(lock_path_for(xlsx_path))
        except OSError:
            pass


def find_conflict_copies(xlsx_path: str) -> list[str]:
    """같은 폴더에서 OneDrive 충돌본으로 의심되는 파일 탐지."""
    folder = os.path.dirname(os.path.abspath(xlsx_path))
    stem, ext = os.path.splitext(os.path.basename(xlsx_path))
    out: list[str] = []
    if not os.path.isdir(folder):
        return out
    for name in os.listdir(folder):
        if name == os.path.basename(xlsx_path):
            continue
        if not name.lower().endswith(ext.lower()):
            continue
        # OneDrive 충돌본: "<stem>-<기기명>.xlsx" / "...conflicted copy..." 형태
        low = name.lower()
        if name.startswith(stem + "-") or "conflict" in low or "충돌" in name or "-사본" in name:
            out.append(os.path.join(folder, name))
    return out


# --------------------------------------------------------------------------
# 저장소(Repository) — 엑셀 한 파일을 다루는 핵심 클래스
# --------------------------------------------------------------------------

class ParamRepository:
    """OneDrive 동기화 폴더의 공용 엑셀 파일 1개를 다루는 저장소."""

    def __init__(self, path: str):
        self.path = path
        self.rows: list[ParamRow] = []
        self.history: list[dict[str, Any]] = []
        self.aoi_ip: dict[str, str] = {}        # 호기 -> IP (관련 자료 시트)
        self.special: list[dict[str, Any]] = []  # 특이사항 시트 행들
        self.reference: list[list[str]] = []     # 참고자료 그리드
        self.sheet_name: str = SHEET_PI          # 데이터 시트 이름(자동 인식)
        self.aoi_units: list = list(AOI_UNITS)   # 호기 목록(자동 인식)
        self._base: dict[str, dict[str, str]] = {}  # row_id -> 편집 전 스냅샷

    # ---- 읽기 --------------------------------------------------------------

    def load(self) -> None:
        wb = openpyxl.load_workbook(self.path, data_only=True)
        self.sheet_name = detect_sheet_name(wb) or SHEET_PI
        self.aoi_units = detect_aoi_units(wb, self.sheet_name)
        self.rows = self._read_pi_all(wb, self.sheet_name, self.aoi_units)
        self.history = self._read_history(wb)
        self.aoi_ip = self._read_aoi_ip(wb)
        self.special = self._read_special(wb)
        self.reference = self._read_reference(wb)
        wb.close()
        self._ensure_row_ids()
        self._capture_base()

    @staticmethod
    def _read_pi_all(wb, sheet_name=None, aoi_units=None) -> list[ParamRow]:
        sheet_name = sheet_name or detect_sheet_name(wb) or SHEET_PI
        if sheet_name not in wb.sheetnames:
            return []
        if aoi_units is None:
            aoi_units = detect_aoi_units(wb, sheet_name)
        ws = wb[sheet_name]
        # 헤더 인덱스 매핑(헤더명 기준; 열 위치가 달라도 동작)
        header = {}
        for c in range(1, ws.max_column + 1):
            name = _s(ws.cell(1, c).value)
            if name:
                header.setdefault(name, c)
        fields = editable_fields(aoi_units)
        rows: list[ParamRow] = []
        for r in range(2, ws.max_row + 1):
            # Parameter 가 비면 데이터 행 아님
            pcol = header.get("Parameter", 5)
            if is_blank(ws.cell(r, pcol).value):
                continue
            pr = ParamRow(aoi_units=list(aoi_units))
            for fname in fields:
                col = header.get(fname)
                pr.values[fname] = ws.cell(r, col).value if col else None
            rid_col = header.get("Row_ID")
            pr.row_id = _s(ws.cell(r, rid_col).value) if rid_col else ""
            ord_col = header.get("Display_Order")
            try:
                pr.display_order = int(ws.cell(r, ord_col).value) if ord_col and ws.cell(r, ord_col).value else r
            except (TypeError, ValueError):
                pr.display_order = r
            rows.append(pr)
        return rows

    @staticmethod
    def _read_history(wb) -> list[dict[str, Any]]:
        if SHEET_LOG not in wb.sheetnames:
            return []
        ws = wb[SHEET_LOG]
        out: list[dict[str, Any]] = []
        for r in range(2, ws.max_row + 1):
            if is_blank(ws.cell(r, 1).value):
                continue
            out.append({h: ws.cell(r, i + 1).value for i, h in enumerate(LOG_HEADERS)})
        return out

    @staticmethod
    def _read_aoi_ip(wb) -> dict[str, str]:
        out: dict[str, str] = {}
        if "관련 자료" not in wb.sheetnames:
            return out
        ws = wb["관련 자료"]
        for r in range(1, ws.max_row + 1):
            a = _s(ws.cell(r, 1).value)
            b = _s(ws.cell(r, 2).value)
            if a.startswith("AOI-") and b:
                out[a] = b
        return out

    @staticmethod
    def _read_reference(wb) -> list[list[str]]:
        """참고자료 그리드를 읽는다. 없으면 기존 '관련 자료'에서 4개 표로 구성."""
        if SHEET_REF in wb.sheetnames:
            ws = wb[SHEET_REF]
            grid: list[list[str]] = []
            for r in range(1, ws.max_row + 1):
                row = [_s(ws.cell(r, c + 1).value) for c in range(REF_COLS)]
                grid.append(row)
            # 끝쪽 빈 행 정리
            while grid and all(v == "" for v in grid[-1]):
                grid.pop()
            return grid
        return ParamRepository._build_reference_from_related(wb)

    @staticmethod
    def _build_reference_from_related(wb) -> list[list[str]]:
        """기존 '관련 자료' 시트를 4개 표(블록)로 분해해 세로로 쌓는다.
        각 표 앞에 비어있는 '제목' 행을 둔다(사용자가 표 이름을 채움)."""
        grid: list[list[str]] = []
        if SHEET_RELATED not in wb.sheetnames:
            # 관련 자료가 없으면 빈 4개 표 골격만
            for i in range(1, 5):
                grid.append([f"표{i} 제목(여기에 입력)", "", "", ""])
                grid.append(["항목", "값", "", ""])
                grid.append(["", "", "", ""])
            return grid
        ws = wb[SHEET_RELATED]
        # 블록 정의: (제목 플레이스홀더, [열 인덱스(1-based)])
        blocks = [
            ("표 제목(여기에 입력)", [1, 2]),       # 장비번호 / IP주소
            ("표 제목(여기에 입력)", [4, 5]),       # 파트장 / 번호
            ("표 제목(여기에 입력)", [7, 8]),       # 공정 / PM 계정
            ("표 제목(여기에 입력)", [10, 11, 12]),  # 기타 매핑
        ]
        for title, cols in blocks:
            grid.append([title] + [""] * (REF_COLS - 1))
            for r in range(1, ws.max_row + 1):
                vals = [_s(ws.cell(r, c).value) for c in cols]
                if all(v == "" for v in vals):
                    continue
                if vals and "Gen5" in vals[0]:   # 시트 상단 제목 줄 스킵
                    continue
                row = vals + [""] * (REF_COLS - len(vals))
                grid.append(row[:REF_COLS])
            grid.append([""] * REF_COLS)
        return grid

    def _write_reference(self, wb) -> None:
        ws = wb.create_sheet(SHEET_REF)
        for row in self.reference:
            ws.append((list(row) + [""] * REF_COLS)[:REF_COLS])

    @staticmethod
    def _read_special(wb) -> list[dict[str, Any]]:
        if SHEET_SPECIAL not in wb.sheetnames:
            return []
        ws = wb[SHEET_SPECIAL]
        out: list[dict[str, Any]] = []
        bool_idx = SPECIAL_HEADERS.index(SPECIAL_BOOL_COL)
        for r in range(2, ws.max_row + 1):
            vals = [ws.cell(r, i + 1).value for i in range(len(SPECIAL_HEADERS))]
            # 종료여부(불리언)를 뺀 나머지가 전부 비면 빈 행으로 간주
            content = [v for i, v in enumerate(vals) if i != bool_idx]
            if all(v is None or _s(v) == "" for v in content):
                continue
            rec = {h: vals[i] for i, h in enumerate(SPECIAL_HEADERS)}
            # 종료 여부는 항상 불리언으로 정규화
            rec[SPECIAL_BOOL_COL] = bool(rec.get(SPECIAL_BOOL_COL)) if rec.get(SPECIAL_BOOL_COL) not in (None, "") else False
            out.append(rec)
        return out

    def _write_special(self, wb) -> None:
        ws = wb.create_sheet(SHEET_SPECIAL)
        ws.append(SPECIAL_HEADERS)
        for rec in self.special:
            row = []
            for h in SPECIAL_HEADERS:
                v = rec.get(h)
                if h == SPECIAL_BOOL_COL:
                    v = bool(v)
                row.append(v)
            ws.append(row)

    def _ensure_row_ids(self) -> None:
        for i, pr in enumerate(self.rows, start=2):
            if not pr.row_id:
                pr.row_id = new_row_id()
            if not pr.display_order:
                pr.display_order = i

    def _capture_base(self) -> None:
        self._base = {pr.row_id: pr.snapshot() for pr in self.rows}

    # ---- 변경 감지 ---------------------------------------------------------

    def detect_changes(self) -> list[ChangeRecord]:
        """현재 self.rows 와 편집 전 스냅샷(self._base)을 비교."""
        aoi_set = set(self.aoi_units)
        fields = editable_fields(self.aoi_units)
        changes: list[ChangeRecord] = []
        for pr in self.rows:
            base = self._base.get(pr.row_id)
            cur = pr.snapshot()
            if base is None:
                # 신규 행: 비어있지 않은 필드를 추가로 기록
                for f in fields:
                    if cur.get(f):
                        changes.append(ChangeRecord(pr.row_id, f, "", pr.values.get(f), f in aoi_set))
                continue
            for f in fields:
                if base.get(f, "") != cur.get(f, ""):
                    changes.append(ChangeRecord(pr.row_id, f, base.get(f, ""), pr.values.get(f), f in aoi_set))
        return changes

    def deleted_row_ids(self) -> list[str]:
        live = {pr.row_id for pr in self.rows}
        return [rid for rid in self._base if rid not in live]

    # ---- 비교/누락 뷰 -------------------------------------------------------

    def comparison_view(self) -> list[dict[str, Any]]:
        out = []
        for pr in self.rows:
            out.append({
                "row_id": pr.row_id,
                "PI": pr.get("PI"),
                "Recipe": pr.get("Recipe"),
                "Zone": pr.get("Zone"),
                "Alg": pr.get("Alg"),
                "Parameter": pr.get("Parameter"),
                "초기 추천값": pr.get("초기 추천값"),
                "missing": pr.missing_units(),
                "value_diff": pr.has_value_diff(),
                "common": pr.is_common(),
                "values": {u: pr.get(u) for u in self.aoi_units},
            })
        return out

    # ---- 저장(병합) --------------------------------------------------------

    def save(self, user: str | None = None) -> dict[str, int]:
        """디스크 최신본을 다시 읽어 Row_ID 기준 병합 후 원자적으로 기록.

        반환: {"changes":n, "added":n, "deleted":n} 통계
        """
        user = user or current_user()
        changes = self.detect_changes()
        deleted = self.deleted_row_ids()
        added_ids = {pr.row_id for pr in self.rows if pr.row_id not in self._base}

        # 1) 디스크 최신 상태(다른 사람 수정 포함)를 병합 베이스로
        if os.path.exists(self.path):
            disk_wb = openpyxl.load_workbook(self.path, data_only=True)
            disk_rows = self._read_pi_all(disk_wb, self.sheet_name, self.aoi_units)
            prev_summary = self._read_summary(disk_wb)
            disk_wb.close()
        else:
            disk_rows, prev_summary = [], {}
        disk_by_id = {pr.row_id: pr for pr in disk_rows if pr.row_id}

        # 2) 내가 바꾼 셀만 디스크 상태 위에 적용 (행 단위 병합)
        for ch in changes:
            target = disk_by_id.get(ch.row_id)
            if target is not None:
                target.set(ch.field, ch.new)
            elif ch.row_id in added_ids:
                pr = next((p for p in self.rows if p.row_id == ch.row_id), None)
                if pr and pr.row_id not in disk_by_id:
                    # display_order 를 그대로 넘겨야 정렬 시 맨 위로 튀지 않음
                    new_pr = ParamRow(values=dict(pr.values), row_id=pr.row_id,
                                      display_order=pr.display_order, aoi_units=list(self.aoi_units))
                    disk_rows.append(new_pr)
                    disk_by_id[new_pr.row_id] = new_pr

        # 3) 내가 삭제한 행 제거
        if deleted:
            disk_rows = [pr for pr in disk_rows if pr.row_id not in deleted]
            disk_by_id = {pr.row_id: pr for pr in disk_rows}

        # 4) 정렬 및 Display_Order/파생열 재계산
        disk_rows.sort(key=lambda p: (pi_order(p.get("PI")), p.display_order or 0))
        for i, pr in enumerate(disk_rows, start=2):
            pr.display_order = i

        # 5) 변경이력 누적
        for ch in changes:
            self._append_history_record(user, ch)

        # 6) 워크북 작성 후 원자적 교체
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        self._write_pi_all(wb, disk_rows)
        self._write_summary(wb, disk_rows, prev_summary, changes, user, deleted)
        self._write_history(wb)
        self._write_snapshot(wb, disk_rows)
        self._write_special(wb)
        self._write_reference(wb)
        self._atomic_save(wb)

        # 7) 메모리 상태 갱신
        self.rows = disk_rows
        self._capture_base()
        return {"changes": len(changes), "added": len(added_ids), "deleted": len(deleted)}

    def _append_history_record(self, user: str, ch: ChangeRecord) -> None:
        pr = next((p for p in self.rows if p.row_id == ch.row_id), None)
        rec = {h: None for h in LOG_HEADERS}
        rec["Version"] = "v" + datetime.now().strftime("%Y%m%d_%H%M%S")
        rec["Update Date"] = today_str()
        rec["Updated By"] = user
        rec["Sheet"] = self.sheet_name
        if pr is not None:
            rec["Recipe"] = pr.get("Recipe")
            rec["Zone"] = pr.get("Zone")
            rec["Alg"] = pr.get("Alg")
            rec["Parameter"] = pr.get("Parameter")
        rec["AOI"] = ch.field if ch.is_aoi else ""
        rec["Old Value"] = "" if ch.old in (None, "") else ch.old
        rec["New Value"] = ch.new
        rec["Change Type"] = "AOI 값 수정" if ch.is_aoi else "메타 정보 수정"
        rec["Reason"] = "사용자 입력"
        rec["Source"] = "프로그램 자동화"
        rec["Comment"] = "기록 시간: " + now_str()
        self.history.append(rec)

    # ---- 시트 쓰기 ---------------------------------------------------------

    def _write_pi_all(self, wb, rows: list[ParamRow]) -> None:
        headers = pi_headers(self.aoi_units)
        fields = editable_fields(self.aoi_units)
        ws = wb.create_sheet(self.sheet_name)
        ws.append(headers)
        idx = {h: i for i, h in enumerate(headers)}
        for pr in rows:
            line: list[Any] = [None] * len(headers)
            for f in fields:
                line[idx[f]] = pr.get(f)
            line[idx["공통 여부"]] = "공통" if pr.is_common() else "미확인"
            line[idx["값 차이 여부"]] = "차이있음" if pr.has_value_diff() else "동일"
            line[idx["누락 호기"]] = ", ".join(pr.missing_units())
            line[idx["Recipe_Filter"]] = pr.get("Recipe")
            line[idx["Zone_Filter"]] = pr.get("Zone")
            line[idx["Alg_Filter"]] = pr.get("Alg")
            line[idx["SourceSheet"]] = self.sheet_name
            line[idx["Parameter_Key"]] = pr.parameter_key()
            line[idx["Row_ID"]] = pr.row_id
            line[idx["Display_Order"]] = pr.display_order
            ws.append(line)

    @staticmethod
    def _read_summary(wb) -> dict[str, dict[str, Any]]:
        """기존 요약 시트를 stable_key(row_id|aoi) 기준 dict 로 (집계 보존용)."""
        out: dict[str, dict[str, Any]] = {}
        if SHEET_SUM not in wb.sheetnames:
            return out
        ws = wb[SHEET_SUM]
        hdr = {_s(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}
        sk_col = hdr.get("Stable_Key")
        if not sk_col:
            return out
        for r in range(2, ws.max_row + 1):
            sk = _s(ws.cell(r, sk_col).value)
            if not sk:
                continue
            out[sk] = {h: ws.cell(r, hdr[h]).value for h in SUM_HEADERS if h in hdr}
        return out

    def _write_summary(self, wb, rows, prev_summary, changes, user, deleted) -> None:
        ws = wb.create_sheet(SHEET_SUM)
        ws.append(SUM_HEADERS)
        idx = {h: i for i, h in enumerate(SUM_HEADERS)}

        # 이번 저장에서 발생한 AOI 변경을 stable_key -> old 값으로
        aoi_changes: dict[str, Any] = {}
        for ch in changes:
            if ch.is_aoi:
                aoi_changes[f"{ch.row_id}|{ch.field}"] = ch.old

        live_ids = {pr.row_id for pr in rows}
        written_keys: set[str] = set()

        for pr in rows:
            for u in self.aoi_units:
                stable_key = f"{pr.row_id}|{u}"
                written_keys.add(stable_key)
                prev = prev_summary.get(stable_key, {})
                line: list[Any] = [None] * len(SUM_HEADERS)
                line[idx["PI"]] = pr.get("PI")
                line[idx["AOI"]] = u
                line[idx["Recipe"]] = pr.get("Recipe")
                line[idx["Zone"]] = pr.get("Zone")
                line[idx["Alg"]] = pr.get("Alg")
                line[idx["Parameter"]] = pr.get("Parameter")
                line[idx["초기 추천값"]] = pr.get("초기 추천값")
                line[idx["비고"]] = pr.get("비고")
                line[idx["현재값"]] = pr.get(u)

                first_v = prev.get("최초값")
                prev_v = prev.get("이전값")
                cnt = prev.get("변경횟수") or 0
                first_dt = prev.get("최초 변경일시")
                last_dt = prev.get("마지막 변경일시")
                last_user = prev.get("마지막 변경자")
                last_type = prev.get("마지막 변경유형")

                if stable_key in aoi_changes:
                    old_val = aoi_changes[stable_key]
                    if first_v in (None, ""):
                        first_v = old_val
                        first_dt = now_str()
                    prev_v = old_val
                    try:
                        cnt = int(cnt) + 1
                    except (TypeError, ValueError):
                        cnt = 1
                    last_dt = now_str()
                    last_user = user
                    last_type = "AOI 값 수정"

                line[idx["최초값"]] = first_v
                line[idx["이전값"]] = prev_v
                line[idx["변경횟수"]] = cnt
                line[idx["최초 변경일시"]] = first_dt
                line[idx["마지막 변경일시"]] = last_dt
                line[idx["마지막 변경자"]] = last_user
                line[idx["마지막 변경유형"]] = last_type
                line[idx["Source"]] = "PI_ALL 자동 연계"
                line[idx["Comment"]] = prev.get("Comment")
                line[idx["Summary_Key"]] = "|".join([
                    _s(pr.get("PI")), u, _s(pr.get("Recipe")), _s(pr.get("Zone")),
                    _s(pr.get("Alg")), _s(pr.get("Parameter")),
                ])
                line[idx["Source_Cell"]] = f"{self.sheet_name}!{u}"
                line[idx["Row_ID"]] = pr.row_id
                line[idx["Stable_Key"]] = stable_key
                line[idx["Display_Order"]] = pr.display_order
                line[idx["Status"]] = "Active"
                line[idx["PI_Order"]] = pi_order(pr.get("PI"))
                line[idx["AOI_Order"]] = aoi_order(u, self.aoi_units)
                ws.append(line)

        # 삭제된 행: 기존 요약을 Deleted 로 보존
        for stable_key, prev in prev_summary.items():
            if stable_key in written_keys:
                continue
            rid = stable_key.split("|")[0]
            if rid in live_ids:
                continue
            line = [prev.get(h) for h in SUM_HEADERS]
            line[idx["Status"]] = "Deleted"
            ws.append(line)

    def _write_history(self, wb) -> None:
        ws = wb.create_sheet(SHEET_LOG)
        ws.append(LOG_HEADERS)
        for rec in self.history:
            ws.append([rec.get(h) for h in LOG_HEADERS])

    def _write_snapshot(self, wb, rows: list[ParamRow]) -> None:
        ws = wb.create_sheet(SHEET_SNAP)
        ws.sheet_state = "hidden"
        ws.append(snap_headers(self.aoi_units))
        for pr in rows:
            line = [pr.row_id] + [pr.get(f) for f in META_FIELDS] + [pr.get(u) for u in self.aoi_units]
            ws.append(line)

    def _atomic_save(self, wb) -> None:
        folder = os.path.dirname(os.path.abspath(self.path)) or "."
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=folder)
        os.close(fd)
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        finally:
            if os.path.exists(tmp):
                try:
                    os.remove(tmp)
                except OSError:
                    pass

    # ---- 행 편집 헬퍼 ------------------------------------------------------

    def add_row(self, values: dict[str, Any] | None = None) -> ParamRow:
        fields = editable_fields(self.aoi_units)
        pr = ParamRow(values={f: None for f in fields}, row_id=new_row_id(),
                      aoi_units=list(self.aoi_units))
        if values:
            pr.values.update(values)
        pr.display_order = (max((p.display_order for p in self.rows), default=1) + 1)
        self.rows.append(pr)
        return pr

    def remove_row(self, row_id: str) -> None:
        self.rows = [pr for pr in self.rows if pr.row_id != row_id]


# --------------------------------------------------------------------------
# 새 빈 파일 생성 / 엑셀 가져오기·내보내기
# --------------------------------------------------------------------------

def create_empty_workbook(path: str) -> None:
    """헤더만 갖춘 새 공용 파일 생성."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet(SHEET_PI).append(PI_HEADERS)
    wb.create_sheet(SHEET_SUM).append(SUM_HEADERS)
    wb.create_sheet(SHEET_LOG).append(LOG_HEADERS)
    wb.create_sheet(SHEET_SPECIAL).append(SPECIAL_HEADERS)
    ref = wb.create_sheet(SHEET_REF)
    for i in range(1, 5):
        ref.append([f"표{i} 제목(여기에 입력)", "", "", ""])
        ref.append(["항목", "값", "", ""])
        ref.append(["", "", "", ""])
    snap = wb.create_sheet(SHEET_SNAP)
    snap.sheet_state = "hidden"
    snap.append(SNAP_HEADERS)
    wb.save(path)


def import_from_xlsm(src_path: str, dest_xlsx: str) -> ParamRepository:
    """기존 엑셀(.xlsm/.xlsx)에서 데이터를 읽어 새 .xlsx 로 변환 저장.
    시트 이름(PI_ALL/RDL_ALL 등)과 호기 열을 자동 인식한다."""
    wb = openpyxl.load_workbook(src_path, data_only=True)
    repo = ParamRepository(dest_xlsx)
    repo.sheet_name = detect_sheet_name(wb) or SHEET_PI
    repo.aoi_units = detect_aoi_units(wb, repo.sheet_name)
    repo.rows = ParamRepository._read_pi_all(wb, repo.sheet_name, repo.aoi_units)
    repo.history = ParamRepository._read_history(wb)
    repo.aoi_ip = ParamRepository._read_aoi_ip(wb)
    repo.special = ParamRepository._read_special(wb)
    repo.reference = ParamRepository._read_reference(wb)
    wb.close()
    repo._ensure_row_ids()
    repo._capture_base()
    # 변경이력 없이 그대로 기록 (변경감지 대상 0)
    out = openpyxl.Workbook()
    out.remove(out.active)
    repo._write_pi_all(out, repo.rows)
    repo._write_summary(out, repo.rows, {}, [], current_user(), [])
    repo._write_history(out)
    repo._write_snapshot(out, repo.rows)
    repo._write_special(out)
    repo._write_reference(out)
    out.save(dest_xlsx)
    repo._capture_base()
    return repo


def export_copy(repo_path: str, dest_path: str) -> None:
    """현재 공용 파일을 다른 경로로 복사(내보내기)."""
    shutil.copy2(repo_path, dest_path)
