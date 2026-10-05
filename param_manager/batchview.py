"""Batch Report 화면(가동률 조사 및 분석 · 찾기 · 취합)에 보낼 데이터 (헤드리스).

사용자가 프로토타입(`tools/batch_prototype`)으로 검토한 계산을 실제 앱으로 옮긴 것이다.
Lot 규칙은 `lotmodel` 을 그대로 쓰고, 여기서는 화면이 그릴 값만 만든다.

- `View(records, overrides)` — 조사 결과 1회분. `payload()` 가 화면에 한 번에 보내는 요약
  (Batch Report 요약 · Lot 요약 · 가동률 일 단위 · 재스캔 전 대기 · WPH 재료),
  `lot(li)` · `raw(g)` 는 누를 때만 보내는 상세(전송 한도 4MB 때문에 원문 행은 요약에 넣지 않는다).
- 시간 3칸(사용자 확정 2026-10-05): 호기마다 하루 24시간 = 웨이퍼 처리 + Error·중단 및 조치 + 유휴.
  · Error · 중단 없는 Batch Report = Batch Time 전체가 웨이퍼 처리. 그 외 = Pass 장수 × 1장 처리 시간(Batch Time 이내),
    나머지는 Error · 중단 손실. 같은 호기에서 같은 Lot 을 다시 스캔하기까지 장비가 아무것도 스캔하지 않은 시간도 손실(조치).
  · 1장 처리 시간 = Error · 중단 없이 25매 모두 Pass 한 Batch Report 의 Batch Time 합 ÷ 장수 합(호기 × 레시피 → 레시피).
  · 작업자 중단(앞 Error 없는 Aborted.)은 Error 가 아니지만 시간은 손실로 센다. 멈출 때 wafer 의 Faults 가 그 레시피
    정상 wafer Faults 상위 1% 보다 크면 Defect 과다 중단(정상 wafer MIN_FAULT_BASE 장 미만이면 판정하지 않음).
  · 점검 스캔은 웨이퍼 처리 시간(그중 점검)에만 넣고 레시피 WPH 에서는 뺀다. 자정을 넘긴 시간은 날마다 나눈다.
- WPH: 정상 = 정상 25매 Batch Report 장수 합 × 3600 ÷ Batch Time 합, 실제 = Pass 장수 × 3600 ÷ (웨이퍼 처리 + 손실)
  — 유휴만 뺀 시간. 같은 wafer 를 다시 스캔해 또 Pass 한 것도 장수에 넣는다(투입 · 출하 정보가 없어 필요한 재스캔인지
  판단 불가 — 재스캔 참고 비율로 표시).
- 중복 Pass 는 추천(가장 나중 Pass) + **저장된 사람 선택(overrides)** 으로만 정한다. 화면의 개발자 기능
  토글은 여기에 들어오지 않는다(켜고 끄는 것만으로 숫자가 바뀌면 안 된다 — CLAUDE.md).

파일 접근은 `write_excel`(로컬 결과 폴더, 경로는 호출자가 정함) 하나뿐이다.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, time, timedelta
from statistics import quantiles

from . import lotmodel as lm, lotreport

DUP = 'Pass 다시 스캔(중복)'          # 이미 Pass 한 wafer 를 다시 스캔한 몫
SPLIT = '오류 없이 나눠 스캔'          # Error 없이 같은 Lot 을 나눠 스캔한 사이 대기
FULL = lm.FULL_SLOTS
MIN_FAULT_BASE = 100                   # Defect 과다 판정에 필요한 정상 wafer 최소 수(사용자 확정 2026-10-05)


def stamp(value):
    return value.strftime('%Y-%m-%d %H:%M') if value else ''


def day_parts(start, end):
    """[start, end] 를 달력 날짜별로 나눈 [(YYYY-MM-DD, 비율)]. 길이가 없으면 시작일 하나."""
    if not start:
        return []
    if not end or end <= start:
        return [(start.strftime('%Y-%m-%d'), 1.0)]
    total = (end - start).total_seconds()
    out, cur = [], start
    while cur < end:
        stop = min(datetime.combine(cur.date() + timedelta(days=1), time.min), end)
        out.append((cur.strftime('%Y-%m-%d'), (stop - cur).total_seconds() / total))
        cur = stop
    return out


def _span(a):
    start = a['start']
    if not start:
        return None, None
    end = a['end'] if a['end'] and a['end'] > start else start + timedelta(seconds=a['batch_sec'] or 0)
    return start, end


def _reason(row):
    """wafer 행의 Error 원문(첫 문구). 앞 Error 뒤 연쇄(Aborted/Skipped)는 그 앞 Error 문구로."""
    return (row['phrase'] if row['kind'] == '원인' else row['trigger_phrase']) or '(빈 칸)'


def _first_error(a):
    """Error 가 있는 Batch Report 의 첫 Error 원문(작업자 중단 · 스캔 안 한 슬롯은 Error 가 아니라 '')."""
    if a.get('outcome') != lm.ERROR:
        return ''
    for row in a['rows']:
        if not row['pass'] and row['kind'] in ('원인', '연쇄'):
            return row['status'] or '(빈 칸)'
    return ''


def recipe_of(a):
    """레시피 = Job · Recipe(s)(3D 스캔은 ' · 3D 스캔'). WPH · 기준값의 단위."""
    return f"{a['job']} · {lm.step_label(a['step'], a['scan'])}"


def is_normal(a):
    """정상 25매 Batch Report — Error · 중단 없이 25매 모두 Pass(정상 WPH · 1장 처리 시간의 기준)."""
    return (a['outcome'] == lm.CLEAN and len(a['rows']) == FULL and all(r['pass'] for r in a['rows'])
            and (a['batch_sec'] or 0) > 0)


def _stats(a):
    s = dict(n=len(a['rows']), ok=0, err=0, chain=0, stop=0, skip=0)
    ph, cph = Counter(), Counter()
    for row in a['rows']:
        if row['pass']:
            s['ok'] += 1
        elif row['kind'] == lm.STOP:
            s['stop'] += 1
        elif row['kind'] == lm.SKIP:
            s['skip'] += 1
        elif row['kind'] == '연쇄':
            s['chain'] += 1
            cph[row['phrase'] or '(빈 칸)'] += 1
        else:
            s['err'] += 1
            ph[row['phrase'] or '(빈 칸)'] += 1
    return s, ph, cph


step_label = lm.step_label


def _num(value):
    return None if value is None else (int(value) if float(value).is_integer() else round(value, 3))


class View:
    """조사 1회분(또는 찾기 1회분)의 Lot 모델 + 화면용 계산."""

    def __init__(self, records, overrides=None, hits=None):
        self.records = list(records)
        self.overrides = dict(overrides or {})
        self.model = lm.build(self.records, self.overrides)
        by_id = {}
        for lot in self.model['lots']:
            for b in lot['bunches']:
                for a in b['attempts']:
                    by_id[a['id']] = a
        for a in self.model['excluded']:
            by_id[a['id']] = a
        self.A = [by_id[r['id']] for r in self.records]
        self.index = {a['id']: g for g, a in enumerate(self.A)}
        self.place = {}                                   # g → (lot, bunch, 위치)
        for li, lot in enumerate(self.model['lots']):
            for bi, b in enumerate(lot['bunches']):
                for p, a in enumerate(b['attempts']):
                    self.place[self.index[a['id']]] = (li, bi, p)
        self.hits = sorted(self.index[h] for h in (hits or ()) if h in self.index)
        self._bases()
        self._payload = None

    # ------------------------------------------------------------------ 요약(한 번에 보냄)
    def payload(self):
        if self._payload is None:
            self._payload = self._build()
        return self._payload

    def _build(self):
        R = []
        for g, (record, a) in enumerate(zip(self.records, self.A)):
            s, _, _ = _stats(a)
            li, bi, _ = self.place.get(g, (None, None, None))
            stop = self.stop_kind(a) if a['outcome'] == lm.STOP else None
            R.append(dict(f=a['file'], m=a['machine'], sm=a['sm'], code=a['code'] or '', job=a['job'], setup=a['setup'],
                          step=a['step'], k=a['scan'], s=stamp(a['start']), e=stamp(a['end']), sec=a['batch_sec'] or 0,
                          lot=li, b=bi, n=s['n'], ok=s['ok'], err=s['err'], chain=s['chain'], fe=_first_error(a),
                          st=s['stop'], sp=s['skip'], o={lm.CLEAN: 'c', lm.ERROR: 'e', lm.STOP: 's'}[a['outcome']],
                          sk=stop[0] if stop else '', ff=_num(a['stop_faults']), r=recipe_of(a)))
        lots = []
        for li, lot in enumerate(self.model['lots']):
            bunches, causes, per, stops = [], Counter(), {}, {}
            for b in lot['bunches']:
                bunches.append(dict(att=[self.index[a['id']] for a in b['attempts']], step=b['step'], k=b['scan'], s=stamp(b['start']),
                                    e=stamp(b['end']), moved=b['moved'], machines=b['machines'], sec=b['batch_sec']))
                for w in b['wafers']:
                    if w['stop']:
                        # 작업자 중단은 Error 가 아니다 — 종류(Defect 과다 / 그 외)별로 따로 센다.
                        k = self.stop_kind(b['attempts'][w['stop_attempt']])[0]
                        t = stops.setdefault(k, [0, False])
                        t[0] += 1
                        t[1] = t[1] or w['verdict'] == lm.UNRESOLVED
                        continue
                    if not w['cause']:
                        continue
                    causes[w['cause']] += 1
                    p = per.setdefault(w['cause'], [0, False])
                    p[0] += 1
                    p[1] = p[1] or w['verdict'] == lm.UNRESOLVED
            re_ = sum(w['verdict'] == lm.RECOVERED for b in lot['bunches'] for w in b['wafers'])
            saved = sum(w['overridden'] for b in lot['bunches'] for w in b['wafers'])
            lots.append(dict(key=lot['key'], label=lot['label'], code=lot['code'], lot_id=lot['lot_id'] or '',
                             sms=lot['sms'], jobs=lot['jobs'], machines=lot['machines'], n=lot['attempts'],
                             s=stamp(lot['start']), e=stamp(lot['end']), state=lot['state'], open=lot['unresolved'],
                             re=re_, dup=lot['duplicates'], moved=lot['moved'], saved=saved, scans=lot['scans'],
                             multi=any(len(b['attempts']) > 1 for b in lot['bunches']),
                             causes=[[k, n] for k, n in causes.most_common(3)],
                             cz=[[k, v[0], v[1]] for k, v in per.items()],
                             sz=[[k, v[0], v[1]] for k, v in stops.items()], bunches=bunches))
        starts = [r['s'] for r in R if r['s']]
        U, waits, stops, N, L, X = self._metrics()
        return dict(R=R, lots=lots, excluded=[self.index[a['id']] for a in self.model['excluded']],
                    U=U, waits=waits, stops=stops, N=N, L=L, X=X, hits=self.hits,
                    range=[min(starts)[:10], max(starts)[:10]] if starts else ['', ''],
                    saved=sum(lot['saved'] for lot in lots),
                    criteria=[[t, d, e] for t, d, e in lotreport.CRITERIA])

    # ------------------------------------------------------------------ 가동률 · WPH (사용자 확정 2026-10-05)
    def _bases(self):
        """레시피 기준값 — 1장 처리 시간(호기 × 레시피, 레시피) · Defect 과다 판정 기준(레시피별 정상 wafer Faults 상위 1%)."""
        n25, clean, faults = defaultdict(lambda: [0, 0.0]), defaultdict(lambda: [0, 0.0]), defaultdict(list)
        for lot in self.model['lots']:
            for b in lot['bunches']:
                for a in b['attempts']:
                    if a['outcome'] != lm.CLEAN or not a['batch_sec']:
                        continue
                    r, npass = recipe_of(a), sum(row['pass'] for row in a['rows'])
                    faults[r] += [row['faults'] for row in a['rows'] if row['pass'] and row['faults'] is not None]
                    for key in ((a['machine'], r), r):
                        if npass:
                            clean[key][0] += npass
                            clean[key][1] += a['batch_sec']
                        if is_normal(a):
                            n25[key][0] += FULL
                            n25[key][1] += a['batch_sec']
        self.unit_n25 = {k: s / w for k, (w, s) in n25.items() if w}
        self.unit_clean = {k: s / w for k, (w, s) in clean.items() if w}
        self.fault_base = {r: (quantiles(v, n=100)[98] if len(v) >= MIN_FAULT_BASE else None, len(v)) for r, v in faults.items()}

    def unit(self, a):
        """이 Batch Report 를 계산할 1장 처리 시간(초). 정상 25매(호기 × 레시피 → 레시피) → Error · 중단 없는 스캔 전체 순."""
        r = recipe_of(a)
        for table in (self.unit_n25, self.unit_clean):
            for key in ((a['machine'], r), r):
                if key in table:
                    return table[key]
        return None

    def stop_kind(self, a):
        """작업자 중단 Batch Report → ('d' Defect 과다 | 'o' 그 외, 기준값 또는 None(기준 wafer 부족))."""
        base, _ = self.fault_base.get(recipe_of(a), (None, 0))
        f = a['stop_faults']
        return ('d' if base is not None and f is not None and f > base else 'o'), base

    def split(self, a):
        """Batch Report 1장의 Batch Time → (웨이퍼 처리 초, Error · 중단 손실 초)."""
        sec = a['batch_sec'] or 0
        if a['outcome'] == lm.CLEAN:
            return float(sec), 0.0
        npass = sum(r['pass'] for r in a['rows'])
        t = self.unit(a)
        if t is None:
            # 그 레시피에 Error · 중단 없는 스캔이 하나도 없으면 Dice 있는 행 몫으로 나눈다(종전 방식).
            dice = sum(1 for r in a['rows'] if r['scanned'] is not None)
            t = sec / dice if dice else 0.0
        proc = min(float(sec), npass * t)
        return proc, sec - proc

    def _metrics(self):
        busy = defaultdict(list)
        for a in self.A:
            s, e = _span(a)
            if s and e and e > s:
                busy[a['machine']].append((s, e))
        for m, spans in busy.items():
            merged = []
            for s0, e0 in sorted(spans):
                if merged and s0 <= merged[-1][1]:
                    merged[-1] = (merged[-1][0], max(merged[-1][1], e0))
                else:
                    merged.append((s0, e0))
            busy[m] = merged

        def idle(machine, s0, e0):
            used = sum(max(0.0, (min(e0, b) - max(s0, a)).total_seconds()) for a, b in busy[machine] if a < e0 and b > s0)
            return max(0.0, (e0 - s0).total_seconds() - used)

        U = defaultdict(lambda: dict(p=0.0, du=0.0, ck=0.0, e=defaultdict(lambda: [0.0, 0, set()]),
                                     sd=[0.0, 0, set()], so=[0.0, 0, set()], ps=0, dn=0, n=0, ne=0, w25=0, s25=0.0, aw=0, asum=0.0))
        stops, waits, N, L = [], [], [], []
        dupn = Counter()
        for lot in self.model['lots']:
            for b in lot['bunches']:
                for w in b['wafers']:
                    passes = [c['attempt'] for c in w['cells'] if c['pass']]
                    for p in passes[1:]:
                        dupn[b['attempts'][p]['id']] += 1          # 이미 Pass 한 wafer 를 다시 스캔해 또 Pass

        def loss(day, a, li, sec, first):
            """Error · 중단 손실 sec 를 그 Batch Report 의 Error 원문 비율(또는 중단 종류)로 나눠 넣는다.

            first = 스캔 시작일 — 건수(중단 Batch Report 수 · Error wafer 수)는 그날 한 번만 센다.
            """
            u = U[(day, a['machine'], recipe_of(a))]
            if a['outcome'] == lm.STOP:
                t = u['sd' if self.stop_kind(a)[0] == 'd' else 'so']
                t[0] += sec
                t[2].add(li)
                t[1] += first
                return
            bad = Counter(_reason(r) for r in a['rows'] if not r['pass'] and r['kind'] in ('원인', '연쇄'))
            total = sum(bad.values()) or 1
            for reason, n in bad.items():
                t = u['e'][reason]
                t[0] += sec * n / total
                t[2].add(li)
                if first:
                    t[1] += n

        for li, lot in enumerate(self.model['lots']):
            for bi, b in enumerate(lot['bunches']):
                bunch_sec, bunch_pass = 0.0, 0
                for p, a in enumerate(b['attempts']):
                    start, end = _span(a)
                    g = self.index[a['id']]
                    npass = sum(r['pass'] for r in a['rows'])
                    if start and a['batch_sec']:
                        proc, lost = self.split(a)
                        bunch_sec += a['batch_sec']
                        bunch_pass += npass
                        for k, (day, frac) in enumerate(day_parts(start, end)):
                            u = U[(day, a['machine'], recipe_of(a))]
                            u['p'] += proc * frac
                            if npass:
                                u['du'] += dupn[a['id']] * proc / npass * frac
                            if a['outcome'] != lm.CLEAN:
                                loss(day, a, li, lost * frac, k == 0)
                            if k == 0:
                                u['n'] += 1
                                u['ps'] += npass
                                u['dn'] += dupn[a['id']]
                                u['ne'] += a['outcome'] == lm.ERROR
                                if is_normal(a):
                                    u['w25'] += FULL
                                    u['s25'] += a['batch_sec']
                                    if a['avg_scan_sec']:
                                        u['aw'] += FULL
                                        u['asum'] += a['avg_scan_sec'] * FULL
                        if is_normal(a):
                            N.append(dict(g=g, m=a['machine'], r=recipe_of(a), d=stamp(start)[:10], s=a['batch_sec'],
                                          a=a['avg_scan_sec'], lot=li))
                    if a['outcome'] == lm.STOP:
                        k, base = self.stop_kind(a)
                        stops.append(dict(g=g, m=a['machine'], r=recipe_of(a), li=li, f=_num(a['stop_faults']), k=k,
                                          base=None if base is None else round(base, 1), ps=npass))
                    if p + 1 >= len(b['attempts']) or a['outcome'] == lm.CLEAN:
                        continue
                    nxt = b['attempts'][p + 1]
                    if nxt['machine'] != a['machine'] or not a['end'] or not nxt['start'] or nxt['start'] <= a['end']:
                        continue
                    # Error · 중단 뒤 같은 호기에서 같은 Lot 을 다시 스캔하기까지 장비가 아무것도 스캔하지 않은 시간 = 조치.
                    kind = 's' if a['outcome'] == lm.STOP else 'e'
                    waits.append(dict(m=a['machine'], s=stamp(a['end']), e=stamp(nxt['start']), li=li, g=g, t=kind))
                    for day, _ in day_parts(a['end'], nxt['start']):
                        seg_s = max(a['end'], datetime.fromisoformat(day + ' 00:00'))
                        seg_e = min(nxt['start'], datetime.fromisoformat(day + ' 00:00') + timedelta(days=1))
                        gap = idle(a['machine'], seg_s, seg_e)
                        if gap:
                            loss(day, a, li, gap, False)
                            bunch_sec += gap
                if bunch_sec > 0:
                    L.append(dict(lot=li, b=bi, m=b['machines'][-1], r=f"{b['attempts'][-1]['job']} · {step_label(b['step'], b['scan'])}",
                                  d=stamp(b['start'])[:10], ps=bunch_pass, s=round(bunch_sec), n=len(b['attempts']),
                                  dn=sum(dupn[a['id']] for a in b['attempts'])))
        for a in self.model['excluded']:
            start, end = _span(a)
            if start and a['batch_sec']:
                for day, frac in day_parts(start, end):
                    U[(day, a['machine'], recipe_of(a))]['ck'] += a['batch_sec'] * frac
        rows = []
        for (d, m, r), u in sorted(U.items()):
            rows.append(dict(d=d, m=m, r=r, p=round(u['p']), du=round(u['du']), ck=round(u['ck']),
                             e={k: [round(x[0]), x[1], sorted(x[2])] for k, x in u['e'].items()},
                             sd=[round(u['sd'][0]), u['sd'][1], sorted(u['sd'][2])],
                             so=[round(u['so'][0]), u['so'][1], sorted(u['so'][2])],
                             ps=u['ps'], dn=u['dn'], n=u['n'], ne=u['ne'], w25=u['w25'], s25=u['s25'],
                             aw=u['aw'], asum=round(u['asum'])))
        X = dict(unit={f'{k[0]}||{k[1]}' if isinstance(k, tuple) else k: round(v, 1) for k, v in self.unit_n25.items()},
                 faults={r: [None if b is None else round(b, 1), n] for r, (b, n) in self.fault_base.items()},
                 min_base=MIN_FAULT_BASE, full=FULL)
        return rows, waits, stops, N, L, X


    # ------------------------------------------------------------------ 누를 때만 보내는 상세
    def lot(self, li):
        if type(li) is not int or not 0 <= li < len(self.model['lots']):
            raise ValueError('Lot 을 찾지 못했습니다. 조사를 다시 시작하세요.')
        lot = self.model['lots'][li]
        bunches, stats = [], {}
        for b in lot['bunches']:
            wafers = []
            for w in b['wafers']:
                cells, rec = [], None
                for c in w['cells']:
                    row = b['attempts'][c['attempt']]['rows'][c['order'] - 1]
                    cells.append([c['attempt'], c['status'] or '(빈 칸)', c['pass'], c['kind'] == '연쇄',
                                  _num(row['scanned']), _num(row['bad']), _num(row['good']), c['kind'] == lm.STOP])
                    if c['pass']:
                        rec = c['attempt']
                wafers.append(dict(k=w['key'], slot=w['slot'], id=w['wafer_id'], v=w['verdict'], pick=w['pick'],
                                   rec=rec, ov=w['overridden'], cause=w['cause'], chain=w['chain_only'], stop=w['stop'],
                                   cells=cells))
            bunches.append(dict(wafers=wafers))
            for a in b['attempts']:
                _, ph, cph = _stats(a)
                stats[self.index[a['id']]] = dict(ph=ph.most_common(), cph=cph.most_common())
        return dict(li=li, bunches=bunches, stats=stats)

    def raw(self, g):
        if type(g) is not int or not 0 <= g < len(self.records):
            raise ValueError('Batch Report 를 찾지 못했습니다. 조사를 다시 시작하세요.')
        record = self.records[g]
        report = record['report']
        headers = list(report['wafers'][0].keys()) if report.get('wafers') else []
        return dict(g=g, f=report.get('file_name', ''), m=record['machine'], folder=record.get('source_folder', ''),
                    meta=[list(kv) for kv in report.get('metadata', [])], h=headers,
                    rows=[[w.get(k, '') for k in headers] for w in report.get('wafers', [])])

    def choice_key(self, li, bi, wafer_key):
        """화면의 (Lot, 공정 단계 줄, wafer) → lotmodel overrides 키 (묶음 key, wafer key)."""
        try:
            b = self.model['lots'][li]['bunches'][bi]
        except (IndexError, TypeError):
            raise ValueError('Lot 을 찾지 못했습니다. 조사를 다시 시작하세요.') from None
        w = next((w for w in b['wafers'] if w['key'] == wafer_key), None)
        if w is None:
            raise ValueError('wafer 를 찾지 못했습니다. 조사를 다시 시작하세요.')
        return b, w

    # ------------------------------------------------------------------ 찾기 · 취합
    def aggregate(self, reports):
        """고른 Batch Report 를 Lot · 공정 단계별로 모아 wafer 마다 후보 행을 만든다.

        추천 = 저장된 사람 선택(그 Batch Report 를 골랐을 때) → 가장 나중 Pass → 가장 나중 스캔.
        """
        if not isinstance(reports, list) or not reports or len(reports) > 5000 or any(
                type(g) is not int or not 0 <= g < len(self.A) for g in reports):
            raise ValueError('취합할 Batch Report 를 고르세요')
        groups = {}
        for g in sorted(set(reports), key=lambda g: (self.A[g]['start'] or datetime.max, g)):
            if g not in self.place:
                continue                                   # 점검 스캔은 취합하지 않는다
            li, bi, _ = self.place[g]
            groups.setdefault((li, bi), []).append(g)
        out = []
        for (li, bi), att in groups.items():
            b = self.model['lots'][li]['bunches'][bi]
            chosen = {w['key']: b['attempts'][w['pick']]['id'] for w in b['wafers'] if w['overridden'] and w['pick'] is not None}
            id_slots = {r['wafer_id']: r['slot'] for g in att for r in self.A[g]['rows']
                        if r['slot'] is not None and lm.is_real_id(r['wafer_id'])}
            cells = defaultdict(list)
            for g in att:
                for r in self.A[g]['rows']:
                    if r['kind'] == lm.SKIP:
                        continue                           # 스캔 안 한 슬롯은 wafer 가 아니다
                    cells[lm._key(r, id_slots)].append(dict(g=g, status=r['status'] or '(빈 칸)', ok=r['pass'],
                                                            id=r['wafer_id'], sc=_num(r['scanned']), bad=_num(r['bad']),
                                                            good=_num(r['good'])))
            wafers = {}
            for key in sorted(cells, reverse=True):
                c = cells[key]
                rec = max((j for j, x in enumerate(c) if x['ok']), default=len(c) - 1)
                saved = chosen.get(key)
                hit = next((j for j, x in enumerate(c) if self.A[x['g']]['id'] == saved and x['ok']), None)
                wafers[key] = dict(cells=c, rec=hit if hit is not None else rec, auto=rec, saved=hit is not None)
            out.append(dict(li=li, bi=bi, att=att, keys=list(wafers), w=wafers))
        if not out:
            raise ValueError('Lot 으로 모인 Batch Report 를 골라 주세요(점검 스캔은 취합하지 않습니다).')
        return out


def write_excel(path, title, stamp_text, lot_rows, wafer_rows):
    """Lot 취합 · wafer 시트 두 장. 1행 = 선택 기준(사람이 결과의 기준을 알 수 있게)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    lot_head = ['Lot', '공정 단계(Recipe(s))', '기간', '호기', 'Batch Report', 'wafer', 'Pass', 'Pass 없음',
                'Scanned', 'Bad', 'Good', 'Yield (%)']
    wafer_head = ['Lot', '공정 단계(Recipe(s))', '슬롯', 'Wafer ID', '결과(쓰는 Batch Report 원문)', 'wafer 결과',
                  '쓰는 Batch Report', '호기', '스캔 시작', 'Scanned', 'Bad', 'Good', 'Yield (%)', '원본 폴더']
    wb = Workbook()
    wb.remove(wb.active)
    for name, head, rows in (('Lot 취합', lot_head, lot_rows), ('wafer', wafer_head, wafer_rows)):
        ws = wb.create_sheet(name)
        ws.append([f'{title} · 선택 기준: {stamp_text}'])
        ws.append(head)
        for row in rows:
            ws.append(list(row))
        for cells in ws.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = 's'           # 원문에서 온 글자는 수식으로 해석하지 않는다
        ws['A1'].font = Font(bold=True, color='183047')
        for cell in ws[2]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='183047')
        for i in range(1, len(head) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 22
        ws.freeze_panes = 'A3'
    wb.save(path)
