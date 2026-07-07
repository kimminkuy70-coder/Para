"""파라미터 이력 확인(3차 재설계) — 취합 파일 2개(멀티시트) 비교.

'파라미터 값 취합' 폴더의 취합 파일 2개를 골라, **달라진 부분만** 뽑아
새 창(표)·엑셀로 보여준다. 취합 파일은 레시피별 시트 구조이므로 시트별로 비교한다.

매칭 키 = (PI, Recipe(변형), norm Zone, norm Alg, norm Parameter). 한글 이름 보존.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from . import collate, engine

DIFF_SHEET = "변경내역"
DIFF_HEADERS = ["레시피(시트)", "PI", "Recipe", "Zone", "Alg", "Parameter", "호기",
                "이전 값", "새 값", "구분", "비고(메모)"]


def norm_key(name: str) -> str:
    return re.sub(r"[^0-9a-z가-힣µ]", "", str(name).lower())


@dataclass
class Change:
    sheet: str
    pi: str
    recipe: str
    zone: str
    alg: str
    param: str
    machine: str
    old: str
    new: str
    kind: str                # 값변경 / 추가 / 삭제


@dataclass
class HistoryDiff:
    changes: list[Change] = field(default_factory=list)
    added_rows: list[dict] = field(default_factory=list)
    removed_rows: list[dict] = field(default_factory=list)
    machines: list[str] = field(default_factory=list)


def _row_key(rd: dict) -> tuple:
    return (engine._s(rd.get("PI")).strip().lower(),
            engine._s(rd.get("Recipe")).strip().lower(),
            norm_key(rd.get("Zone")), norm_key(rd.get("Alg")),
            norm_key(rd.get("Parameter")))


def _meta(rd: dict, sheet: str) -> dict:
    return {"sheet": sheet, "PI": engine._s(rd.get("PI")),
            "Recipe": engine._s(rd.get("Recipe")), "Zone": engine._s(rd.get("Zone")),
            "Alg": engine._s(rd.get("Alg")), "Parameter": engine._s(rd.get("Parameter"))}


def diff_files(old_path: str, new_path: str) -> HistoryDiff:
    """이전(old) → 최신(new) 취합 파일 비교(레시피 시트별)."""
    old_sheets, om = collate.load_collation(old_path)
    new_sheets, nm = collate.load_collation(new_path)
    machines = list(dict.fromkeys(list(om) + list(nm)))
    diff = HistoryDiff(machines=machines)
    recipes = list(dict.fromkeys(list(old_sheets) + list(new_sheets)))
    for recipe in recipes:
        oidx = {_row_key(rd): rd for rd in old_sheets.get(recipe, [])}
        nidx = {_row_key(rd): rd for rd in new_sheets.get(recipe, [])}
        for k, nrd in nidx.items():
            ord_ = oidx.get(k)
            if ord_ is None:
                diff.added_rows.append(_meta(nrd, recipe))
                continue
            for m in machines:
                ov, nv = engine._s(ord_.get(m)), engine._s(nrd.get(m))
                if ov == nv:
                    continue
                kind = "추가" if ov == "" else "삭제" if nv == "" else "값변경"
                meta = _meta(nrd, recipe)
                diff.changes.append(Change(
                    sheet=recipe, pi=meta["PI"], recipe=meta["Recipe"],
                    zone=meta["Zone"], alg=meta["Alg"], param=meta["Parameter"],
                    machine=m, old=ov, new=nv, kind=kind))
        for k, ord_ in oidx.items():
            if k not in nidx:
                diff.removed_rows.append(_meta(ord_, recipe))
    return diff


def write_diff_excel(diff: HistoryDiff, dest_xlsx: str,
                     old_label: str = "", new_label: str = "",
                     memos: dict | None = None) -> str:
    """변경내역 엑셀 생성(비고 메모 열 포함)."""
    memos = memos or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = DIFF_SHEET
    ws.append(DIFF_HEADERS)
    for c in diff.changes:
        mkey = (c.sheet, c.param, c.machine)
        ws.append([c.sheet, c.pi, c.recipe, c.zone, c.alg, c.param, c.machine,
                   c.old, c.new, c.kind, engine._s(memos.get(mkey, ""))])
    fill = PatternFill("solid", fgColor="1F4E78")
    white = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = white
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJK", (14, 8, 12, 16, 14, 24, 10, 16, 16, 8, 34)):
        ws.column_dimensions[col].width = w
    fills = {"값변경": "FFF2CC", "추가": "D9EAD3", "삭제": "F4CCCC"}
    for r in range(2, ws.max_row + 1):
        k = ws.cell(r, 10).value
        if k in fills:
            f = PatternFill("solid", fgColor=fills[k])
            for cc in ws[r]:
                cc.fill = f

    sm = wb.create_sheet("행 추가·삭제")
    sm.append(["구분", "레시피(시트)", "PI", "Recipe", "Zone", "Alg", "Parameter"])
    for r in diff.added_rows:
        sm.append(["추가", r["sheet"], r["PI"], r["Recipe"], r["Zone"], r["Alg"], r["Parameter"]])
    for r in diff.removed_rows:
        sm.append(["삭제", r["sheet"], r["PI"], r["Recipe"], r["Zone"], r["Alg"], r["Parameter"]])

    info = wb.create_sheet("정보", 0)
    info.append(["항목", "값"])
    for k, v in [("이전(old)", old_label), ("최신(new)", new_label),
                 ("값 변경 셀", len(diff.changes)), ("행 추가", len(diff.added_rows)),
                 ("행 삭제", len(diff.removed_rows)),
                 ("안내", "'변경내역' 시트의 '비고(메모)'에 특이사항을 적으세요.")]:
        info.append([k, engine._s(v)])
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 60
    wb.save(dest_xlsx)
    return dest_xlsx
