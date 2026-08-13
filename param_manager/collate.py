"""파라미터 값 취합(3차 재설계) — 레시피별 시트 · 전체 호기 · 직전본 이어받기.

값 업데이트 흐름:
  - 취합할 레시피마다 `양식/{레시피}/` 최신 확정 양식을 기준으로,
  - 참고자료의 **모든 호기**를 열로 두고,
  - 이번에 수집한 장비 값으로 해당 호기 열만 갱신하며,
  - 이번에 수집하지 않은 호기 값은 **직전 취합본에서 이어받는다**(누적).
결과는 `파라미터 값 취합_{생성시간}.xlsx` 한 파일에 **레시피별 시트**로 저장한다.

매칭: 양식의 Parameter 는 사람이 바꿨을 수 있으므로 표시 이름이 아니라
`_EXTRACT_MAP`(설정 Section/Parameter)+Zone 기준(이름 변경에 견고), 없으면 이름 폴백.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from . import engine, extract_io, ini_parser, workdirs


def norm_key(name: str) -> str:
    """매칭용 정규화 — 한글 보존(사람이 붙인 한글 파라미터명 대응)."""
    return re.sub(r"[^0-9a-z가-힣µ]", "", str(name).lower())


def _safe_sheet(name: str) -> str:
    return re.sub(r"[\\/*?:\[\]]", "_", str(name))[:31] or "Recipe"


@dataclass
class CollateRecipe:
    recipe: str
    records: list[dict] = field(default_factory=list)     # META + 전체 호기 값
    machines: list[str] = field(default_factory=list)
    mismatches: list[dict] = field(default_factory=list)  # 이번 파싱값에 없는 양식 행
    matched_rows: int = 0
    filled_cells: int = 0
    missing_form: bool = False                            # 해당 레시피 양식이 없음
    carried: bool = False                                 # 이번에 안 고르고 직전본에서 유지


# --------------------------------------------------------------------------
# 파싱 인덱스 / 키
# --------------------------------------------------------------------------
def _ext_key(pi, variant, zone, section, param) -> tuple:
    return (engine._s(pi).strip().lower(), engine._s(variant).strip().lower(),
            norm_key(zone), norm_key(section), norm_key(param))


def _name_key(pi, variant, zone, alg, param) -> tuple:
    return (engine._s(pi).strip().lower(), engine._s(variant).strip().lower(),
            norm_key(zone), norm_key(alg), norm_key(param))


def _build_parsed_index(pivot_rows: list[dict]) -> tuple[dict, dict]:
    by_ext, by_name = {}, {}
    for r in pivot_rows:
        ext = r.get("extract") or {}
        by_ext.setdefault(_ext_key(r.get("recipe"), r.get("mag"), r.get("zone"),
                                   ext.get("section"), ext.get("key")), r)
        by_name.setdefault(_name_key(r.get("recipe"), r.get("mag"), r.get("zone"),
                                     r.get("alg"), r.get("param")), r)
    return by_ext, by_name


# --------------------------------------------------------------------------
# 취합(레시피 1개)
# --------------------------------------------------------------------------
def collate_recipe(recipe: str, form_path: str, pivot_rows: list[dict],
                   machines_all: list[str],
                   prev_values: dict | None = None,
                   coef_lookup=None) -> CollateRecipe:
    """양식(레시피 1개) + 파싱 + 전체 호기 + 직전값 → CollateRecipe.

    값은 **양식의 변환방식**(_EXTRACT_MAP transform)을 수집 raw 에 재적용해 채운다
    (사람이 양식에서 고친 변환방식·계수가 그대로 반영됨). 계수 우선순위:
      1) coef_lookup(호기, 변형)  — 호기별·변형별 변환계수.xlsx,
      2) 변환방식 라벨에 박힌 계수(AREA_0.77..^2),
      3) 기본(DEFAULT_SCALE).
    """
    repo = engine.ParamRepository(form_path)
    repo.load()
    emap = extract_io.read_extract_map(form_path)
    by_ext, by_name = _build_parsed_index(pivot_rows)
    prev_values = prev_values or {}
    res = CollateRecipe(recipe=recipe, machines=list(machines_all))

    for pr in repo.rows:
        level = engine._s(pr.get("PI"))
        variant = engine._s(pr.get("Recipe"))
        zone = engine._s(pr.get("Zone"))
        rkey = _name_key(level, variant, zone, pr.get("Alg"), pr.get("Parameter"))
        rec = {f: pr.get(f) for f in engine.META_FIELDS}
        # 1) 직전 취합본 값 이어받기(누적)
        for m, v in prev_values.get(rkey, {}).items():
            if m in machines_all and engine._s(v) != "":
                rec[m] = v
        # 2) 이번 수집값으로 갱신
        meta = emap.get(pr.row_id, {})
        match = None
        if meta.get("section") or meta.get("key"):
            match = by_ext.get(_ext_key(level, variant, zone,
                                        meta.get("section"), meta.get("key")))
        if match is None:
            match = by_name.get(rkey)
        if match is None:
            res.mismatches.append({
                "recipe": recipe, "zone": zone, "alg": engine._s(pr.get("Alg")),
                "param": engine._s(pr.get("Parameter")),
                "reason": "이번 수집 장비에서 설정키를 찾지 못함"})
        else:
            res.matched_rows += 1
            stored_t = engine._s(meta.get("transform")).strip() or "RAW"
            pname = engine._s(pr.get("Parameter"))
            # µ(마이크로) 규칙: 이름에 µ 있으면 **항상 변환**(이름 기준). 이름과 저장된
            # 변환방식이 어긋나도(옛 양식·지도 미매칭·LIGHT 고정 RAW) 일관되게 변환.
            # µ 없으면 저장된 변환방식 존중(수동 편집·BOOL/REGION/CLASSIFY/RAW).
            ftrans = (ini_parser.resolve_transform(pname, stored_t)
                      if ini_parser._has_micron(pname) else stored_t)
            label_coef = ini_parser.scale_from_label(stored_t)
            raws = match.get("raws") or {}
            values = match.get("values") or {}
            for m in machines_all:
                raw = raws.get(m)
                if engine._s(raw) == "":
                    v = values.get(m)                 # raw 없으면 파싱값 폴백
                else:
                    coef = None
                    if coef_lookup is not None:
                        try:
                            coef = coef_lookup(m, variant)   # (호기, 변형)
                        except Exception:  # noqa: BLE001
                            coef = None
                    if coef is None:
                        coef = label_coef
                    if coef is None:
                        coef = ini_parser.DEFAULT_SCALE
                    v = ini_parser.transform_value(raw, ftrans, coef)
                if engine._s(v) != "":
                    rec[m] = v
                    res.filled_cells += 1
        res.records.append(rec)
    return res


# --------------------------------------------------------------------------
# 하위 레시피(변형) 이름 매칭
#   장비마다 하위 레시피 폴더 이름이 다를 수 있다(2D+3D_CAMTEK / 2D+3D CAMTEK …).
#   변형은 행 키의 일부라 이름이 다르면 값이 다른 행으로 가서 채워지지 않는다.
#   그래서 값 업데이트 전에 **양식에 있는 이름과 매칭**할 수 있게 한다.
# --------------------------------------------------------------------------
def form_variants(form_path: str) -> list[str]:
    """확정 양식의 하위 레시피(Recipe 열) 목록 — 등장 순서."""
    repo = engine.ParamRepository(form_path)
    repo.load()
    out: list[str] = []
    for pr in repo.rows:
        v = engine._s(pr.get("Recipe")).strip()
        if v not in out:
            out.append(v)
    return out


def parsed_variants(pivot_rows: list[dict]) -> list[str]:
    """이번에 수집·파싱한 하위 레시피(변형) 목록 — 등장 순서."""
    out: list[str] = []
    for r in pivot_rows or []:
        v = engine._s(r.get("mag")).strip()
        if v not in out:
            out.append(v)
    return out


def unmatched_variants(pivot_rows: list[dict], form_path: str) -> list[str]:
    """양식에 없는(=이대로 두면 값이 안 채워지는) 수집 변형 목록.

    비교는 `norm_key` 정규화(대소문자·구분자 무시)라 `2D+3D_CAMTEK` 과
    `2d+3d camtek` 은 같은 것으로 본다. 정말 이름이 다른 것만 돌려준다.
    """
    known = {norm_key(v) for v in form_variants(form_path)}
    return [v for v in parsed_variants(pivot_rows) if norm_key(v) not in known]


def auto_variant_map(form_vars: list[str], parsed_vars: list[str]) -> dict:
    """수집 변형 → 양식 변형 **자동 매핑**(norm_key 일치 = 대소문자·구분자만 무시).

    이름이 정규화로 같은데 문자열만 다르면(예: `2d+3d camtek` → `2D+3D_CAMTEK`)
    수집 이름을 양식 이름으로 바꿔 준다. 정확히 같은 이름은 넣지 않는다(불필요).
    양식에 대응이 없는 수집 변형은 매핑에 없다(그대로 둠 = 안 채워짐).
    무인 회차(자동 감시)가 창 없이 안전하게 쓰는 범위다."""
    fmap: dict = {}
    for f in form_vars or []:
        fmap.setdefault(norm_key(f), engine._s(f).strip())
    out: dict = {}
    for p in parsed_vars or []:
        f = fmap.get(norm_key(p))
        if f is not None and f != engine._s(p).strip():
            out[engine._s(p).strip()] = f
    return out


def variant_match_table(form_vars: list[str], parsed_vars: list[str]) -> dict:
    """하위 레시피 이름 **매칭 확인창**용 표.

    왼쪽=양식 변형, 오른쪽 기본선택=자동매칭(norm_key)된 수집(복사) 변형.
    반환: {
      "rows": [(양식변형, 자동선택 수집변형 or "")],
      "parsed": [수집 변형 전체],                # 오른쪽 콤보 후보
      "unmatched_parsed": [어느 양식에도 안 붙는 수집 변형],   # 하단 안내(무시됨)
    }
    """
    pfirst: dict = {}
    for p in parsed_vars or []:
        pfirst.setdefault(norm_key(p), engine._s(p).strip())
    rows = [(engine._s(f).strip(), pfirst.get(norm_key(f), "")) for f in (form_vars or [])]
    fnorms = {norm_key(f) for f in (form_vars or [])}
    unmatched = [engine._s(p).strip() for p in (parsed_vars or [])
                 if norm_key(p) not in fnorms]
    return {"rows": rows, "parsed": [engine._s(p).strip() for p in (parsed_vars or [])],
            "unmatched_parsed": unmatched}


def apply_variant_map(pivot_rows: list[dict], mapping: dict) -> list[dict]:
    """수집 변형 → 양식 변형으로 이름을 바꿔 준다(사람이 매칭한 결과).

    mapping 값이 빈 문자열이면 **그 변형은 이번 취합에서 제외**한다
    (양식에 대응이 없는 레시피를 억지로 끼워 넣지 않기 위해).
    원본 리스트는 건드리지 않고 새 리스트를 돌려준다.
    """
    if not mapping:
        return list(pivot_rows or [])
    out: list[dict] = []
    for r in pivot_rows or []:
        v = engine._s(r.get("mag")).strip()
        if v in mapping:
            to = engine._s(mapping[v]).strip()
            if not to:
                continue                       # 무시(제외)
            r = dict(r)
            r["mag"] = to
        out.append(r)
    return out


def build_collation(save_dir: str, recipes: list[str], pivot_rows: list[dict],
                    machines_all: list[str],
                    prev_collate_path: str | None = None,
                    coef_lookup=None) -> dict[str, CollateRecipe]:
    """레시피별 취합 결과. 양식은 각 레시피의 최신 확정본에서 가져온다.
    coef_lookup(호기, 변형)→계수: 값 재적용 시 호기별·변형별 변환계수 적용(없으면 라벨/기본)."""
    prev = load_prev_values(prev_collate_path) if prev_collate_path else {}
    out: dict[str, CollateRecipe] = {}
    for recipe in recipes:
        form = workdirs.latest_form(save_dir, recipe)
        if not form:
            out[recipe] = CollateRecipe(recipe=recipe, machines=list(machines_all),
                                        missing_form=True)
            continue
        out[recipe] = collate_recipe(recipe, form, pivot_rows, machines_all,
                                     prev.get(recipe, {}), coef_lookup=coef_lookup)
    # 누적: 직전 취합본에 있던 **다른 레시피 시트**는 그대로 유지(이번에 안 고른 레시피가
    # 사라지지 않게). 이번에 고른 레시피는 위에서 새로 취합한 결과가 우선.
    if prev_collate_path:
        try:
            sheets, _ = load_collation(prev_collate_path)
        except Exception:  # noqa: BLE001
            sheets = {}
        for r, rows in sheets.items():
            if r not in recipes:
                out.setdefault(r, CollateRecipe(recipe=r, records=rows,
                                                machines=list(machines_all),
                                                carried=True))
    return out


# --------------------------------------------------------------------------
# 저장 / 로드(멀티시트)
# --------------------------------------------------------------------------
def delete_recipe(path: str, recipe: str) -> int:
    """취합 파일에서 특정 레시피(레벨)의 **시트를 삭제**한다(실제 엑셀 반영).
    시트 판별: 시트명 == _safe_sheet(recipe), 또는 그 시트의 데이터 PI 값이 전부 recipe.
    반환: 삭제한 시트 수. (시트가 0개가 되면 '취합없음' 빈 시트를 남긴다.)"""
    wb = openpyxl.load_workbook(path)
    target_title = _safe_sheet(recipe)
    lvl = engine._s(recipe).strip().lower()
    removed = 0
    for ws in list(wb.worksheets):
        if ws.title in _AUX_SHEETS:
            continue
        heads = [engine._s(c.value).strip() for c in ws[1]]
        if heads[:1] != ["PI"] or "Parameter" not in heads:
            continue
        pi_i = heads.index("PI")
        pis = {engine._s(row[pi_i]).strip().lower()
               for row in ws.iter_rows(min_row=2, values_only=True)
               if row and any(v not in (None, "") for v in row)}
        match = (ws.title == target_title) or (pis and pis <= {lvl})
        if match:
            wb.remove(ws)
            removed += 1
    if removed:
        if not [ws for ws in wb.worksheets if ws.title not in _AUX_SHEETS]:
            wb.create_sheet("취합없음")
        wb.save(path)
    wb.close()
    return removed


def write_collation(dest_xlsx: str, results: dict[str, CollateRecipe],
                    machines_all: list[str]) -> str:
    """레시피별 시트로 취합 파일 저장. 헤더 = META_FIELDS + 전체 호기."""
    headers = list(engine.META_FIELDS) + list(machines_all)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    fill = PatternFill("solid", fgColor="1F4E78")
    white = Font(color="FFFFFF", bold=True)
    for recipe, res in results.items():
        if res.missing_form:
            continue
        ws = wb.create_sheet(_safe_sheet(recipe))
        ws.append(headers)
        for rec in res.records:
            ws.append([rec.get(h) for h in headers])
        for c in ws[1]:
            c.fill = fill
            c.font = white
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
    if not wb.sheetnames:
        wb.create_sheet("취합없음")
    wb.save(dest_xlsx)
    return dest_xlsx


# 양식 파일(engine.create_from_records)이 남기는 부가 시트 — 취합 로드 시 제외.
# (요약/스냅샷/이력 시트도 'Parameter' 헤더가 있어 중복 로드 원인이 됨)
_AUX_SHEETS = {
    engine.SHEET_SUM, engine.SHEET_SNAP, engine.SHEET_LOG, engine.SHEET_SPECIAL,
    engine.SHEET_REF, engine.SHEET_RELATED, engine.SHEET_COLORS, engine.SHEET_BORDERS,
    extract_io.SHEET_MAP, extract_io.SHEET_SUMMARY,
}


def load_collation(path: str) -> tuple[dict[str, list[dict]], list[str]]:
    """취합 파일 → ({시트(레시피): [행dict]}, 호기목록). 행dict=헤더명→값.
    양식 파일의 부가 시트(요약/스냅샷/이력 등)는 제외해 중복 로드를 막는다."""
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets: dict[str, list[dict]] = {}
    machines: list[str] = []
    meta = set(engine.META_FIELDS)
    for ws in wb.worksheets:
        if ws.title in _AUX_SHEETS:
            continue
        heads = [engine._s(c.value).strip() for c in ws[1]]
        # 데이터 시트 판별: 앞부분이 META 스키마(PI·Zone·Alg·Parameter)로 시작
        if "Parameter" not in heads or "Zone" not in heads or "PI" not in heads:
            continue
        if heads[:1] != ["PI"]:              # 요약(PI,AOI…)·스냅샷(Row_ID…) 제외
            continue
        for h in heads:
            if h and h not in meta and h not in machines:
                machines.append(h)
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(v in (None, "") for v in row):
                continue
            rows.append({heads[i]: row[i] for i in range(len(heads)) if i < len(row)})
        sheets[ws.title] = rows
    wb.close()
    return sheets, machines


def load_prev_values(path: str | None) -> dict[str, dict]:
    """직전 취합본 → {레시피: {row_key: {호기: 값}}} (이어받기용)."""
    if not path:
        return {}
    try:
        sheets, machines = load_collation(path)
    except Exception:  # noqa: BLE001
        return {}
    out: dict[str, dict] = {}
    for recipe, rows in sheets.items():
        rmap = {}
        for rd in rows:
            k = _name_key(rd.get("PI"), rd.get("Recipe"), rd.get("Zone"),
                          rd.get("Alg"), rd.get("Parameter"))
            rmap[k] = {m: rd.get(m) for m in machines if engine._s(rd.get(m)) != ""}
        out[recipe] = rmap
    return out


def load_as_repo(path: str, machines_all: list[str]) -> engine.ParamRepository:
    """취합 파일(멀티시트) → 값 확인 화면용 in-memory ParamRepository(모든 시트 병합).

    호기 열 = 참고자료(machines_all) + **취합 파일에 실제로 있는 호기 열의 합집합**.
    (참고자료에 없는 호기로 취합된 값이 화면에서 빠지던 버그 방지 — 파일의 값은 항상 표시.)
    """
    sheets, file_machines = load_collation(path)
    machines = list(machines_all) + [m for m in file_machines if m not in machines_all]
    repo = engine.ParamRepository(path)
    repo.aoi_units = machines
    repo.rows = []
    i = 0
    for recipe, rows in sheets.items():
        for rd in rows:
            vals = {f: rd.get(f) for f in engine.META_FIELDS}
            for m in machines:
                vals[m] = rd.get(m)
            repo.rows.append(engine.ParamRow(
                values=vals, row_id=engine.new_row_id(), display_order=i + 2,
                aoi_units=list(machines)))
            i += 1
    repo.cell_colors = {}
    return repo
