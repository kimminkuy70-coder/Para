"""Batch Report status summaries; no additional file/network reads."""
import re
from collections import Counter

TYPES = (
    ('맵 Import 오류', 'Wafer Map Import failed.'),
    ('2D Scan 오류', 'Scan 2D Error.'),
    ('Alignment 오류', 'Alignment Error.'),
    ('Wafer ID 판독 오류', 'Failed to read wafer id.'),
    ('검사 대상 없음', 'Nothing to Scan.'),
    ('작업 중단', 'Aborted.'),
    ('검사 제외', 'Skipped.'),
)


def classify(raw):
    text = re.sub(r'\s+', ' ', str(raw or '')).strip()
    if not text:
        return [('상태 확인 불가', '(빈 상태)')]
    if re.fullmatch(r'(pass|passed|ok|success|successful|completed|normal)[.!]?', text, re.I):
        return []
    found = []
    rest = text
    for category, status in TYPES:
        pattern = r'\b' + re.escape(status.rstrip('.')) + r'\b\.?'
        if re.search(pattern, text, re.I):
            found.append((category, status))
            rest = re.sub(pattern, '', rest, flags=re.I)
    if rest.strip(' .;,+/|\t\n'):
        found.append(('기타 상태', rest.strip(' .;,+/|\t\n')))
    return found


def analyze(rows, errors):
    summary, details, reports = {}, [], []
    counts = Counter()
    for row in rows:
        machine, filename = row.get('machine', ''), row.get('source_file', '')
        statuses = row.get('wafer_statuses') or []
        occurrences = Counter()
        if not statuses:
            statuses = [{'status': '', 'slot': '', 'wafer_id': ''}]
        for status in statuses:
            for category, original in classify(status.get('status')):
                occurrences[(category, original)] += 1
                details.append([machine, filename, status.get('slot', ''),
                                status.get('wafer_id', ''), category, original,
                                status.get('status', '')])
        issues = [key for key in occurrences if key[0] != '상태 확인 불가']
        state = '이슈 포함' if issues else '상태 확인 불가' if occurrences else '정상 상태만'
        counts[state] += 1
        reports.append([machine, filename, state, len(row.get('wafer_statuses') or []),
                        ' / '.join(dict.fromkeys(k[0] for k in occurrences))])
        for key, n in occurrences.items():
            tally = summary.setdefault(key, [0, 0])
            tally[0] += 1
            tally[1] += n if key[0] != '상태 확인 불가' or row.get('wafer_statuses') else 0
    return counts, summary, reports, details


def add_sheets(wb, rows, errors):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    counts, summary, reports, details = analyze(rows, errors)
    overview = wb.create_sheet('07_상태요약')
    overview.append(['Batch Report 상태 조사', '건수'])
    for label, number in [('전체 Batch Report', len(rows) + len(errors)),
                          ('이슈 상태 포함', counts['이슈 포함']),
                          ('정상 상태만', counts['정상 상태만']),
                          ('상태 확인 불가', counts['상태 확인 불가']),
                          ('원본 Report 파싱 오류', len(errors))]:
        overview.append([label, number])
    overview.append(['유형별 Report 수는 중복 집계됩니다. 기타 상태는 원문을 확인하세요.'])
    overview.append(['생성 시점 집계입니다. Raw Data를 수정해도 상태 통계는 바뀌지 않습니다.'])
    overview.append(['대분류', '원본 상태', 'Report 수', 'Wafer/Slot 발생 건수'])
    ordered = list(TYPES) + [k for k in summary if k not in TYPES]
    for key in ordered:
        overview.append([*key, *summary.get(key, [0, 0])])
    report_sheet = wb.create_sheet('08_Report목록')
    report_sheet.append(['호기', 'Report 파일명', '상태 구분', 'Wafer/Slot 행 수', '포함 유형'])
    for record in reports:
        report_sheet.append(record)
    detail_sheet = wb.create_sheet('09_상태상세')
    detail_sheet.append(['호기', 'Report 파일명', 'Slot/No', 'Wafer ID', '대분류', '분류된 원본 상태', '실제 상태 원문'])
    for record in details:
        detail_sheet.append(record)
    error_sheet = wb.create_sheet('10_파싱오류')
    error_sheet.append(['호기', 'Report 파일명', '오류 내용'])
    for error in errors:
        error_sheet.append([error.get('machine', ''), error.get('source_file', ''), error.get('error', '')])
    for sheet in (overview, report_sheet, detail_sheet, error_sheet):
        for row in sheet:
            for cell in row:
                # Report names/statuses are untrusted text, never Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        header_rows = [1, 9] if sheet is overview else [1]
        for r in header_rows:
            for cell in sheet[r]:
                cell.fill = PatternFill('solid', fgColor='234863')
                cell.font = Font(color='FFFFFF', bold=True)
        sheet.freeze_panes = 'C10' if sheet is overview else 'C2'
        sheet.auto_filter.ref = f'A9:D{sheet.max_row}' if sheet is overview else sheet.dimensions
        for column, width in [('A', 24), ('B', 72), ('C', 24), ('D', 25), ('E', 32), ('F', 38), ('G', 55)]:
            sheet.column_dimensions[column].width = width
    # Derive both frequency series from the exact rows written to 09_상태상세.
    frequency, report_ids = Counter(), {}
    for machine, filename, slot, wafer_id, category, original, raw in details:
        if category == '상태 확인 불가':
            continue
        key = (category, original)
        frequency[key] += 1
        report_ids.setdefault(key, set()).add((machine, filename))
    chart_rows = [[f'{category} / {original}', len(report_ids[(category, original)]), count]
                  for (category, original), count in sorted(frequency.items(), key=lambda item: -item[1])]
    if not chart_rows:
        chart_rows = [['상세에 이슈 상태 없음', 0, 0]]
    for c, label in enumerate(['09 상태상세 기준', 'Report 수 (중복 제외)', 'Wafer/Slot 발생 건수'], 9):
        overview.cell(1, c, label)
    for r, values in enumerate(chart_rows, 2):
        for c, value in enumerate(values, 9):
            cell = overview.cell(r, c, value)
            if isinstance(value, str):
                cell.data_type = 's'
    # Native charts over concrete numeric cells: no dynamic formula/cached value dependency.
    data_end = 9 + len(ordered)
    overview.cell(data_end + 2, 1, '오류·중단·제외 상태 빈도 그래프 ↓ (유형별 중복 포함)')
    from .wph_charts import format_chart, value_labels
    chart_index = 0
    for column, title in [(10, '유형별 Report 수'), (11, '유형별 Wafer/Slot 발생 빈도')]:
        # Long state labels get at most 12 categories per large chart.
        for offset in range(0, len(chart_rows), 12):
            first = 2 + offset
            last = min(offset + 12, len(chart_rows)) + 1
            chart = BarChart()
            chart.type = 'bar'
            chart.style = 10
            chart.title = title + (f' ({offset // 12 + 1})' if len(chart_rows) > 12 else '')
            format_chart(chart, horizontal=True)
            chart.height = 20
            chart.add_data(Reference(overview, min_col=column, min_row=first, max_row=last))
            chart.set_categories(Reference(overview, min_col=9, min_row=first, max_row=last))
            value_labels(chart)
            chart.y_axis.scaling.orientation = 'maxMin'
            anchor_row = data_end + 4 + chart_index * 44
            for r in range(anchor_row, anchor_row + 44):
                overview.row_dimensions[r].height = 18
            overview.add_chart(chart, f'A{anchor_row}')
            chart_index += 1
